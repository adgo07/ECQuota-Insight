"""UEBench 进程入口：桌面界面与无头自检（ECQ-RS05 §14）。

顶层**不得**导入 PySide6：``--self-check`` 必须在完全没有 Qt、没有显示器、
也不需要 ``QT_QPA_PLATFORM=offscreen`` 的前提下完成。Qt 只在真正要启动桌面
窗口的 ``_run_desktop_application()`` 里延迟导入。

CLI 契约：

.. code-block:: text

    UEBench.exe                                    启动桌面界面（原行为）
    UEBench.exe --self-check --data-dir <DIR> [--output <JSON>]
                                                   无头自检

退出码：``0`` 全部检查通过；``1`` 至少一项检查失败；``2`` 用法/参数错误。
详见 ``uebench.application.self_check`` 的模块文档。
"""

from __future__ import annotations

import argparse
import json
import logging
import sys
from collections.abc import Mapping
from pathlib import Path


def configure_application_font(application: "QApplication") -> None:  # noqa: F821
    """使用系统中文字体配置 Qt 应用字体（仅在桌面路径调用）。"""
    from PySide6.QtGui import QFont, QFontDatabase

    candidates = [
        Path("C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/simhei.ttf"),
    ]
    for candidate in candidates:
        if not candidate.exists():
            continue
        font_id = QFontDatabase.addApplicationFont(str(candidate))
        families = QFontDatabase.applicationFontFamilies(font_id)
        if families:
            application.setFont(QFont(families[0], 10))
            return


def bundled_package_directory() -> Path:
    """内置标准包所在目录。

    PyInstaller 会把 ``release/standard-packages/initial-standard-package-published.uebench``
    收集到 ``uebench/resources``（见 ``uebench.spec``）；打包运行时模块的
    ``__file__`` 已位于 ``<_MEIPASS>/uebench/``，因此与源码运行使用同一相对位置。
    """
    return Path(__file__).resolve().parent / "resources"


def reconcile_standard_package(context, bundled_directory: Path | None = None):
    """ECQ-RS05 启动对账：每次正常启动都比对内置包与已安装包（原一次性初始化）。

    决策与执行都在 application 层（``uebench.application.package_reconciliation``）；
    组合根只负责注入唯一版本排序实现与内置包目录。返回值同时记录到
    ``context.application``，供界面与自检读取。

    真正的 I/O / 校验异常会向上传播（``StandardPackageService.install`` 已保证
    失败回滚、不留半安装状态），由调用方决定如何提示。
    """
    from .application.package_reconciliation import PackageReconciliationService
    from .infrastructure.packages import _data_version_key

    service = PackageReconciliationService(
        context.package_service,
        data_version_key=_data_version_key,
    )
    outcome = service.reconcile(bundled_directory or bundled_package_directory())
    context.application.record_package_reconciliation(outcome)
    return outcome


# ---------------------------------------------------------------------------
# 无头自检的适配器（组合根职责：application 层被禁止直接依赖这些实现）
# ---------------------------------------------------------------------------


def create_self_check_context(data_dir: Path):
    """为自检构造真实应用上下文（真实 DB / 仓储 / ApplicationFacade）。"""
    from .application.self_check import locate_public_key
    from .bootstrap import create_context

    return create_context(data_dir, public_key_path=locate_public_key())


def write_workbook_rows(path: Path, sheet_name: str, rows: Mapping[int, object]) -> None:
    """把 ``rows``（行号 → 值）写入 ``sheet_name`` 表的 B 列。

    这是自检使用的 Excel 填写适配器；它只负责“怎么写”，业务值由
    ``uebench.application.self_check`` 决定。
    """
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    try:
        sheet = workbook[sheet_name]
        for row, value in rows.items():
            sheet.cell(row=int(row), column=2, value=value)
        workbook.save(path)
    finally:
        workbook.close()


def check_output_writable(path: Path) -> str | None:
    """返回不可写原因；可写时返回 ``None``。

    在真正执行检查前先验证，使 ``--output`` 不可写能以用法错误（退出码 2）
    结束，而不是在检查跑完很久之后才失败。
    """
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        with path.open("a", encoding="utf-8"):
            pass
    except OSError as exc:
        return f"{type(exc).__name__}: {exc}"
    return None


def _write_text(stream, text: str) -> None:
    """向可能为 ``None`` 的标准流写文本（windowed EXE 没有控制台）。"""
    if stream is None:
        return
    try:
        stream.write(text)
        stream.flush()
    except Exception:  # pragma: no cover - 控制台编码/管道异常不应改变退出码
        pass


def _summary_text(report: Mapping[str, object]) -> str:
    lines = [
        f"UEBench 自检：{report['overall_status']}（exit_code={report['exit_code']}）",
        f"产品版本 {report['product_version']}，数据目录 {report['data_dir']}",
    ]
    for check in report["checks"]:  # type: ignore[index]
        lines.append(f"  [{check['status']}] {check['id']} {check['title']}")
        for error in check["errors"]:
            lines.append(f"        失败：{error}")
    for warning in report["warnings"]:  # type: ignore[index]
        lines.append(f"  警告：{warning}")
    return "\n".join(lines) + "\n"


def run_self_check_command(data_dir: Path, output: Path | None) -> int:
    """执行 ``--self-check``；返回进程退出码。"""
    from .application.self_check import EXIT_CHECKS_FAILED, run_self_check
    from .infrastructure.logging import close_logging

    try:
        report = run_self_check(
            data_dir=data_dir,
            context_factory=create_self_check_context,
            workbook_writer=write_workbook_rows,
        )
    finally:
        # 自检使用隔离数据目录内的滚动日志；Windows 上必须先释放日志句柄，
        # 调用方才能安全删除该目录。
        close_logging()

    document = json.dumps(report, ensure_ascii=False, indent=2) + "\n"
    if output is None:
        _write_text(sys.stdout, document)
    else:
        try:
            output.parent.mkdir(parents=True, exist_ok=True)
            output.write_text(document, encoding="utf-8")
        except OSError as exc:
            _write_text(sys.stderr, f"自检报告写入失败：{output}（{exc}）\n")
            return EXIT_CHECKS_FAILED
    _write_text(sys.stdout, _summary_text(report))
    return int(report["exit_code"])


# ---------------------------------------------------------------------------
# 桌面路径
# ---------------------------------------------------------------------------


def _run_desktop_application() -> int:
    from PySide6.QtGui import QIcon
    from PySide6.QtWidgets import QApplication

    from .bootstrap import create_context
    from .infrastructure.logging import close_logging
    from .ui.main_window import create_main_window

    application = QApplication.instance() or QApplication(sys.argv)
    application.setApplicationName("单位产品能耗对标软件")
    application.setOrganizationName("UEBench")
    bundled_icon = Path(__file__).resolve().parent / "resources" / "uebench.ico"
    if bundled_icon.exists():
        application.setWindowIcon(QIcon(str(bundled_icon)))
    configure_application_font(application)
    resource_key = Path(__file__).resolve().parent / "resources" / "update_public_key.pem"
    context = create_context(public_key_path=resource_key)
    try:
        try:
            reconcile_standard_package(context)
        except Exception:
            # 对账失败绝不能阻止界面启动：standards 安装已由 install 自行回滚，
            # 存量数据保持可用（旧行为同样是记录日志后继续启动）。
            logging.getLogger(__name__).exception("内置标准包对账失败")
        window = create_main_window(context)
        window.show()
        return application.exec()
    finally:
        context.database.dispose()
        close_logging()


# ---------------------------------------------------------------------------
# 参数解析（必须在任何 Qt 导入之前完成）
# ---------------------------------------------------------------------------


def build_argument_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        prog="UEBench",
        description="单位产品能耗对标软件（UEBench）",
    )
    parser.add_argument(
        "--self-check",
        action="store_true",
        help="运行无头自检并写出 JSON 报告（不启动界面，不导入 Qt）",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=None,
        metavar="PATH",
        help="自检 JSON 报告输出路径；省略时写到标准输出",
    )
    parser.add_argument(
        "--data-dir",
        type=Path,
        default=None,
        metavar="DIR",
        help="自检使用的隔离数据目录（不会访问 %%LOCALAPPDATA%%\\UEBench）",
    )
    return parser


def main(argv: list[str] | None = None) -> int:
    """进程入口；``argv`` 默认取 ``sys.argv[1:]``。"""
    parser = build_argument_parser()
    arguments = parser.parse_args(sys.argv[1:] if argv is None else argv)

    if arguments.self_check:
        if arguments.data_dir is None:
            parser.error("--self-check 必须同时指定 --data-dir（隔离数据目录）")
        if arguments.output is not None:
            problem = check_output_writable(arguments.output)
            if problem is not None:
                parser.error(f"--output 不可写：{arguments.output}（{problem}）")
        return run_self_check_command(arguments.data_dir, arguments.output)

    if arguments.data_dir is not None or arguments.output is not None:
        parser.error("--data-dir / --output 仅在 --self-check 模式下有效")
    return _run_desktop_application()


if __name__ == "__main__":
    raise SystemExit(main())
