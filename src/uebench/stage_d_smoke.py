"""候选专用实包验收；预期取独立手算，不把Calculator输出作为oracle。"""
from __future__ import annotations

from decimal import Decimal
import hashlib
import json
import os
from pathlib import Path
import sqlite3
import subprocess
import sys
from uuid import uuid4

from .domain.models import (CementDetails, CementFuelLine, EvaluationRequest,
    EvaluationResult, InputMode, InputValue, StandardDefinition)

EXPECTED = {
    'cement.comprehensive_energy': ('80.95225', 'LEVEL_2'),
    'clinker.comprehensive': ('102.6702', 'LEVEL_2'),
    'clinker.coal': ('96.771', 'LEVEL_2'), 'clinker.electricity': ('48', 'LEVEL_1'),
    'cement-preparation.electricity': ('23', 'LEVEL_1')}
MODE_IDS = {'integrated': set(EXPECTED), 'partial_purchased': set(EXPECTED),
    'clinker_only': {'clinker.comprehensive', 'clinker.coal', 'clinker.electricity'},
    'grinding_station': {'cement-preparation.electricity'}}


def require(condition, message):
    # PyInstaller optimize=1 removes assert statements; these checks must execute.
    if not condition:
        raise RuntimeError(message)


def sample_request(mode='integrated', *, dual=False):
    fields = {
        'production_mode': (mode,None), 'evaluation_period': ('2025年',None), 'production_line_id': ('A',None),
        'altitude_m': ('1600' if dual else '1500','m'), 'clinker_production_t': ('1000','t'),
        'clinker_consumption_t': ('750','t'), 'cement_production_t': ('1000','t'),
        'clinker_ratio_percent': ('80' if dual else '75','%'), 'clinker_electricity_kwh': ('50000','kW·h'),
        'waste_heat_generation_kwh': ('12000','kW·h'), 'waste_heat_self_use_kwh': ('2000','kW·h'),
        'cement_grinding_packaging_kwh': ('20000','kW·h'), 'admixture_preparation_kwh': ('1000','kW·h'),
        'gypsum_preparation_kwh': ('1000','kW·h'), 'auxiliary_electricity_kwh': ('2000','kW·h'),
        'cement_fuel_kgce': ('1000','kgce'), 'grinding_electricity_kwh': ('23000','kW·h'),
        'deduction.co_processing_fuel_tce': ('1','tce'), 'deduction.alternative_fuel_tce': ('1','tce'),
        'deduction.co_processing_electricity_kwh': ('1000','kW·h'),
        'deduction.alternative_fuel_electricity_kwh': ('1000','kW·h')}
    if mode == 'clinker_only':
        for key in ['clinker_consumption_t','cement_production_t','clinker_ratio_percent',
            'cement_grinding_packaging_kwh','admixture_preparation_kwh','gypsum_preparation_kwh',
            'auxiliary_electricity_kwh','cement_fuel_kgce','grinding_electricity_kwh']:
            fields.pop(key)
    elif mode == 'grinding_station':
        fields = {key:value for key,value in fields.items() if key in {
            'production_mode','evaluation_period','production_line_id','cement_production_t','grinding_electricity_kwh'}}
    if dual:
        fields['pressure_pa'] = ('81060','Pa')
    return EvaluationRequest(evaluation_date='2026-10-29',standard_id='gb-16780-2021',product_id='cement',
        input_mode=InputMode.DETAIL, inputs={key:InputValue(value=value,unit=unit) for key,(value,unit) in fields.items()},
        cement_details=CementDetails(fuel_lines=[] if mode=='grinding_station' else [CementFuelLine(
            line_id='f1',name='燃煤',amount_kg='100000',low_heat_kj_per_kg='29307.6')]))


def check_result(result, mode, dual=False):
    items = {item.indicator_id:item for item in result.results if item.grade.value!='NOT_APPLICABLE'}
    require(set(items)==MODE_IDS[mode], '适用指标集合错误')
    expected = dict(EXPECTED)
    if dual:
        expected['cement.comprehensive_energy'] = ('85.7908','LEVEL_2')
    for key,item in items.items():
        value,grade = expected[key]
        require(item.actual_value==Decimal(value) and item.grade.value==grade, f'{key}独立预期不符')
    if mode in ('integrated','partial_purchased'):
        require(items['cement.comprehensive_energy'].display_values['cement_electricity']==Decimal('60'), '中间量W_s错误')
    if dual:
        require(items['cement.comprehensive_energy'].corrected_thresholds=={
            'LEVEL_1':Decimal('85.50'),'LEVEL_2':Decimal('92.75'),'LEVEL_3':Decimal('100.9588')}, '双修正阈值错误')
        for key,upper in [('clinker.comprehensive','118.1934'),('clinker.coal','110.1118'),('clinker.electricity','61.6222'),('cement-preparation.electricity','34')]:
            require(items[key].corrected_thresholds['LEVEL_3']==Decimal(upper), '有效三级阈值错误')
    return {key:{'actual':str(item.actual_value),'grade':item.grade.value} for key,item in items.items()}


def digests(row):
    return [hashlib.sha256(value.encode('utf-8')).hexdigest() for value in row]


def readback(root: Path) -> int:
    """独立进程只读原JSON及投影，不构造Calculator、不启动规则对账。"""
    from .application.cement_reporting import build_cement_report_view
    from .domain.engine import EvaluationEngine
    ledger = json.loads((root/'smoke-ledger.json').read_text(encoding='utf-8'))
    with sqlite3.connect((root/'uebench.sqlite3').as_uri()+'?mode=ro',uri=True) as connection:
        connection.execute('PRAGMA query_only=ON')
        require(connection.execute('SELECT COUNT(*) FROM evaluations').fetchone()[0]==len(ledger['records']), '记录数量变化')
        for record in ledger['records']:
            row = connection.execute('SELECT request_json,result_json,rule_snapshot_json FROM evaluations WHERE evaluation_id=?',(record['evaluation_id'],)).fetchone()
            require(row is not None and digests(row)==record['snapshot_digests'], '输入/结果/规则快照字节漂移')
            request,result,standard = EvaluationRequest.model_validate_json(row[0]),EvaluationResult.model_validate_json(row[1]),StandardDefinition.model_validate_json(row[2])
            require(result.rule_snapshot_sha256 == EvaluationEngine._snapshot_sha256(standard), '结果与规则快照关联hash不符')
            check_result(result,record['mode'],record['dual'])
            require(standard.rule_revision==2 and build_cement_report_view(request,result,standard).complete, '保存快照投影无效')
    output = {'status':'passed','process_id':os.getpid(),'read_only':True,'recalculation':False,'records':len(ledger['records'])}
    (root/'smoke-readback.json').write_text(json.dumps(output,ensure_ascii=False,indent=2),encoding='utf-8')
    return 0


def run_smoke(parent: Path) -> int:
    from .stage_d_candidate import CANDIDATE_STANDARD_IDS, prepare_isolated_root
    from .bootstrap import create_context
    from .main import reconcile_standard_package
    from .infrastructure.logging import close_logging
    from .application.cement_reporting import ALTITUDE_DISCLOSURE
    from sqlalchemy import text
    from docx import Document
    root = prepare_isolated_root(parent/('smoke-'+uuid4().hex))
    context = create_context(root,public_key_path=Path(__file__).parent/'resources/update_public_key.pem',formal_standard_ids=CANDIDATE_STANDARD_IDS)
    report = {'status':'running','frozen':bool(getattr(sys,'frozen',False)),'process_id':os.getpid(),'data_directory':str(root),
        'build_identity':json.loads((parent/'candidate-runtime.json').read_text(encoding='utf-8')),
        'records':[],'visual_acceptance':'未验','installation_acceptance':'未验'}
    try:
        reconcile_standard_package(context)
        for mode,dual in [*[(mode,False) for mode in MODE_IDS],('integrated',True)]:
            result = context.application.evaluate(sample_request(mode,dual=dual))
            actual = check_result(result,mode,dual)
            docx = root/('水泥-'+mode+('-双修正' if dual else '')+'.docx')
            context.application.export_word_report(result.evaluation_id,docx)
            document = Document(docx)
            contents = '\n'.join([p.text for p in document.paragraphs]+[cell.text for table in document.tables for row in table.rows for cell in row.cells])
            require('GB 16780' in contents or 'GB16780' in contents, 'Word标准身份缺失')
            if dual:
                require(ALTITUDE_DISCLOSURE in contents,'Word必要海拔说明缺失')
            view = context.application.get_report_view(result.evaluation_id)
            require(len(view.indicators)==len(MODE_IDS[mode]) and all(item.label in contents for item in view.indicators),'Word适用指标缺失')
            with context.database.engine.connect() as connection:
                row = connection.execute(text('SELECT request_json,result_json,rule_snapshot_json FROM evaluations WHERE evaluation_id=:id'),{'id':result.evaluation_id}).one()
            report['records'].append({'mode':mode,'dual':dual,'evaluation_id':result.evaluation_id,'snapshot_digests':digests(row),
                'actual':actual,'word_path':str(docx),'word_sha256':hashlib.sha256(docx.read_bytes()).hexdigest()})
        partial = sample_request()
        partial.inputs.pop('grinding_electricity_kwh')
        before = context.evaluations.count()
        preview = context.application.preview_evaluation(partial)
        require(sum(item.grade.value=='INCOMPLETE' for item in preview.results)==1,'部分预览依赖隔离错误')
        try:
            context.application.evaluate(partial)
        except ValueError:
            pass
        else:
            raise RuntimeError('部分预览错误保存')
        require(context.evaluations.count()==before,'拒绝保存后记录数变化')
        try:
            context.application.export_word_report(preview.evaluation_id,root/'不得生成的部分预览.docx')
        except LookupError:
            pass
        else:
            raise RuntimeError('未保存预览错误导出')
        require(not (root/'不得生成的部分预览.docx').exists(),'部分预览写出Word')
        report['partial_save_and_word_rejected'] = True
    finally:
        context.database.dispose()
        close_logging()
    (root/'smoke-ledger.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    command = [sys.executable] if getattr(sys,'frozen',False) else [sys.executable,str(Path(__file__).resolve().parents[2]/'run_stage_d_cement.py')]
    with (root/'readback-stdout.log').open('wb') as stdout,(root/'readback-stderr.log').open('wb') as stderr:
        child = subprocess.run([*command,'--data-dir',str(root),'--acceptance-readback'],stdin=subprocess.DEVNULL,
            stdout=stdout,stderr=stderr,timeout=120,creationflags=0x08000000 if sys.platform=='win32' else 0)
    require(child.returncode==0,f'独立读回失败：{root}')
    report['readback'] = json.loads((root/'smoke-readback.json').read_text(encoding='utf-8'))
    require(report['readback']['process_id']!=report['process_id'],'未使用独立进程')
    report['status'] = 'passed'
    (root/'acceptance-smoke.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    return 0
