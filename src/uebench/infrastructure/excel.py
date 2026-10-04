from __future__ import annotations

import hashlib
import json
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from uuid import uuid4

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import BaseModel, Field

from uebench.application import gb29446 as gb
from uebench.domain.models import (
    EnergyLine,
    GRADE_LABELS,
    EvaluationRequest,
    Grade,
    InputMode,
    InputValue,
    ProductionLine,
    StandardDefinition,
    StandardSelectionMode,
)

from .database import DatabaseManager, ImportBatchRow
from .repositories import AuditRepository, SqlEvaluationRepository


TEMPLATE_VERSION = "1.1"
IMPORT_SHEETS = ["填写说明", "评价信息", "适用条件", "实际值", "能源明细", "产量与分摊"]
EXPORT_SHEETS = ["对标结论", "评价信息", "输入数据", "计算过程", "标准依据", "校验与日志"]
IMPORT_HEADERS = {
    "评价信息": ["字段", "值", "说明"],
    "适用条件": ["键", "名称", "值", "单位", "数据来源/备注"],
    "实际值": ["键", "名称", "值", "单位", "数据来源/备注"],
    "能源明细": ["行ID", "能源名称", "分类键", "方向", "实物量", "实物量单位", "折标系数", "系数单位", "分摊比例", "数据来源/备注"],
    "产量与分摊": ["行ID", "产品名称", "分类键", "产量", "单位", "折算系数", "是否合格", "数据来源/备注"],
}

# Import batch lifecycle.  ``evaluated`` means a formal record now exists;
# ``committed`` is the legacy generic flow and stays supported for the other
# standards, but the reference-standard path must not use it because a failed
# evaluation would leave the batch unrecoverable.
STATUS_VALIDATED = "validated"
STATUS_INVALID = "invalid"
STATUS_COMMITTED = "committed"
STATUS_EVALUATED = "evaluated"

GENERIC_PROFILE_ID = "ecq.generic.excel.v1"
GB29446_PROFILE_ID = "ecq.gb29446.excel.v1"

# GB29446 template layout ---------------------------------------------------
GB29446_TEMPLATE_TITLE = "GB 29446—2019 选煤电力消耗限额 评价数据导入模板"
GB29446_INSTRUCTIONS_SHEET = "填写说明"
GB29446_DATA_SHEET = "评价数据"
GB29446_META_SHEET = "_meta"
GB29446_META_HEADERS = ["字段", "值"]
GB29446_META_FIELDS = (
    "adapter_id",
    "template_version",
    "standard_id",
    "standard_version",
    "rule_revision",
)

#: Excel rows of the 评价数据 sheet (field label is column A, value is column B).
GB29446_DATA_ROWS: tuple[tuple[str, str], ...] = (
    (gb.FIELD_ORGANIZATION_NAME, "可选填写"),
    (gb.FIELD_EVALUATION_DATE, "日期，例如 2026-06-01；不得早于标准实施日期"),
    (gb.FIELD_PERIOD, "下拉选择；选“自定义”时填写下一行"),
    (gb.FIELD_CUSTOM_PERIOD, "仅当核算周期为“自定义”时填写"),
    (gb.FIELD_COAL_TYPE, "下拉选择煤种"),
    (gb.FIELD_WASHING_PROCESS, "下拉选择选煤工艺；软件按煤种与工艺校验合法性"),
    (gb.FIELD_ELECTRICITY, "文本格式，录入十进制字符串，例如 560 或 560.25；不要使用 Excel 公式或数值单元格"),
    (gb.FIELD_RAW_COAL, "文本格式，录入十进制字符串，必须大于 0；不要使用 Excel 公式或数值单元格"),
    (gb.FIELD_NOTES, "可选"),
)

#: Authoritative decimal inputs of the GB29446 template: must be text cells.
GB29446_DECIMAL_FIELDS = (gb.FIELD_ELECTRICITY, gb.FIELD_RAW_COAL)
#: Free-text fields that must be text (or empty).
GB29446_TEXT_FIELDS = (gb.FIELD_ORGANIZATION_NAME, gb.FIELD_CUSTOM_PERIOD, gb.FIELD_NOTES)


@dataclass(frozen=True)
class WorkbookProfile:
    """Immutable description of one workbook layout.

    The generic profile keeps the historical multi-standard template byte for
    byte.  A standard-specific profile (currently only GB29446) shares the same
    Excel infrastructure but owns its sheet structure and parsing rules, so the
    two never depend on a single module-level constant.
    """

    profile_id: str
    template_version: str
    sheet_names: tuple[str, ...]
    headers: dict[str, list[str]] = field(default_factory=dict)
    kind: str = "generic"
    requires_meta_sheet: bool = False


GENERIC_PROFILE = WorkbookProfile(
    profile_id=GENERIC_PROFILE_ID,
    template_version=TEMPLATE_VERSION,
    sheet_names=tuple(IMPORT_SHEETS),
    headers=IMPORT_HEADERS,
    kind="generic",
)

GB29446_PROFILE = WorkbookProfile(
    profile_id=GB29446_PROFILE_ID,
    template_version=gb.GB29446_EXCEL_TEMPLATE_VERSION,
    sheet_names=(GB29446_INSTRUCTIONS_SHEET, GB29446_DATA_SHEET, GB29446_META_SHEET),
    headers={
        GB29446_DATA_SHEET: ["字段", "值", "说明"],
        GB29446_META_SHEET: GB29446_META_HEADERS,
    },
    kind="gb29446",
    requires_meta_sheet=True,
)


class ImportIssue(BaseModel):
    severity: str
    sheet: str
    cell: str
    message: str


class ImportReport(BaseModel):
    import_id: str
    valid: bool
    profile_id: str = GENERIC_PROFILE_ID
    #: Standard version resolved at validation time.  Formal submission re-resolves
    #: the definition and refuses to proceed when it no longer matches, so a rule
    #: revision installed after validation cannot silently change the result.
    standard_version: str | None = None
    #: Rule revision resolved at validation time (see ``standard_version``).
    rule_revision: int | None = None
    issues: list[ImportIssue] = Field(default_factory=list)
    request: EvaluationRequest | None = None


class EvaluationDraft(BaseModel):
    import_id: str
    request: EvaluationRequest


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _excel_value(value):
    if isinstance(value, float):
        return str(value)
    return value


def _finite_decimal(value, *, sheet: str, cell: str, issues: list[ImportIssue]) -> Decimal | None:
    """Parse a decimal and reject non-finite results.

    ``NaN``/``sNaN``/``Infinity`` are accepted by ``Decimal(str(...))`` but are
    not business values.  The domain guards authoritative inputs too; this keeps
    the failure inside the adapter where a Chinese, cell-anchored message can be
    produced.
    """
    try:
        parsed = Decimal(str(value))
    except (InvalidOperation, ValueError):
        issues.append(ImportIssue(severity="error", sheet=sheet, cell=cell, message="不是有效十进制数"))
        return None
    if not parsed.is_finite():
        issues.append(
            ImportIssue(
                severity="error",
                sheet=sheet,
                cell=cell,
                message="正式数值必须为有限十进制数；不接受 NaN / Infinity。",
            )
        )
        return None
    return parsed


def _decimal_from_cell(value, *, sheet: str, cell: str, issues: list[ImportIssue]) -> Decimal | None:
    """Generic-profile decimal ingress.

    Rejects XLSX numeric float cells and non-finite values; a whole-number
    numeric cell is still accepted here because the generic template uses
    ``openpyxl``-style numeric defaults (for example an allocation ratio of
    ``1``) and the other 46 standards rely on that behaviour.  The GB29446
    profile uses :func:`_lexical_decimal_from_cell`, which requires text.
    """
    if value is None or value == "":
        issues.append(ImportIssue(severity="error", sheet=sheet, cell=cell, message="数值不能为空"))
        return None
    if isinstance(value, float):
        issues.append(
            ImportIssue(
                severity="error",
                sheet=sheet,
                cell=cell,
                message=(
                    "正式数值不能使用 XLSX 数值单元格导入：openpyxl 已将其物化为二进制浮点。"
                    "请将单元格设为文本并录入十进制字符串。"
                ),
            )
        )
        return None
    return _finite_decimal(value, sheet=sheet, cell=cell, issues=issues)


def _formula_text(cell) -> str | None:
    """Return the formula text when a cell holds a formula, else ``None``.

    ``data_only=True`` reads the cached result and cannot report that a formula
    exists, so authoritative parsing always consults a ``data_only=False`` view.
    """
    return cell.value if getattr(cell, "data_type", None) == "f" else None


def _lexical_decimal_from_cell(
    cell,
    *,
    sheet: str,
    cell_ref: str,
    issues: list[ImportIssue],
) -> Decimal | None:
    """GB29446 authoritative decimal ingress: text cells only.

    The contract is a single, stable, auditable decimal lexical text ingress:

    * accepted: text ``"560"``, ``"560.25"``, ``"1e3"`` (finite lexical decimal);
    * rejected: XLSX numeric cells (``int`` or ``float``), Excel formulas (with
      or without a cached value), non-finite text, blanks, and booleans.

    Rejecting a numeric *integer* is not a claim that integers are imprecise.
    It removes the last path by which a workbook could hand the domain an
    adapter-manufactured value, so type, provenance and audit stay unambiguous.
    """
    formula = _formula_text(cell)
    if formula is not None:
        issues.append(
            ImportIssue(
                severity="error",
                sheet=sheet,
                cell=cell_ref,
                message="正式数值不能使用 Excel 公式，请将计算结果粘贴为文本值后重新导入。",
            )
        )
        return None
    value = cell.value
    if value is None or value == "":
        issues.append(ImportIssue(severity="error", sheet=sheet, cell=cell_ref, message="数值不能为空"))
        return None
    if isinstance(value, bool):
        issues.append(
            ImportIssue(severity="error", sheet=sheet, cell=cell_ref, message="正式数值必须为文本格式的十进制数")
        )
        return None
    if not isinstance(value, str):
        issues.append(
            ImportIssue(
                severity="error",
                sheet=sheet,
                cell=cell_ref,
                message=(
                    "正式数值必须使用文本单元格录入十进制字符串。"
                    "当前单元格是 XLSX 数值单元格或公式；请将单元格设为文本后重新录入。"
                ),
            )
        )
        return None
    return _finite_decimal(value.strip(), sheet=sheet, cell=cell_ref, issues=issues)


def _text_from_cell(cell, *, sheet: str, cell_ref: str, issues: list[ImportIssue], allow_blank: bool = True) -> str | None:
    """GB29446 text field ingress: text cells only, formulas rejected."""
    if _formula_text(cell) is not None:
        issues.append(
            ImportIssue(
                severity="error",
                sheet=sheet,
                cell=cell_ref,
                message="本字段不能使用 Excel 公式，请粘贴为文本值后重新导入。",
            )
        )
        return None
    value = cell.value
    if value is None or value == "":
        if allow_blank:
            return None
        issues.append(ImportIssue(severity="error", sheet=sheet, cell=cell_ref, message="该字段不能为空"))
        return None
    if not isinstance(value, str):
        issues.append(
            ImportIssue(
                severity="error",
                sheet=sheet,
                cell=cell_ref,
                message="该字段必须使用文本单元格；请将单元格设为文本后重新录入。",
            )
        )
        return None
    return value.strip()


def _json_cell(value) -> str:
    """Render a metadata value deterministically for comparison."""
    return json.dumps(value, ensure_ascii=False, sort_keys=True)


def _gb29446_process_options(standard: StandardDefinition) -> list[str]:
    """Declared washing-process choices, read from the definition.

    The adapter must not keep its own ``process -> product`` truth table, so the
    option list comes from the domain lookup rows exactly as the GUI reads it.
    """
    options: list[str] = []
    for product in standard.products:
        for indicator in product.indicators:
            for display in indicator.display_calculations:
                if display.key != "process_factor" or display.formula.op != "lookup":
                    continue
                for row in display.formula.rows:
                    value = row.condition.value
                    if isinstance(value, str) and value not in options:
                        options.append(value)
            for definition in product.input_definitions:
                if definition.key == gb.INPUT_WASHING_PROCESS:
                    for choice in definition.choices:
                        if choice not in options:
                            options.append(choice)
    return options


def _gb29446_coal_options(standard: StandardDefinition) -> list[str]:
    """Coal-type choices declared by the definition's selection schema."""
    options: list[str] = []
    selection_keys = {level.key for level in standard.selection_schema}
    for product in standard.products:
        for key, value in product.selection_values.items():
            if selection_keys and key not in selection_keys:
                continue
            if value not in options:
                options.append(value)
    return options


def _gb29446_product_for_coal(standard: StandardDefinition, coal_type: str | None):
    """Resolve coal type onto a product using the definition, never a hardcoded map."""
    if not coal_type:
        return None
    return next(
        (product for product in standard.products if coal_type in product.selection_values.values()),
        None,
    )


def _gb29446_process_is_valid(product, process: str) -> bool:
    """A process is valid when the definition declares its lookup coefficient."""
    for indicator in product.indicators:
        for display in indicator.display_calculations:
            if display.key != "process_factor" or display.formula.op != "lookup":
                continue
            for row in display.formula.rows:
                if row.condition.value == process:
                    return True
    return False


def _has_error(issues: list[ImportIssue]) -> bool:
    return any(issue.severity == "error" for issue in issues)


def _first_non_empty_row(sheet, column: int) -> int | None:
    """Lowest row index in ``column`` that holds a non-empty value.

    ``None`` means the column carries no content at all (pure formatting
    widening).  Row 1 is included so a non-blank header is reported too.
    """
    for row in range(1, sheet.max_row + 1):
        value = sheet.cell(row=row, column=column).value
        if value is not None and value != "":
            return row
    return None


class WorkbookTemplateService:
    """Create import templates.

    ``standard_definition`` selects the profile.  The application layer resolves
    the definition, so this service stays free of repository dependencies.
    """

    def create_template(self, path: Path, standard_definition: StandardDefinition | None = None) -> Path:
        if standard_definition is None:
            return self._create_generic(path)
        if standard_definition.id == gb.GB29446_STANDARD_ID:
            return self._create_gb29446(path, standard_definition)
        raise ValueError(f"标准 {standard_definition.number} 尚未提供专用 Excel 模板；请使用通用模板。")

    # -- generic ----------------------------------------------------------
    def _create_generic(self, path: Path) -> Path:
        workbook = Workbook()
        workbook.remove(workbook.active)
        for name in GENERIC_PROFILE.sheet_names:
            workbook.create_sheet(name)
        self._instructions(workbook["填写说明"])
        self._evaluation_info(workbook["评价信息"])
        self._generic_inputs(workbook["适用条件"], "适用条件")
        self._generic_inputs(workbook["实际值"], "实际值")
        self._energy_lines(workbook["能源明细"])
        self._production_lines(workbook["产量与分摊"])
        for sheet in workbook.worksheets:
            sheet.sheet_view.showGridLines = False
            sheet.freeze_panes = "A2"
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(path)
        return path

    # -- GB29446 reference standard ---------------------------------------
    def _create_gb29446(self, path: Path, standard: StandardDefinition) -> Path:
        workbook = Workbook()
        workbook.remove(workbook.active)
        workbook.create_sheet(GB29446_INSTRUCTIONS_SHEET)
        workbook.create_sheet(GB29446_DATA_SHEET)
        meta_sheet = workbook.create_sheet(GB29446_META_SHEET)

        self._gb29446_instructions(workbook[GB29446_INSTRUCTIONS_SHEET])
        self._gb29446_data(workbook[GB29446_DATA_SHEET], standard)
        self._gb29446_meta(meta_sheet, standard)

        # Hidden is a usability aid only; the importer validates metadata
        # strictly, so hiding is never treated as a security measure.
        meta_sheet.sheet_state = "hidden"

        for sheet in workbook.worksheets:
            sheet.sheet_view.showGridLines = False
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(path)
        return path

    def _gb29446_instructions(self, sheet) -> None:
        sheet["A1"] = GB29446_TEMPLATE_TITLE
        sheet["A1"].font = Font(size=16, bold=True, color="1F4E78")
        sheet.merge_cells("A1:D1")
        notes = [
            "本模板只用于 GB 29446—2019 一次评价。模板版本、工作表名称和 _meta 内容不可修改。",
            "请在“评价数据”表按字段填写；带下拉箭头的字段请使用下拉选择。",
            f"{gb.FIELD_ELECTRICITY} 与 {gb.FIELD_RAW_COAL} 必须使用文本格式录入十进制字符串（例如 560、560.25、1e3）。",
            "不要使用 Excel 公式或 XLSX 数值单元格填写正式数值；请把计算结果粘贴为文本值。",
            "本表不提供折算系数 k、单位产品电耗 e_d 或等级判断；这些结果全部由软件按标准正式计算。",
            "填写完成后在软件中选择该文件并点击校验，确认识别摘要后再提交评价。",
        ]
        for row, note in enumerate(notes, start=3):
            sheet.cell(row=row, column=1, value=f"{row - 2}. {note}")
        sheet.column_dimensions["A"].width = 96

    def _gb29446_data(self, sheet, standard: StandardDefinition) -> None:
        self._header(sheet, ["字段", "值", "说明"])
        for label, hint in GB29446_DATA_ROWS:
            sheet.append([label, None, hint])

        # Dropdowns: periods, coal types and all declared processes.  The same
        # full process list is offered for both coal types; the importer then
        # validates the coal/process combination against the definition.
        from openpyxl.utils.cell import coordinate_from_string  # local import keeps top tidy
        row_of = {label: index for index, (label, _hint) in enumerate(GB29446_DATA_ROWS, start=2)}

        period_validation = DataValidation(
            type="list",
            formula1='"' + ",".join(gb.GB29446_PERIOD_OPTIONS) + '"',
            allow_blank=False,
        )
        sheet.add_data_validation(period_validation)
        period_validation.add(sheet.cell(row=row_of[gb.FIELD_PERIOD], column=2))

        coal_options = [
            value
            for level in standard.selection_schema
            for product in standard.products
            for key, value in product.selection_values.items()
            if key == level.key
        ]
        coal_options = list(dict.fromkeys(coal_options))
        if coal_options:
            coal_validation = DataValidation(
                type="list", formula1='"' + ",".join(coal_options) + '"', allow_blank=False
            )
            sheet.add_data_validation(coal_validation)
            coal_validation.add(sheet.cell(row=row_of[gb.FIELD_COAL_TYPE], column=2))

        process_options = _gb29446_process_options(standard)
        if process_options:
            process_validation = DataValidation(
                type="list",
                formula1='"' + ",".join(process_options) + '"',
                allow_blank=False,
            )
            sheet.add_data_validation(process_validation)
            process_validation.add(sheet.cell(row=row_of[gb.FIELD_WASHING_PROCESS], column=2))

        sheet.cell(row=row_of[gb.FIELD_EVALUATION_DATE], column=2).number_format = "yyyy-mm-dd"
        # Authoritative decimal fields plus free-text fields are text cells, so
        # Excel does not silently materialise a numeric cell.
        for label in GB29446_DECIMAL_FIELDS + GB29446_TEXT_FIELDS:
            sheet.cell(row=row_of[label], column=2).number_format = "@"
        sheet.column_dimensions["A"].width = 30
        sheet.column_dimensions["B"].width = 40
        sheet.column_dimensions["C"].width = 62

    def _gb29446_meta(self, sheet, standard: StandardDefinition) -> None:
        self._header(sheet, GB29446_META_HEADERS)
        values = {
            "adapter_id": gb.GB29446_EXCEL_ADAPTER_ID,
            "template_version": gb.GB29446_EXCEL_TEMPLATE_VERSION,
            "standard_id": standard.id,
            "standard_version": standard.version,
            "rule_revision": standard.rule_revision,
        }
        for key in GB29446_META_FIELDS:
            sheet.append([key, values[key]])
        sheet.column_dimensions["A"].width = 24
        sheet.column_dimensions["B"].width = 34

    # -- shared styling ---------------------------------------------------
    @staticmethod
    def _header(sheet, headers: list[str]) -> None:
        sheet.append(headers)
        fill = PatternFill("solid", fgColor="1F4E78")
        for cell in sheet[1]:
            cell.fill = fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.row_dimensions[1].height = 26

    def _instructions(self, sheet) -> None:
        sheet["A1"] = "单位产品能耗对标数据导入模板"
        sheet["A1"].font = Font(size=16, bold=True, color="1F4E78")
        sheet.merge_cells("A1:D1")
        notes = [
            "模板版本和工作表名称不可修改。",
            "直接录入模式填写“实际值”；明细计算模式填写“能源明细”和“产量与分摊”。",
            "适用条件和实际值中的“键”应由软件生成的标准专用模板提供。",
            "正式数值请使用文本格式录入十进制字符串（模板数值输入列已设为文本）；不要使用 XLSX 浮点数值单元格，也不要录入带单位的文本。",
            "direction 只能填写 input 或 output；分摊比例范围为 0~1。",
        ]
        for row, note in enumerate(notes, start=3):
            sheet.cell(row=row, column=1, value=f"{row - 2}. {note}")
        sheet.column_dimensions["A"].width = 90

    def _evaluation_info(self, sheet) -> None:
        self._header(sheet, ["字段", "值", "说明"])
        rows = [
            ("template_version", TEMPLATE_VERSION, "固定值"),
            ("evaluation_date", date.today(), "日期，且不得早于标准实施日期"),
            ("standard_id", "", "标准唯一ID"),
            ("product_id", "", "产品或工序唯一ID"),
            ("input_mode", InputMode.DIRECT.value, "DIRECT 或 DETAIL"),
            ("organization_name", "", "单位名称，可空"),
            ("project_name", "", "项目名称，可空"),
            ("notes", "", "备注，可空"),
        ]
        for item in rows:
            sheet.append(item)
        sheet["B3"].number_format = "yyyy-mm-dd"
        validation = DataValidation(type="list", formula1='"DIRECT,DETAIL"', allow_blank=False)
        sheet.add_data_validation(validation)
        validation.add(sheet["B6"])
        sheet.column_dimensions["A"].width = 26
        sheet.column_dimensions["B"].width = 34
        sheet.column_dimensions["C"].width = 48

    def _generic_inputs(self, sheet, _title: str) -> None:
        self._header(sheet, ["键", "名称", "值", "单位", "数据来源/备注"])
        for _ in range(20):
            sheet.append(["", "", "", "", ""])
        for row in range(2, 22):
            sheet.cell(row=row, column=3).number_format = "@"
        widths = [28, 34, 20, 18, 48]
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[chr(64 + index)].width = width

    def _energy_lines(self, sheet) -> None:
        self._header(
            sheet,
            ["行ID", "能源名称", "分类键", "方向", "实物量", "实物量单位", "折标系数", "系数单位", "分摊比例", "数据来源/备注"],
        )
        for _ in range(50):
            sheet.append(["", "", "", "input", "", "", "", "", 1, ""])
        for row in range(2, 52):
            for column in (5, 7, 9):
                sheet.cell(row=row, column=column).number_format = "@"
        validation = DataValidation(type="list", formula1='"input,output"', allow_blank=False)
        sheet.add_data_validation(validation)
        validation.add("D2:D51")
        widths = [16, 24, 18, 12, 16, 16, 16, 20, 14, 40]
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[chr(64 + index)].width = width

    def _production_lines(self, sheet) -> None:
        self._header(sheet, ["行ID", "产品名称", "分类键", "产量", "单位", "折算系数", "是否合格", "数据来源/备注"])
        for _ in range(30):
            sheet.append(["", "", "", "", "", 1, "是", ""])
        for row in range(2, 32):
            for column in (4, 6):
                sheet.cell(row=row, column=column).number_format = "@"
        validation = DataValidation(type="list", formula1='"是,否"', allow_blank=False)
        sheet.add_data_validation(validation)
        validation.add("G2:G31")
        widths = [16, 30, 18, 16, 14, 16, 14, 44]
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[chr(64 + index)].width = width


class WorkbookImportService:
    """Validate workbooks and hand a canonical request to the application layer.

    This service never runs the evaluation engine and never creates a formal
    record.  Formal calculation stays with ``EvaluationService``; the adapter
    only reads a workbook and maps it onto ``EvaluationRequest``.
    """

    def __init__(self, database: DatabaseManager, audit: AuditRepository | None = None, standards=None) -> None:
        self.database = database
        self.audit = audit or AuditRepository(database)
        self.standards = standards

    # -- validation -------------------------------------------------------
    def validate(self, path: Path) -> ImportReport:
        import_id = str(uuid4())
        issues: list[ImportIssue] = []
        request: EvaluationRequest | None = None
        profile = GENERIC_PROFILE
        path = path.resolve()
        try:
            # load_workbook is used without read_only here: the authoritative
            # ingress must read cell.data_type to detect formulas, and read-only
            # worksheets cannot report that reliably.  Structure correctness is
            # worth more than the memory saving.
            workbook = load_workbook(path, data_only=False)
            profile, sheet_ok = self._detect_profile(workbook, issues)
            if sheet_ok:
                if profile.kind == "gb29446":
                    self._validate_gb29446_structure(workbook, issues)
                    if not _has_error(issues):
                        request = self._parse_gb29446(workbook, issues)
                else:
                    self._validate_fixed_structure(workbook, issues)
                    if not _has_error(issues):
                        request = self._parse_generic(workbook, issues)
        except Exception as exc:
            issues.append(ImportIssue(severity="error", sheet="工作簿", cell="-", message=str(exc)))
        valid = request is not None and not _has_error(issues)
        resolved = self._resolved_identity(profile, workbook) if valid else (None, None)
        report = ImportReport(
            import_id=import_id,
            valid=valid,
            profile_id=profile.profile_id,
            standard_version=resolved[0],
            rule_revision=resolved[1],
            issues=issues,
            request=request if valid else None,
        )
        with self.database.session() as session:
            session.add(
                ImportBatchRow(
                    import_id=import_id,
                    source_file=str(path),
                    source_sha256=_sha256(path) if path.exists() else "0" * 64,
                    status=STATUS_VALIDATED if valid else STATUS_INVALID,
                    payload_json=request.model_dump_json() if request is not None else "{}",
                    validation_json=report.model_dump_json(),
                )
            )
            self.audit.append(
                "WORKBOOK_VALIDATE",
                "import_batch",
                import_id,
                {"valid": valid, "issue_count": len(issues), "profile_id": profile.profile_id},
                session=session,
            )
        return report

    def _resolved_identity(self, profile: WorkbookProfile, workbook) -> tuple[str | None, int | None]:
        """Identity of the definition this workbook's evaluation would use.

        Recorded at validation time so formal submission can detect that the
        currently applicable rule revision changed in the meantime.
        """
        if profile.kind != "gb29446" or self.standards is None:
            return None, None
        issues: list[ImportIssue] = []
        standard, _meta_values, _date = self._resolve_gb29446_standard(workbook, issues)
        if standard is None:
            return None, None
        return standard.version, standard.rule_revision

    @staticmethod
    def _detect_profile(workbook, issues: list[ImportIssue]) -> tuple[WorkbookProfile, bool]:
        """Resolve the workbook profile from its sheet names.

        Detection is structural, not metadata-based, so a corrupted ``_meta``
        sheet still reports the specific metadata error instead of an opaque
        "unknown template" one.
        """
        actual = list(workbook.sheetnames)
        if actual == list(GB29446_PROFILE.sheet_names):
            return GB29446_PROFILE, True
        if actual == list(GENERIC_PROFILE.sheet_names):
            return GENERIC_PROFILE, True
        issues.append(
            ImportIssue(
                severity="error",
                sheet="工作簿",
                cell="-",
                message=(
                    f"工作表必须严格为 {', '.join(GENERIC_PROFILE.sheet_names)}"
                    f"（通用模板）或 {', '.join(GB29446_PROFILE.sheet_names)}（GB 29446—2019 专用模板）"
                ),
            )
        )
        return GENERIC_PROFILE, False

    @staticmethod
    def _validate_fixed_structure(workbook, issues: list[ImportIssue]) -> None:
        """Reject renamed, reordered, missing, or extra generic template columns.

        The import contract is intentionally strict: a workbook with a custom
        schema must be converted to the official template before it can enter
        the audit trail.  Empty data rows remain allowed; only the header and
        instruction sheet structure is checked here.
        """
        instructions = workbook["填写说明"]
        if instructions["A1"].value != "单位产品能耗对标数据导入模板":
            issues.append(ImportIssue(severity="error", sheet="填写说明", cell="A1", message="模板标题不正确"))
        WorkbookImportService._check_headers(workbook, GENERIC_PROFILE, issues)

    @staticmethod
    def _check_headers(workbook, profile: WorkbookProfile, issues: list[ImportIssue]) -> None:
        for sheet_name, expected in profile.headers.items():
            sheet = workbook[sheet_name]
            actual = [sheet.cell(row=1, column=index).value for index in range(1, len(expected) + 1)]
            if actual != expected:
                issues.append(
                    ImportIssue(
                        severity="error",
                        sheet=sheet_name,
                        cell="A1",
                        message=f"表头必须严格为：{'、'.join(expected)}",
                    )
                )
            # Anything beyond the declared width must be empty everywhere, not
            # just in the header row.  A column whose header is blank but that
            # carries data below is still a real extra column, so scanning only
            # row 1 would silently accept it.  Merely widening columns by
            # formatting (no content at all) stays compatible.
            for column in range(len(expected) + 1, sheet.max_column + 1):
                offending_row = _first_non_empty_row(sheet, column)
                if offending_row is not None:
                    issues.append(
                        ImportIssue(
                            severity="error",
                            sheet=sheet_name,
                            cell=f"{get_column_letter(column)}{offending_row}",
                            message="模板不允许增加自定义列",
                        )
                    )

    @staticmethod
    def _validate_gb29446_structure(workbook, issues: list[ImportIssue]) -> None:
        instructions = workbook[GB29446_INSTRUCTIONS_SHEET]
        if instructions["A1"].value != GB29446_TEMPLATE_TITLE:
            issues.append(
                ImportIssue(severity="error", sheet=GB29446_INSTRUCTIONS_SHEET, cell="A1", message="模板标题不正确")
            )
        WorkbookImportService._check_headers(workbook, GB29446_PROFILE, issues)
        data = workbook[GB29446_DATA_SHEET]
        for offset, (label, _hint) in enumerate(GB29446_DATA_ROWS, start=2):
            if data.cell(row=offset, column=1).value != label:
                issues.append(
                    ImportIssue(
                        severity="error",
                        sheet=GB29446_DATA_SHEET,
                        cell=f"A{offset}",
                        message=f"字段行必须严格为：{label}",
                    )
                )

    # -- lifecycle --------------------------------------------------------
    def prepare(self, import_id: str) -> EvaluationDraft:
        """Return the canonical request, re-verifying it before formal evaluation.

        Two independent checks run here, and both must pass:

        1. **source integrity** - the file still exists with the same SHA-256;
        2. **applicable rule identity** - the workbook is re-parsed through the
           formal resolver, so a standard/rule revision installed after
           validation cannot be applied to a template that was generated for an
           older revision.  An unchanged file hash does *not* prove that the
           template still matches the definition that will actually run.

        The batch status is deliberately left at ``validated``: only a successful
        formal evaluation may advance it, so a failed evaluation stays retryable.
        """
        with self.database.session() as session:
            row = session.get(ImportBatchRow, import_id)
            if row is None:
                raise LookupError("导入批次不存在")
            if row.status not in (STATUS_VALIDATED, STATUS_COMMITTED):
                raise ValueError("只有校验通过的导入批次可以评价")
            self._verify_source(row)
            self._verify_applicable_revision(row)
            request = EvaluationRequest.model_validate_json(row.payload_json)
            return EvaluationDraft(import_id=import_id, request=request)

    def _verify_applicable_revision(self, row: ImportBatchRow) -> None:
        """Re-run the resolution + metadata gate immediately before evaluation.

        Uses the same ``_resolve_gb29446_standard`` helper as validation, so the
        adapter still owns no second standard-selection rule.  The recorded
        ``rule_revision`` is compared as well, which catches a definition swap
        that left every other field equal.
        """
        profile = self._profile_of(row)
        if profile is None or profile.kind != "gb29446":
            return
        recorded = self._validation_record(row)
        issues: list[ImportIssue] = []
        try:
            workbook = load_workbook(Path(row.source_file), data_only=False)
        except Exception as exc:  # pragma: no cover - unreadable file already rejected
            raise ValueError("校验时使用的 Excel 文件已无法读取，请重新选择文件并重新校验。") from exc
        standard, _meta_values, _date = self._resolve_gb29446_standard(workbook, issues)
        if standard is None or _has_error(issues):
            raise ValueError(
                "该模板对应的标准/规则修订已不是当前本次评价适用版本，请重新生成模板后填写。"
            )
        if recorded.get("rule_revision") is not None and int(recorded["rule_revision"]) != standard.rule_revision:
            raise ValueError(
                "该模板对应的标准/规则修订已不是当前本次评价适用版本，请重新生成模板后填写。"
            )

    @staticmethod
    def _validation_record(row: ImportBatchRow) -> dict:
        try:
            return json.loads(row.validation_json)
        except (TypeError, ValueError):
            return {}

    def _profile_of(self, row: ImportBatchRow) -> WorkbookProfile | None:
        profile_id = self._validation_record(row).get("profile_id")
        if profile_id == GB29446_PROFILE.profile_id:
            return GB29446_PROFILE
        if profile_id == GENERIC_PROFILE.profile_id:
            return GENERIC_PROFILE
        return None

    @staticmethod
    def _verify_source(row: ImportBatchRow) -> None:
        """Re-check that the validated file still exists with the same content.

        ``source_file`` is only a locator; ``source_sha256`` is the content
        identity.  A moved file cannot be rediscovered from the old locator, so
        that case is rejected and the user is asked to re-select the file.  A
        path change is never itself treated as a content change.
        """
        path = Path(row.source_file)
        if not path.exists() or not path.is_file():
            raise ValueError("校验时使用的 Excel 文件已不存在或无法读取，请重新选择文件并重新校验。")
        try:
            current = _sha256(path)
        except OSError as exc:
            raise ValueError("校验时使用的 Excel 文件已无法读取，请重新选择文件并重新校验。") from exc
        if current != row.source_sha256:
            raise ValueError("文件在校验后已发生变化，请重新校验。")

    def mark_evaluated(self, import_id: str, evaluation_id: str) -> None:
        with self.database.session() as session:
            row = session.get(ImportBatchRow, import_id)
            if row is None:
                raise LookupError("导入批次不存在")
            row.status = STATUS_EVALUATED
            details = {
                "import_id": import_id,
                "evaluation_id": evaluation_id,
                "source_sha256": row.source_sha256,
                "profile_id": self._profile_id_for(row),
            }
            self.audit.append("WORKBOOK_EVALUATION", "import_batch", import_id, details, session=session)

    @staticmethod
    def _profile_id_for(row: ImportBatchRow) -> str:
        return WorkbookImportService._validation_record(row).get("profile_id") or GENERIC_PROFILE_ID

    def commit(self, import_id: str) -> EvaluationDraft:
        """Legacy generic submission path, retained for the other standards.

        The reference-standard path uses :meth:`prepare` + :meth:`mark_evaluated`
        instead, because marking ``committed`` before the evaluation succeeds
        would strand the batch when the evaluation fails.
        """
        with self.database.session() as session:
            row = session.get(ImportBatchRow, import_id)
            if row is None:
                raise LookupError("导入批次不存在")
            if row.status != STATUS_VALIDATED:
                raise ValueError("只有验证通过的导入批次可以提交")
            row.status = STATUS_COMMITTED
            request = EvaluationRequest.model_validate_json(row.payload_json)
            self.audit.append("WORKBOOK_COMMIT", "import_batch", import_id, session=session)
            return EvaluationDraft(import_id=import_id, request=request)

    # -- parsing ----------------------------------------------------------
    def _resolve_gb29446_standard(
        self, workbook, issues: list[ImportIssue]
    ) -> tuple[StandardDefinition | None, dict[str, object], date | None]:
        """Resolve the definition this evaluation would actually use.

        Selection goes through the正式 resolver only; the adapter never
        re-implements standard selection (for example ``max(rule_revision)``).
        Returns the resolved definition together with the template metadata that
        the workbook claims, so callers can compare the two.
        """
        if self.standards is None:
            issues.append(
                ImportIssue(severity="error", sheet="工作簿", cell="-", message="标准仓库未配置，无法校验 GB29446 模板")
            )
            return None, {}, None
        data = workbook[GB29446_DATA_SHEET]
        row_of = {label: index for index, (label, _hint) in enumerate(GB29446_DATA_ROWS, start=2)}
        meta_values = self._read_meta(workbook[GB29446_META_SHEET], issues)
        if _has_error(issues):
            return None, meta_values, None

        evaluation_date = self._read_gb29446_date(data.cell(row=row_of[gb.FIELD_EVALUATION_DATE], column=2), issues)
        if evaluation_date is None:
            return None, meta_values, None

        standard = self.standards.get_for_evaluation(
            str(meta_values.get("standard_id") or gb.GB29446_STANDARD_ID),
            evaluation_date,
            StandardSelectionMode.CURRENT,
        )
        if standard is None:
            issues.append(
                ImportIssue(
                    severity="error",
                    sheet=GB29446_DATA_SHEET,
                    cell=f"B{row_of[gb.FIELD_EVALUATION_DATE]}",
                    message="该评价日期没有适用的已发布 GB 29446—2019 规则。",
                )
            )
            return None, meta_values, evaluation_date

        self._check_meta_against_standard(meta_values, standard, issues)
        if _has_error(issues):
            return None, meta_values, evaluation_date
        return standard, meta_values, evaluation_date

    def _parse_gb29446(self, workbook, issues: list[ImportIssue]) -> EvaluationRequest | None:
        data = workbook[GB29446_DATA_SHEET]
        row_of = {label: index for index, (label, _hint) in enumerate(GB29446_DATA_ROWS, start=2)}

        standard, _meta_values, evaluation_date = self._resolve_gb29446_standard(workbook, issues)
        if standard is None or evaluation_date is None:
            return None

        period = self._read_gb29446_choice(
            data.cell(row=row_of[gb.FIELD_PERIOD], column=2),
            gb.GB29446_PERIOD_OPTIONS,
            sheet=GB29446_DATA_SHEET,
            cell_ref=f"B{row_of[gb.FIELD_PERIOD]}",
            issues=issues,
            field=gb.FIELD_PERIOD,
        )
        custom_period = _text_from_cell(
            data.cell(row=row_of[gb.FIELD_CUSTOM_PERIOD], column=2),
            sheet=GB29446_DATA_SHEET,
            cell_ref=f"B{row_of[gb.FIELD_CUSTOM_PERIOD]}",
            issues=issues,
        ) or ""
        if period == gb.PERIOD_CUSTOM and not custom_period:
            issues.append(
                ImportIssue(
                    severity="error",
                    sheet=GB29446_DATA_SHEET,
                    cell=f"B{row_of[gb.FIELD_CUSTOM_PERIOD]}",
                    message="核算周期为“自定义”时必须填写自定义周期。",
                )
            )
        if period != gb.PERIOD_CUSTOM:
            custom_period = ""

        organization = _text_from_cell(
            data.cell(row=row_of[gb.FIELD_ORGANIZATION_NAME], column=2),
            sheet=GB29446_DATA_SHEET,
            cell_ref=f"B{row_of[gb.FIELD_ORGANIZATION_NAME]}",
            issues=issues,
        )
        note = _text_from_cell(
            data.cell(row=row_of[gb.FIELD_NOTES], column=2),
            sheet=GB29446_DATA_SHEET,
            cell_ref=f"B{row_of[gb.FIELD_NOTES]}",
            issues=issues,
        ) or ""

        coal_type = self._read_gb29446_choice(
            data.cell(row=row_of[gb.FIELD_COAL_TYPE], column=2),
            _gb29446_coal_options(standard),
            sheet=GB29446_DATA_SHEET,
            cell_ref=f"B{row_of[gb.FIELD_COAL_TYPE]}",
            issues=issues,
            field=gb.FIELD_COAL_TYPE,
        )
        process = self._read_gb29446_choice(
            data.cell(row=row_of[gb.FIELD_WASHING_PROCESS], column=2),
            _gb29446_process_options(standard),
            sheet=GB29446_DATA_SHEET,
            cell_ref=f"B{row_of[gb.FIELD_WASHING_PROCESS]}",
            issues=issues,
            field=gb.FIELD_WASHING_PROCESS,
        )

        product = _gb29446_product_for_coal(standard, coal_type)
        if coal_type and product is None:
            issues.append(
                ImportIssue(
                    severity="error",
                    sheet=GB29446_DATA_SHEET,
                    cell=f"B{row_of[gb.FIELD_COAL_TYPE]}",
                    message=f"标准中不存在煤种：{coal_type}",
                )
            )
        if product is not None and process and not _gb29446_process_is_valid(product, process):
            issues.append(
                ImportIssue(
                    severity="error",
                    sheet=GB29446_DATA_SHEET,
                    cell=f"B{row_of[gb.FIELD_WASHING_PROCESS]}",
                    message=f"煤种“{coal_type}”不支持选煤工艺“{process}”，请核对标准附录A。",
                )
            )

        electricity = _lexical_decimal_from_cell(
            data.cell(row=row_of[gb.FIELD_ELECTRICITY], column=2),
            sheet=GB29446_DATA_SHEET,
            cell_ref=f"B{row_of[gb.FIELD_ELECTRICITY]}",
            issues=issues,
        )
        raw_coal = _lexical_decimal_from_cell(
            data.cell(row=row_of[gb.FIELD_RAW_COAL], column=2),
            sheet=GB29446_DATA_SHEET,
            cell_ref=f"B{row_of[gb.FIELD_RAW_COAL]}",
            issues=issues,
        )
        if _has_error(issues) or product is None or evaluation_date is None:
            return None

        inputs: dict[str, InputValue] = {}
        if process:
            inputs[gb.INPUT_WASHING_PROCESS] = InputValue(value=process)
        if electricity is not None:
            inputs[gb.INPUT_ELECTRICITY] = InputValue(value=str(electricity), unit=gb.UNIT_ELECTRICITY)
        if raw_coal is not None:
            inputs[gb.INPUT_RAW_COAL] = InputValue(value=str(raw_coal), unit=gb.UNIT_RAW_COAL)

        try:
            return EvaluationRequest(
                evaluation_date=evaluation_date,
                standard_id=standard.id,
                product_id=product.id,
                selection_mode=StandardSelectionMode.CURRENT,
                input_mode=InputMode.DETAIL,
                inputs=inputs,
                organization_name=organization or None,
                notes=gb.encode_period_notes(period or gb.PERIOD_FULL_YEAR, custom_period, note),
            )
        except ValueError as exc:
            issues.append(ImportIssue(severity="error", sheet=GB29446_DATA_SHEET, cell="-", message=str(exc)))
            return None

    @staticmethod
    def _read_meta(meta, issues: list[ImportIssue]) -> dict[str, object]:
        values: dict[str, object] = {}
        for row_number, row in enumerate(meta.iter_rows(min_row=2), start=2):
            key = row[0].value
            if not key:
                continue
            if key in values:
                issues.append(
                    ImportIssue(severity="error", sheet=GB29446_META_SHEET, cell=f"A{row_number}", message=f"_meta 字段重复：{key}")
                )
                continue
            values[str(key)] = row[1].value
        for field_name in GB29446_META_FIELDS:
            if field_name not in values:
                issues.append(
                    ImportIssue(
                        severity="error",
                        sheet=GB29446_META_SHEET,
                        cell="A1",
                        message=f"_meta 缺少必需字段：{field_name}",
                    )
                )
                continue
            if values[field_name] in (None, ""):
                issues.append(
                    ImportIssue(
                        severity="error",
                        sheet=GB29446_META_SHEET,
                        cell="B1",
                        message=f"_meta 字段不能为空：{field_name}",
                    )
                )
        if values.get("adapter_id") not in (None, "", gb.GB29446_EXCEL_ADAPTER_ID):
            issues.append(
                ImportIssue(
                    severity="error",
                    sheet=GB29446_META_SHEET,
                    cell="B2",
                    message=f"adapter_id 必须为 {gb.GB29446_EXCEL_ADAPTER_ID}",
                )
            )
        if values.get("template_version") not in (None, "", gb.GB29446_EXCEL_TEMPLATE_VERSION):
            issues.append(
                ImportIssue(
                    severity="error",
                    sheet=GB29446_META_SHEET,
                    cell="B3",
                    message=f"模板版本不受支持，必须为 {gb.GB29446_EXCEL_TEMPLATE_VERSION}",
                )
            )
        return values

    @staticmethod
    def _check_meta_against_standard(
        meta_values: dict[str, object], standard: StandardDefinition, issues: list[ImportIssue]
    ) -> None:
        """Compare template metadata with the definition this evaluation would use."""
        expected = {
            "standard_id": standard.id,
            "standard_version": standard.version,
            "rule_revision": standard.rule_revision,
        }
        mismatch = False
        for key, value in expected.items():
            if meta_values.get(key) in (None, ""):
                continue
            if str(meta_values[key]) != str(value):
                mismatch = True
        if mismatch:
            issues.append(
                ImportIssue(
                    severity="error",
                    sheet=GB29446_META_SHEET,
                    cell="B4",
                    message=(
                        "该模板对应的标准/规则修订已不是当前本次评价适用版本，请重新生成模板后填写。"
                    ),
                )
            )

    @staticmethod
    def _read_gb29446_date(cell, issues: list[ImportIssue]) -> date | None:
        if _formula_text(cell) is not None:
            issues.append(
                ImportIssue(
                    severity="error",
                    sheet=GB29446_DATA_SHEET,
                    cell=cell.coordinate,
                    message="评价日期不能使用 Excel 公式，请填写日期值。",
                )
            )
            return None
        value = cell.value
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        if isinstance(value, str):
            try:
                return date.fromisoformat(value.strip())
            except ValueError:
                issues.append(
                    ImportIssue(
                        severity="error",
                        sheet=GB29446_DATA_SHEET,
                        cell=cell.coordinate,
                        message="评价日期格式应为 yyyy-mm-dd",
                    )
                )
                return None
        issues.append(
            ImportIssue(
                severity="error",
                sheet=GB29446_DATA_SHEET,
                cell=cell.coordinate,
                message="评价日期不能为空",
            )
        )
        return None

    @staticmethod
    def _read_gb29446_choice(
        cell,
        options: list[str],
        *,
        sheet: str,
        cell_ref: str,
        issues: list[ImportIssue],
        field: str,
    ) -> str | None:
        value = _text_from_cell(cell, sheet=sheet, cell_ref=cell_ref, issues=issues, allow_blank=False)
        if value is None:
            return None
        if options and value not in options:
            issues.append(
                ImportIssue(
                    severity="error",
                    sheet=sheet,
                    cell=cell_ref,
                    message=f"{field} 取值不在标准允许范围内：{value}",
                )
            )
            return None
        return value

    def _parse(self, workbook, issues: list[ImportIssue]) -> EvaluationRequest | None:
        """Backwards-compatible alias for the generic parser."""
        return self._parse_generic(workbook, issues)

    def _parse_generic(self, workbook, issues: list[ImportIssue]) -> EvaluationRequest | None:
        info_sheet = workbook["评价信息"]
        info = {}
        for row_number, row in enumerate(info_sheet.iter_rows(min_row=2), start=2):
            key = row[0].value
            if not key:
                continue
            if key in info:
                issues.append(
                    ImportIssue(severity="error", sheet="评价信息", cell=f"A{row_number}", message=f"字段重复：{key}")
                )
                continue
            info[key] = _excel_value(row[1].value)
        if str(info.get("template_version")) != TEMPLATE_VERSION:
            issues.append(
                ImportIssue(severity="error", sheet="评价信息", cell="B2", message="模板版本不受支持")
            )
        evaluation_date = info.get("evaluation_date")
        if isinstance(evaluation_date, datetime):
            evaluation_date = evaluation_date.date()
        elif isinstance(evaluation_date, str):
            try:
                evaluation_date = date.fromisoformat(evaluation_date)
            except ValueError:
                issues.append(
                    ImportIssue(severity="error", sheet="评价信息", cell="B3", message="日期格式应为 yyyy-mm-dd")
                )
        if not isinstance(evaluation_date, date):
            issues.append(ImportIssue(severity="error", sheet="评价信息", cell="B3", message="评价日期不能为空"))
        for key, cell in (("standard_id", "B4"), ("product_id", "B5"), ("input_mode", "B6")):
            if not info.get(key):
                issues.append(ImportIssue(severity="error", sheet="评价信息", cell=cell, message=f"{key} 不能为空"))
        try:
            mode = InputMode(str(info.get("input_mode")))
        except ValueError:
            issues.append(
                ImportIssue(severity="error", sheet="评价信息", cell="B6", message="input_mode 必须为 DIRECT 或 DETAIL")
            )
            mode = InputMode.DIRECT

        inputs: dict[str, InputValue] = {}
        for sheet_name in ("适用条件", "实际值"):
            sheet = workbook[sheet_name]
            for row_number, row in enumerate(sheet.iter_rows(min_row=2), start=2):
                key = row[0].value
                if not key:
                    continue
                raw_value = row[2].value
                if isinstance(raw_value, float):
                    issues.append(
                        ImportIssue(
                            severity="error",
                            sheet=sheet_name,
                            cell=f"C{row_number}",
                            message=(
                                "正式输入不能使用 XLSX 数值单元格：openpyxl 已将其物化为二进制浮点。"
                                "请将单元格设为文本并录入十进制字符串。"
                            ),
                        )
                    )
                    continue
                value = _excel_value(raw_value)
                if value is None or value == "":
                    issues.append(
                        ImportIssue(severity="error", sheet=sheet_name, cell=f"C{row_number}", message="值不能为空")
                    )
                    continue
                if str(key) in inputs:
                    issues.append(
                        ImportIssue(
                            severity="error",
                            sheet=sheet_name,
                            cell=f"A{row_number}",
                            message=f"输入键重复或同时出现在两个输入表：{key}",
                        )
                    )
                    continue
                inputs[str(key)] = InputValue(value=value, unit=row[3].value, source_note=row[4].value)

        energy_lines: list[EnergyLine] = []
        energy_line_ids: set[str] = set()
        for row_number, row in enumerate(workbook["能源明细"].iter_rows(min_row=2), start=2):
            if not row[0].value:
                continue
            line_id = str(row[0].value)
            if line_id in energy_line_ids:
                issues.append(ImportIssue(severity="error", sheet="能源明细", cell=f"A{row_number}", message=f"行ID重复：{line_id}"))
                continue
            energy_line_ids.add(line_id)
            amount = _decimal_from_cell(row[4].value, sheet="能源明细", cell=f"E{row_number}", issues=issues)
            coefficient = _decimal_from_cell(
                row[6].value, sheet="能源明细", cell=f"G{row_number}", issues=issues
            )
            ratio = _decimal_from_cell(row[8].value, sheet="能源明细", cell=f"I{row_number}", issues=issues)
            if None in (amount, coefficient, ratio):
                continue
            try:
                energy_lines.append(
                    EnergyLine(
                        line_id=line_id,
                        energy_name=str(row[1].value or ""),
                        category_key=str(row[2].value) if row[2].value else None,
                        direction=str(row[3].value or "input"),
                        amount=amount,
                        unit=str(row[5].value or ""),
                        standard_coal_coefficient=coefficient,
                        coefficient_unit=str(row[7].value or ""),
                        allocation_ratio=ratio,
                        source_note=row[9].value,
                    )
                )
            except ValueError as exc:
                issues.append(ImportIssue(severity="error", sheet="能源明细", cell=f"A{row_number}", message=str(exc)))

        production_lines: list[ProductionLine] = []
        production_line_ids: set[str] = set()
        for row_number, row in enumerate(workbook["产量与分摊"].iter_rows(min_row=2), start=2):
            if not row[0].value:
                continue
            line_id = str(row[0].value)
            if line_id in production_line_ids:
                issues.append(ImportIssue(severity="error", sheet="产量与分摊", cell=f"A{row_number}", message=f"行ID重复：{line_id}"))
                continue
            production_line_ids.add(line_id)
            quantity = _decimal_from_cell(row[3].value, sheet="产量与分摊", cell=f"D{row_number}", issues=issues)
            factor = _decimal_from_cell(row[5].value, sheet="产量与分摊", cell=f"F{row_number}", issues=issues)
            if None in (quantity, factor):
                continue
            try:
                production_lines.append(
                    ProductionLine(
                        line_id=line_id,
                        product_name=str(row[1].value or ""),
                        category_key=str(row[2].value) if row[2].value else None,
                        quantity=quantity,
                        unit=str(row[4].value or ""),
                        conversion_factor=factor,
                        qualified=str(row[6].value or "是") == "是",
                        source_note=row[7].value,
                    )
                )
            except ValueError as exc:
                issues.append(
                    ImportIssue(severity="error", sheet="产量与分摊", cell=f"A{row_number}", message=str(exc))
                )
        if any(issue.severity == "error" for issue in issues):
            return None
        try:
            request = EvaluationRequest(
                evaluation_date=evaluation_date,
                standard_id=str(info["standard_id"]),
                product_id=str(info["product_id"]),
                input_mode=mode,
                inputs=inputs,
                energy_lines=energy_lines,
                production_lines=production_lines,
                organization_name=info.get("organization_name") or None,
                project_name=info.get("project_name") or None,
                notes=info.get("notes") or None,
            )
        except ValueError as exc:
            issues.append(ImportIssue(severity="error", sheet="工作簿", cell="-", message=str(exc)))
            return None
        if self.standards is not None:
            standard = self.standards.get_published(request.standard_id)
            if standard is None:
                issues.append(
                    ImportIssue(
                        severity="error",
                        sheet="评价信息",
                        cell="B4",
                        message=f"未找到已发布标准：{request.standard_id}",
                    )
                )
                return None
            if not any(product.id == request.product_id for product in standard.products):
                issues.append(
                    ImportIssue(
                        severity="error",
                        sheet="评价信息",
                        cell="B5",
                        message=f"产品/工序不属于标准 {standard.number}：{request.product_id}",
                    )
                )
                return None
        return request


class WorkbookExportService:
    def __init__(self, evaluations: SqlEvaluationRepository, audit: AuditRepository | None = None) -> None:
        self.evaluations = evaluations
        self.audit = audit or evaluations.audit

    def export(self, evaluation_id: str, path: Path) -> Path:
        # A record whose stored payload no longer parses propagates
        # ``StorageCorruptionError`` from the repository untouched.  Its message is
        # Chinese and names the corruption, so the UI reports a storage defect
        # instead of the generic "check your input data and units" wording that
        # ``_friendly_error`` produces for ASCII-only library messages.
        loaded = self.evaluations.get(evaluation_id)
        if loaded is None:
            raise LookupError("评价记录不存在")
        request, result, standard = loaded
        workbook = Workbook()
        workbook.remove(workbook.active)
        for name in EXPORT_SHEETS:
            workbook.create_sheet(name)
        self._summary(workbook["对标结论"], result)
        self._info(workbook["评价信息"], request, result)
        self._inputs(workbook["输入数据"], request)
        self._trace(workbook["计算过程"], result)
        self._sources(workbook["标准依据"], standard, result)
        self._quality(workbook["校验与日志"], result)
        for sheet in workbook.worksheets:
            self._format_sheet(sheet)
        path = path.resolve()
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(path)
        self.audit.append("WORKBOOK_EXPORT", "evaluation", evaluation_id, {"path": str(path)})
        return path

    @staticmethod
    def _title(sheet, title: str, columns: int) -> None:
        sheet.append([title])
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns)
        cell = sheet.cell(1, 1)
        cell.font = Font(size=16, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="left", vertical="center")
        sheet.row_dimensions[1].height = 32

    @staticmethod
    def _table_header(sheet, row: int, headers: list[str]) -> None:
        for column, value in enumerate(headers, start=1):
            cell = sheet.cell(row=row, column=column, value=value)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="5B9BD5")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def _summary(self, sheet, result) -> None:
        self._title(sheet, "单位产品能耗对标结论", 14)
        sheet.append(["标准", f"{result.standard_number} {result.standard_title}"])
        sheet.append(["产品/工序", result.product_name])
        sheet.append([])
        headers = [
            "指标", "实际值", "单位", "1级基础限额", "2级基础限额", "3级基础限额",
            "判定结果", "警告", "1级修正后限额", "2级修正后限额", "3级修正后限额",
            "依据页码", "条款/表号", "规则快照哈希",
        ]
        self._table_header(sheet, 5, headers)
        for item in result.results:
            base = item.base_thresholds
            thresholds = item.corrected_thresholds
            references = "; ".join(
                f"p.{ref.page}" + (f" {ref.clause}" if ref.clause else "") + (f" {ref.table}" if ref.table else "")
                for ref in item.source_references
            )
            sheet.append(
                [
                    item.indicator_name,
                    item.actual_value,
                    item.unit,
                    base.get("LEVEL_1"),
                    base.get("LEVEL_2"),
                    base.get("LEVEL_3"),
                    GRADE_LABELS[item.grade],
                    "；".join(item.warnings),
                    thresholds.get("LEVEL_1"),
                    thresholds.get("LEVEL_2"),
                    thresholds.get("LEVEL_3"),
                    ", ".join(str(ref.page) for ref in item.source_references),
                    "; ".join(filter(None, {ref.clause for ref in item.source_references} | {ref.table for ref in item.source_references})),
                    result.rule_snapshot_sha256,
                ]
            )
        red_fill = PatternFill("solid", fgColor="F4CCCC")
        green_fill = PatternFill("solid", fgColor="D9EAD3")
        sheet.conditional_formatting.add(
            f"G6:G{max(6, sheet.max_row)}",
            FormulaRule(formula=["G6=\"未达标\""], fill=red_fill),
        )
        sheet.conditional_formatting.add(
            f"G6:G{max(6, sheet.max_row)}",
            FormulaRule(formula=["OR(G6=\"1级\",G6=\"2级\",G6=\"3级\")"], fill=green_fill),
        )

    def _info(self, sheet, request, result) -> None:
        self._title(sheet, "评价信息", 2)
        rows = [
            ("评价ID", result.evaluation_id),
            ("评价日期", request.evaluation_date),
            ("计算时间", result.evaluated_at.astimezone().replace(tzinfo=None)),
            ("单位名称", request.organization_name),
            ("项目名称", request.project_name),
            ("标准编号", result.standard_number),
            ("标准版本", result.standard_version),
            ("产品/工序", result.product_name),
            ("录入模式", request.input_mode.value),
            ("备注", request.notes),
        ]
        for row in rows:
            sheet.append(row)
        sheet["B3"].number_format = "yyyy-mm-dd"
        sheet["B4"].number_format = "yyyy-mm-dd hh:mm:ss"

    def _inputs(self, sheet, request) -> None:
        self._title(sheet, "输入数据", 9)
        self._table_header(sheet, 3, ["类型", "键/行ID", "名称", "数值", "单位", "方向/合格", "系数", "分摊/折算", "来源/备注"])
        for key, value in request.inputs.items():
            sheet.append(["输入值", key, "", value.value, value.unit, "", "", "", value.source_note])
        for line in request.energy_lines:
            sheet.append(
                [
                    "能源明细",
                    line.line_id,
                    line.energy_name,
                    line.amount,
                    line.unit,
                    line.direction,
                    line.standard_coal_coefficient,
                    line.allocation_ratio,
                    line.source_note,
                ]
            )
        for line in request.production_lines:
            sheet.append(
                [
                    "产量",
                    line.line_id,
                    line.product_name,
                    line.quantity,
                    line.unit,
                    "是" if line.qualified else "否",
                    "",
                    line.conversion_factor,
                    line.source_note,
                ]
            )

    def _trace(self, sheet, result) -> None:
        self._title(sheet, "计算过程", 8)
        self._table_header(sheet, 3, ["指标", "序号", "步骤", "操作", "表达式", "结果", "单位", "警告"])
        for item in result.results:
            if not item.calculation_trace:
                sheet.append([item.indicator_name, "", "", "", "", "", item.unit, "；".join(item.warnings)])
            for step in item.calculation_trace:
                sheet.append(
                    [
                        item.indicator_name,
                        step.sequence,
                        step.label,
                        step.operation,
                        step.expression,
                        step.value,
                        step.unit,
                        "；".join(item.warnings),
                    ]
                )

    def _sources(self, sheet, standard, result) -> None:
        self._title(sheet, "标准依据", 8)
        self._table_header(sheet, 3, ["标准编号", "标准名称", "原文文件", "原文SHA-256", "页码", "条款", "表号", "备注"])
        seen = set()
        for item in result.results:
            for reference in item.source_references:
                key = (reference.page, reference.clause, reference.table, reference.note)
                if key in seen:
                    continue
                seen.add(key)
                sheet.append(
                    [
                        standard.number,
                        standard.title,
                        reference.source_file,
                        reference.source_sha256,
                        reference.page,
                        reference.clause,
                        reference.table,
                        reference.note,
                    ]
                )

    def _quality(self, sheet, result) -> None:
        self._title(sheet, "校验与日志", 3)
        self._table_header(sheet, 3, ["类型", "对象", "信息"])
        sheet.append(["规则快照", result.standard_number, result.rule_snapshot_sha256])
        for warning in result.warnings:
            sheet.append(["评价警告", result.evaluation_id, warning])
        for item in result.results:
            for warning in item.warnings:
                sheet.append(["指标警告", item.indicator_name, warning])
        if sheet.max_row == 4:
            sheet.append(["校验", "评价结果", "未发现警告"])

    @staticmethod
    def _format_sheet(sheet) -> None:
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "A4"
        sheet.auto_filter.ref = sheet.dimensions if sheet.max_row >= 3 else None
        thin = Side(style="thin", color="D9E2F3")
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if cell.row > 1:
                    cell.border = Border(bottom=thin)
                if isinstance(cell.value, Decimal):
                    cell.number_format = "0.########"
        for column_index, column_cells in enumerate(sheet.columns, start=1):
            letter = get_column_letter(column_index)
            maximum = max((len(str(cell.value)) for cell in column_cells if cell.value is not None), default=8)
            sheet.column_dimensions[letter].width = min(max(maximum + 2, 12), 48)
