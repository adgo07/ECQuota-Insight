"""只消费应用层报告视图的 DOCX 写入器。"""

from __future__ import annotations

import os
import re
import tempfile
from pathlib import Path

from uebench.application.reporting import ReportField, ReportView
from uebench.application.cement_reporting import CementReportView


def _add_subscript(paragraph, text: str) -> None:
    """把 E_d / e_d 写成真正的 Word 下标，而非靠 HTML 解释。"""
    for part in re.split(r"([EeWwPpQ]_[A-Za-z0-9]+)", text):
        if not part:
            continue
        if re.fullmatch(r"[EeWwPpQ]_[A-Za-z0-9]+", part):
            paragraph.add_run(part[0])
            suffix = paragraph.add_run(part.split("_", 1)[1])
            suffix.font.subscript = True
        else:
            paragraph.add_run(part)


def _row(table, field: ReportField) -> None:
    cells = table.add_row().cells
    _add_subscript(cells[0].paragraphs[0], field.label)
    value = f"{field.value} {field.unit}".strip()
    _add_subscript(cells[1].paragraphs[0], value)
    for cell in cells:
        cell.paragraphs[0].paragraph_format.space_after = 0


class WordReportWriter:
    """创建文件时先写同目录临时文件，成功后原子替换目标。"""

    def write(self, view: ReportView | CementReportView, path: Path) -> Path:
        if isinstance(view, CementReportView):
            return self._write_cement(view, path)
        from docx import Document
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Cm, Pt, RGBColor
        from docx.oxml.ns import qn

        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        document = Document()
        section = document.sections[0]
        section.top_margin = Cm(1.6)
        section.bottom_margin = Cm(1.6)
        section.left_margin = Cm(1.9)
        section.right_margin = Cm(1.9)
        section.page_height = Cm(29.7)
        section.page_width = Cm(21)
        normal = document.styles["Normal"]
        normal.font.name = "Microsoft YaHei"
        normal.font.size = Pt(9)
        normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        normal.paragraph_format.space_after = Pt(2)
        heading = document.styles["Heading 1"]
        heading.font.name = "Microsoft YaHei"
        heading._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        heading.font.size = Pt(11)
        heading.font.color.rgb = RGBColor(24, 55, 83)
        heading.paragraph_format.space_before = Pt(7)
        heading.paragraph_format.space_after = Pt(3)
        title_style = document.styles["Title"]
        title_style.font.name = "Microsoft YaHei"
        title_style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        title_style.font.size = Pt(15)
        title_style.font.bold = True
        title_style.font.color.rgb = RGBColor(0, 0, 0)
        title_style.paragraph_format.space_after = Pt(6)
        # Word 内置 Title 有时携带蓝色下边框；清除样式与段落边框。
        for element in (title_style._element,):
            ppr = element.pPr
            if ppr is not None:
                border = ppr.find(qn("w:pBdr"))
                if border is not None:
                    ppr.remove(border)
        title = document.add_paragraph(style="Title")
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _add_subscript(title, view.title)
        if title._p.pPr is not None:
            border = title._p.pPr.find(qn("w:pBdr"))
            if border is not None:
                title._p.pPr.remove(border)

        document.add_heading(view.headings[0], level=1)
        first = document.add_table(rows=0, cols=2)
        first.style = "Table Grid"
        first.alignment = WD_TABLE_ALIGNMENT.CENTER
        for field in (
            view.standard, view.evaluation_date, view.evaluated_at,
            view.organization, view.period, view.enterprise_type, view.process_type,
            view.electricity, view.raw_coal, view.factor,
        ):
            if field is not None:
                _row(first, field)

        document.add_heading(view.headings[1], level=1)
        result = document.add_table(rows=0, cols=2)
        result.style = "Table Grid"
        for field in (view.actual, view.grade, *view.thresholds):
            _row(result, field)

        document.add_heading(view.headings[2], level=1)
        for line in view.explanation:
            _add_subscript(document.add_paragraph(), line)
        for warning in view.warnings:
            _add_subscript(document.add_paragraph(), warning)

        document.add_heading(view.headings[3], level=1)
        scope = document.add_paragraph()
        _add_subscript(scope, f"{view.scope.label}：{view.scope.value}")
        for basis in view.basis:
            document.add_paragraph(basis)
        # 历史记录不得宣称原评价已采用后来批准的映射。
        document.add_paragraph(view.mapping_adoption)

        temporary: Path | None = None
        try:
            with tempfile.NamedTemporaryFile(prefix=f".{path.stem}-", suffix=".docx", dir=path.parent, delete=False) as handle:
                temporary = Path(handle.name)
            document.save(temporary)
            os.replace(temporary, path)
        finally:
            if temporary is not None:
                temporary.unlink(missing_ok=True)
        return path

    def _write_cement(self, view: CementReportView, path: Path) -> Path:
        from docx import Document
        from docx.shared import Cm, Pt, RGBColor
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        if not view.complete:
            raise ValueError("部分计算预览不能导出正式 Word 报告。")
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        document = Document()
        section = document.sections[0]
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.top_margin = section.bottom_margin = Cm(1.7)
        section.left_margin = section.right_margin = Cm(1.8)
        for name, size in (("Normal", 9), ("Title", 15), ("Heading 1", 11)):
            style = document.styles[name]
            style.font.name = "Microsoft YaHei"
            style.font.size = Pt(size)
            style.font.color.rgb = RGBColor(0, 0, 0)
            style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
            style.paragraph_format.space_after = Pt(4)
        title = document.add_paragraph(view.title, "Title")
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        document.add_paragraph("；".join(f"{f.label}：{f.value}" for f in view.information))
        table = document.add_table(rows=1, cols=6)
        table.style = "Table Grid"
        for cell, label in zip(table.rows[0].cells, ("指标", "实际值", "等级", "1级上限", "2级上限", "3级上限")):
            cell.text = label
        header = OxmlElement("w:tblHeader")
        table.rows[0]._tr.get_or_add_trPr().append(header)
        for item in view.indicators:
            thresholds = {field.key: field.value for field in item.thresholds}
            values = (item.label + f"（{item.unit}）", item.actual, item.grade,
                      *(thresholds.get(key, "—") for key in ("LEVEL_1", "LEVEL_2", "LEVEL_3")))
            for cell, value in zip(table.add_row().cells, values):
                cell.text = value
        document.add_paragraph(view.scope_note)
        for line in view.disclosures:
            document.add_paragraph(line)
        document.add_heading("输入数据与计算依据", level=1)
        raw = document.add_table(rows=0, cols=2)
        raw.style = "Table Grid"
        for field in view.inputs:
            _row(raw, field)
        if view.calculations:
            document.add_heading("计算依据", level=1)
            for field in view.calculations:
                document.add_paragraph(f"{field.label}：{field.value} {field.unit}".strip())
        for line in view.basis:
            document.add_paragraph(line)
        temporary = None
        try:
            with tempfile.NamedTemporaryFile(prefix=f".{path.stem}-", suffix=".docx", dir=path.parent, delete=False) as handle:
                temporary = Path(handle.name)
            document.save(temporary)
            os.replace(temporary, path)
        finally:
            if temporary is not None:
                temporary.unlink(missing_ok=True)
        return path
