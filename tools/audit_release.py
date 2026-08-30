from __future__ import annotations

"""Audit the user-facing release directory without modifying any project data."""

import argparse
import hashlib
import json
import zipfile
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


REQUIRED_FILES = (
    "UEBench-0.1.0-win-x64.zip",
    "UEBench-Setup-0.1.0-x64.exe",
    "UEBench-source-0.1.0.zip",
    "initial-standard-package-published.uebench",
    "统一标准规则确认表.xlsx",
    "单位产品能耗对标导入模板.xlsx",
)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def audit_release(root: Path) -> dict:
    errors: list[str] = []
    hashes_path = root / "SHA256SUMS.txt"
    expected_hashes: dict[str, str] = {}
    if not hashes_path.exists():
        errors.append("缺少 SHA256SUMS.txt")
    else:
        for line_number, line in enumerate(hashes_path.read_text(encoding="utf-8-sig").splitlines(), start=1):
            parts = line.split()
            if len(parts) != 2:
                errors.append(f"SHA256SUMS.txt 第{line_number}行格式错误")
                continue
            expected_hashes[parts[1]] = parts[0].lower()

    file_report: dict[str, dict[str, object]] = {}
    for name in REQUIRED_FILES:
        path = root / name
        exists = path.exists()
        actual = sha256(path) if exists else None
        expected = expected_hashes.get(name)
        file_report[name] = {"exists": exists, "sha256": actual, "matches": exists and actual == expected}
        if not exists:
            errors.append(f"缺少交付文件：{name}")
        elif expected is None:
            errors.append(f"SHA256SUMS.txt 未登记：{name}")
        elif actual != expected:
            errors.append(f"文件哈希不匹配：{name}")

    manifest_report: dict[str, object] = {}
    expected_current_rules = None
    package = root / "initial-standard-package-published.uebench"
    if package.exists():
        try:
            with zipfile.ZipFile(package) as archive:
                manifest = json.loads(archive.read("manifest.json"))
                definitions = [json.loads(archive.read(name)) for name in archive.namelist() if name.startswith("definitions/") and name.endswith(".json")]
            current_definitions = [d for d in definitions if d.get("lifecycle_status", "active") != "obsolete"]
            history_definitions = [d for d in definitions if d.get("lifecycle_status") == "obsolete"]
            expected_current_rules = sum(len(product.get("indicators", [])) for definition in current_definitions for product in definition.get("products", []))
            manifest_report = {
                "standard_count": manifest.get("standard_count"),
                "rule_count": manifest.get("rule_count"),
                "data_version": manifest.get("data_version"),
                "current_standard_count": len(current_definitions),
                "historical_standard_count": len(history_definitions),
                "current_rule_count": expected_current_rules,
            }
            if not isinstance(manifest.get("standard_count"), int) or not isinstance(manifest.get("rule_count"), int):
                errors.append("正式标准包清单缺少有效的标准数量或规则数量")
        except Exception as exc:
            errors.append(f"候选标准包无法读取：{exc}")

    confirmation_report: dict[str, object] = {}
    confirmation = root / "统一标准规则确认表.xlsx"
    if confirmation.exists():
        try:
            workbook = load_workbook(confirmation, read_only=True, data_only=True)
            if workbook.sheetnames != ["确认汇总", "规则确认", "订正记录"]:
                errors.append("规则确认表工作表结构不匹配")
            sheet = workbook["规则确认"]
            headers = {cell.value: index for index, cell in enumerate(next(sheet.iter_rows(min_row=3, max_row=3)))}
            required_headers = {"指标ID", "确认结论", "确认人", "确认日期"}
            if not required_headers <= set(headers):
                errors.append("规则确认表缺少必要列")
            rows = [row for row in sheet.iter_rows(min_row=4) if row[headers["指标ID"]].value]
            conclusions = Counter(str(row[headers["确认结论"]].value) for row in rows)
            confirmation_report = {
                "indicator_rows": len(rows),
                "conclusions": dict(conclusions),
                "published_ready": (
                    expected_current_rules is not None
                    and len(rows) == expected_current_rules
                    and conclusions == Counter({"同意发布": expected_current_rules})
                    and all(row[headers["确认人"]].value for row in rows)
                    and all(row[headers["确认日期"]].value for row in rows)
                ),
            }
            if expected_current_rules is not None and len(rows) != expected_current_rules:
                errors.append(f"规则确认表指标行数与当前标准包不一致：确认表{len(rows)}，当前规则{expected_current_rules}")
            if not confirmation_report["published_ready"]:
                errors.append("规则确认表尚未完成全部指标的同意发布确认，当前交付只能作为候选包")
        except Exception as exc:
            errors.append(f"规则确认表无法读取：{exc}")

    return {
        "valid": not errors,
        "release_dir": str(root),
        "files": file_report,
        "published_package": manifest_report,
        "confirmation": confirmation_report,
        "errors": errors,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="审计UEBench最终交付目录，不修改任何数据")
    parser.add_argument("release_dir", type=Path)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    report = audit_release(args.release_dir.resolve())
    payload = json.dumps(report, ensure_ascii=False, indent=2)
    print(payload)
    if args.output:
        args.output.resolve().parent.mkdir(parents=True, exist_ok=True)
        args.output.resolve().write_text(payload + "\n", encoding="utf-8")
    if not report["valid"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
