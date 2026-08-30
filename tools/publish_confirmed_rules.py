from __future__ import annotations

import argparse
import json
import sys
from datetime import datetime, timezone
from pathlib import Path

from cryptography.hazmat.primitives import serialization
from cryptography.hazmat.primitives.asymmetric.ed25519 import Ed25519PrivateKey
from openpyxl import load_workbook

from uebench.domain.models import PublicationStatus, StandardDefinition
from uebench.infrastructure.packages import StandardPackageBuilder


def validate_confirmation(
    path: Path,
    reviewed_ids: set[str],
    published_ids: set[str],
    all_ids: set[str],
    *,
    scope_count: int,
) -> set[str]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if "规则确认" not in workbook.sheetnames:
        raise ValueError("确认文件缺少“规则确认”工作表")
    sheet = workbook["规则确认"]
    headers = {cell.value: index for index, cell in enumerate(next(sheet.iter_rows(min_row=3, max_row=3)))}
    required = {"指标ID", "确认结论", "确认人", "确认日期"}
    missing = required - set(headers)
    if missing:
        raise ValueError(f"确认文件缺少列：{', '.join(sorted(missing))}")
    confirmed: set[str] = set()
    seen: set[str] = set()
    errors: list[str] = []
    for row_number, row in enumerate(sheet.iter_rows(min_row=4), start=4):
        indicator_id = row[headers["指标ID"]].value
        if not indicator_id:
            continue
        indicator_key = str(indicator_id)
        if indicator_key in seen:
            errors.append(f"指标ID重复：{indicator_id}")
        seen.add(indicator_key)
        conclusion = row[headers["确认结论"]].value
        reviewer = row[headers["确认人"]].value
        confirmed_at = row[headers["确认日期"]].value
        if indicator_key not in all_ids:
            errors.append(f"第{row_number}行 {indicator_id} 不属于当前{scope_count}项规则")
            continue
        # A published row may remain in a historical confirmation workbook.
        # It is immutable and does not need to be re-confirmed on later runs.
        if indicator_key in published_ids:
            continue
        if indicator_key in reviewed_ids:
            if conclusion != "同意发布":
                errors.append(f"第{row_number}行 {indicator_id} 尚未同意发布")
            if not reviewer:
                errors.append(f"第{row_number}行 {indicator_id} 缺少确认人")
            if not confirmed_at:
                errors.append(f"第{row_number}行 {indicator_id} 缺少确认日期")
        elif conclusion == "同意发布":
            errors.append(f"第{row_number}行 {indicator_id} 当前仍为draft，完成原文复核和规则录入后才能确认发布")
        if indicator_key in reviewed_ids and conclusion == "同意发布":
            confirmed.add(indicator_key)
    if errors:
        raise ValueError("\n".join(errors))
    return confirmed


def main() -> None:
    parser = argparse.ArgumentParser(description="核验规则确认表并生成published离线标准包")
    parser.add_argument("confirmation", type=Path)
    parser.add_argument("source_dir", type=Path)
    parser.add_argument("--data-dir", type=Path, default=Path("data"), help="规则数据目录，默认data")
    parser.add_argument("--private-key", type=Path, default=Path("work/signing/development-private-key.pem"))
    parser.add_argument("--output", type=Path, default=Path("dist/standard-packages/initial-standard-package-published.uebench"))
    parser.add_argument("--data-version", default="2026.08-published.1", help="写入标准包清单的数据版本")
    parser.add_argument("--apply", action="store_true", help="确认验证通过后实际写入published状态并构建标准包")
    args = parser.parse_args()

    data_dir = args.data_dir.resolve()
    definition_paths = sorted((data_dir / "definitions").glob("*.json"))
    definitions = [StandardDefinition.model_validate_json(path.read_text(encoding="utf-8")) for path in definition_paths]
    if not definitions:
        raise ValueError(f"数据目录没有规则定义：{data_dir}")
    scope_count = len(definitions)
    reviewed_ids = {
        indicator.id
        for definition in definitions
        if definition.publication_status is PublicationStatus.REVIEWED
        for product in definition.products
        for indicator in product.indicators
    }
    published_ids = {
        indicator.id
        for definition in definitions
        if definition.publication_status is PublicationStatus.PUBLISHED
        for product in definition.products
        for indicator in product.indicators
    }
    all_ids = {
        indicator.id
        for definition in definitions
        for product in definition.products
        for indicator in product.indicators
    }
    try:
        confirmed = validate_confirmation(args.confirmation.resolve(), reviewed_ids, published_ids, all_ids, scope_count=scope_count)
    except (OSError, ValueError) as exc:
        lines = str(exc).splitlines()
        preview = lines[:20]
        if len(lines) > len(preview):
            preview.append(f"……其余{len(lines) - len(preview)}条错误已省略")
        print("确认表校验失败：\n" + "\n".join(preview), file=sys.stderr)
        raise SystemExit(2) from exc
    if confirmed != reviewed_ids:
        missing = sorted(reviewed_ids - confirmed)
        extra = sorted(confirmed - reviewed_ids)
        raise ValueError(f"确认指标集合不一致；缺少={missing}；多出={extra}")
    print(f"确认表验证通过：{len(confirmed)}个指标。")
    if not args.apply:
        print("未指定 --apply，未修改规则状态、未生成published标准包。")
        return

    published_definitions: list[StandardDefinition] = []
    for path, definition in zip(definition_paths, definitions, strict=True):
        if definition.publication_status is PublicationStatus.REVIEWED:
            definition = definition.model_copy(update={"publication_status": PublicationStatus.PUBLISHED})
            path.write_text(definition.model_dump_json(indent=2) + "\n", encoding="utf-8")
        published_definitions.append(definition)
    catalog_path = data_dir / "catalog.json"
    catalog = json.loads(catalog_path.read_text(encoding="utf-8"))
    published_numbers = {item.number for item in published_definitions if item.publication_status is PublicationStatus.PUBLISHED}
    for item in catalog["standards"]:
        if item["number"] in published_numbers:
            item["status"] = "published"
    catalog_path.write_text(json.dumps(catalog, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    private_key = serialization.load_pem_private_key(args.private_key.resolve().read_bytes(), password=None)
    if not isinstance(private_key, Ed25519PrivateKey):
        raise TypeError("签名私钥不是Ed25519私钥")
    source_dir = args.source_dir.resolve()
    source_files = {definition.source_file: source_dir / definition.source_file for definition in published_definitions}
    corrections_path = data_dir / "corrections" / "catalog-corrections.json"
    corrections = json.loads(corrections_path.read_text(encoding="utf-8"))["corrections"] if corrections_path.exists() else []
    StandardPackageBuilder(private_key).build(
        args.output.resolve(),
        published_definitions,
        source_files,
        data_version=args.data_version,
        minimum_app_version="0.1.0",
        corrections=corrections,
        package_id=f"initial-{scope_count}-standards-published-{datetime.now():%Y%m%d%H%M%S}",
        issued_at=datetime.now(timezone.utc),
    )
    print(args.output.resolve())


if __name__ == "__main__":
    main()
