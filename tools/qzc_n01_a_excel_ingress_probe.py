# -*- coding: utf-8 -*-
"""ECQ N01-A / RS03 Excel authoritative-ingress probe.

Proves the current Excel ingress contract through the real adapter services
rather than through private helpers:

* an XLSX numeric **float** cell is rejected;
* an XLSX numeric **integer** cell is rejected;
* XLSX decimal **lexical text** is accepted, and the original decimal text is
  preserved without a ``float -> str`` round trip;
* an Excel **formula** cell is rejected even when a cached value exists.

This is a deliberate tightening of the Excel ingress contract.  It does not
change the Numeric Contract, the project Numeric Profile, or any business
threshold.
"""
from __future__ import annotations

import hashlib
import re
import sys
import zipfile
from datetime import date
from decimal import Decimal
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

from uebench.application import gb29446 as gb
from uebench.domain.models import StandardDefinition
from uebench.infrastructure.database import DatabaseManager
from uebench.infrastructure.excel import GB29446_DATA_SHEET, WorkbookImportService, WorkbookTemplateService
from uebench.infrastructure.paths import AppPaths
from uebench.infrastructure.repositories import AuditRepository, SqlStandardRepository

# CI runs this probe on a Windows console whose default encoding may be a legacy
# code page (for example cp1252), which cannot encode CJK issue messages.  Force
# UTF-8 with a backslashreplace fallback so the probe reports findings instead of
# dying on its own diagnostic output.
for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(encoding="utf-8", errors="backslashreplace")  # type: ignore[union-attr]
    except (AttributeError, ValueError):  # pragma: no cover - non-reconfigurable stream
        pass

ROOT = Path(__file__).resolve().parents[1]
STANDARD_PATH = ROOT / "data" / "definitions" / "gb-29446-2019.json"


def _services(root: Path):
    paths = AppPaths.from_root(root / "appdata")
    paths.ensure()
    database = DatabaseManager(paths.database)
    database.initialize()
    audit = AuditRepository(database)
    standards = SqlStandardRepository(database, audit)
    standard = StandardDefinition.model_validate_json(STANDARD_PATH.read_text(encoding="utf-8"))
    standards.install(standard)
    return database, audit, standards, standard


def _fill(path: Path, electricity, raw_coal="100", *, date_value=date(2026, 6, 1)) -> Path:
    workbook = load_workbook(path)
    sheet = workbook[GB29446_DATA_SHEET]
    positions = {
        gb.FIELD_ORGANIZATION_NAME: 2,
        gb.FIELD_EVALUATION_DATE: 3,
        gb.FIELD_PERIOD: 4,
        gb.FIELD_CUSTOM_PERIOD: 5,
        gb.FIELD_COAL_TYPE: 6,
        gb.FIELD_WASHING_PROCESS: 7,
        gb.FIELD_ELECTRICITY: 8,
        gb.FIELD_RAW_COAL: 9,
        gb.FIELD_NOTES: 10,
    }
    sheet.cell(row=positions[gb.FIELD_EVALUATION_DATE], column=2, value=date_value)
    sheet.cell(row=positions[gb.FIELD_PERIOD], column=2, value=gb.PERIOD_FULL_YEAR)
    sheet.cell(row=positions[gb.FIELD_COAL_TYPE], column=2, value="炼焦煤")
    sheet.cell(row=positions[gb.FIELD_WASHING_PROCESS], column=2, value="重介")
    electricity_cell = sheet.cell(row=positions[gb.FIELD_ELECTRICITY], column=2)
    if isinstance(electricity, str) or electricity is None:
        electricity_cell.value = electricity
    else:
        # Write a genuine XLSX numeric cell (no number_format override).
        electricity_cell.value = electricity
        electricity_cell.number_format = "General"
    raw_cell = sheet.cell(row=positions[gb.FIELD_RAW_COAL], column=2)
    raw_cell.value = raw_coal
    if isinstance(raw_coal, str):
        raw_cell.number_format = "@"
    workbook.save(path)
    return path


def _inject_formula_with_cache(path: Path, formula: str, cached: str) -> Path:
    """Store a cached result for the formula cell in the sheet XML.

    openpyxl can author the formula but never stores a cached value, so the
    sheet XML is edited directly to reproduce a workbook saved by Excel itself.
    The importer must reject it based on the formula view alone, regardless of
    what the cached value says.
    """
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        payload = {name: archive.read(name) for name in names}
    target = "xl/worksheets/sheet2.xml"
    text = payload[target].decode("utf-8")
    # openpyxl writes the formula cell as <c r="B8" s="3"><f>FORMULA</f><v /></c>
    pattern = re.compile(r'(<c r="B8"[^>]*>)(<f>)([^<]*)(</f>)(<v\s*/>)')
    match = pattern.search(text)
    if not match:
        pattern = re.compile(r'(<c r="B8"[^>]*>)(<f>)([^<]*)(</f>)(<v>[^<]*</v>)')
        match = pattern.search(text)
    if not match:
        raise RuntimeError("unexpected sheet XML layout for the electricity cell")
    replacement = f'{match.group(1)}<f>{formula}</f><v>{cached}</v>'
    payload[target] = (text[: match.start()] + replacement + text[match.end():]).encode("utf-8")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name in names:
            archive.writestr(name, payload[target] if name == target else payload[name])
    return path


def main() -> int:
    with TemporaryDirectory() as temp_dir:
        root = Path(temp_dir)
        database, audit, standards, _standard = _services(root)
        service = WorkbookImportService(database, audit, standards)
        template_service = WorkbookTemplateService()

        results: list[tuple[str, bool]] = []

        def check(name: str, condition: bool, detail: str = "") -> None:
            results.append((name, bool(condition)))
            suffix = f" [{detail}]" if detail else ""
            print(f"  [{'PASS' if condition else 'FAIL'}] {name}{suffix}")

        print("=" * 70)
        print("RS03 Excel authoritative ingress probe")
        print("=" * 70)

        print("\n-- accepted: decimal lexical text --")
        for text in ("560", "560.25", "1e3"):
            path = template_service.create_template(root / f"ok-{text}.xlsx", _standard)
            _fill(path, text)
            report = service.validate(path)
            check(f"text {text!r} accepted", report.valid, _first_error(report))

        print("\n-- lexical preservation (no float round trip) --")
        path = template_service.create_template(root / "lexical.xlsx", _standard)
        _fill(path, "5.0000004")
        report = service.validate(path)
        preserved = report.valid and report.request is not None and (
            report.request.inputs[gb.INPUT_ELECTRICITY].value == "5.0000004"
        )
        check("text '5.0000004' keeps its lexical value", preserved,
              "" if preserved else _first_error(report))
        # The lexical contract matters because a binary float cannot represent
        # every decimal literal a user may type.  Use a value whose float form
        # really does lose information, so the assertion is not at the mercy of
        # how a given interpreter formats floats.
        lossy = "5.0000000000000001"
        check(
            "a binary float could not preserve every decimal literal",
            Decimal(lossy) != Decimal(str(float(lossy))),
            f"float({lossy!r}) -> {str(float(lossy))!r}",
        )

        print("\n-- rejected: XLSX numeric cells --")
        for label, value in (("numeric int 560", 560), ("numeric float 560.25", 560.25)):
            path = template_service.create_template(root / f"num-{label.split()[-1]}.xlsx", _standard)
            _fill(path, value)
            report = service.validate(path)
            check(f"{label} rejected", not report.valid, _first_error(report))

        print("\n-- rejected: Excel formula --")
        path = template_service.create_template(root / "formula.xlsx", _standard)
        _fill(path, None)
        workbook = load_workbook(path)
        workbook[GB29446_DATA_SHEET]["B8"] = "=500+60"
        workbook.save(path)
        report = service.validate(path)
        check("formula cell rejected (no cached value)", not report.valid, _first_error(report))

        path = template_service.create_template(root / "formula-cached.xlsx", _standard)
        _fill(path, None)
        workbook = load_workbook(path)
        workbook[GB29446_DATA_SHEET]["B8"] = "=500+60"
        workbook.save(path)
        _inject_formula_with_cache(path, "500+60", "560")
        loaded = load_workbook(path, data_only=False)
        has_formula = loaded[GB29446_DATA_SHEET]["B8"].data_type == "f"
        check("test workbook really contains <f>", has_formula)
        cached_view = load_workbook(path, data_only=True)
        check("cached value view shows 560", cached_view[GB29446_DATA_SHEET]["B8"].value == 560,
              repr(cached_view[GB29446_DATA_SHEET]["B8"].value))
        report = service.validate(path)
        check("formula cell rejected even with cached value", not report.valid, _first_error(report))

        print("\n-- rejected: non-finite text --")
        for text in ("NaN", "sNaN", "Infinity", "-Infinity", "inf", "-inf"):
            path = template_service.create_template(root / f"nf-{text.strip('-')}.xlsx", _standard)
            _fill(path, text)
            report = service.validate(path)
            check(f"non-finite {text!r} rejected", not report.valid, _first_error(report))

        print("\n-- rejected: malformed numeric text --")
        for text in ("1,000", "560 kWh", "abc"):
            path = template_service.create_template(root / f"bad-{hashlib.md5(text.encode()).hexdigest()[:6]}.xlsx", _standard)
            _fill(path, text)
            report = service.validate(path)
            check(f"malformed {text!r} rejected", not report.valid, _first_error(report))

        database.dispose()

        failed = [name for name, ok in results if not ok]
        print("\n" + "=" * 70)
        print(f"checks: {len(results)}  failed: {len(failed)}")
        for name in failed:
            print(f"  FAILED: {name}")
        if not failed:
            print("finding: GB29446 authoritative Excel ingress is decimal-lexical-text only;")
            print("         XLSX numeric cells (int or float) and Excel formulas are rejected,")
            print("         so the adapter can never manufacture a business value.")
        return_code = 1 if failed else 0

    return return_code


def _first_error(report) -> str:
    for issue in report.issues:
        if issue.severity == "error":
            return f"{issue.sheet}!{issue.cell}: {issue.message}"
    return ""


if __name__ == "__main__":
    raise SystemExit(main())
