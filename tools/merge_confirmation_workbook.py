from __future__ import annotations

import argparse
from pathlib import Path
from openpyxl import load_workbook


def merge(base: Path, source: Path, output: Path) -> dict:
    wb = load_workbook(base)
    src = load_workbook(source, read_only=True, data_only=False)
    dst_sheet = wb['规则确认']
    src_sheet = src['规则确认']
    headers = [cell.value for cell in next(dst_sheet.iter_rows(min_row=3, max_row=3))]
    source_headers = [cell.value for cell in next(src_sheet.iter_rows(min_row=3, max_row=3))]
    index = {value: pos for pos, value in enumerate(headers)}
    source_index = {value: pos for pos, value in enumerate(source_headers)}
    source_rows = {
        str(row[source_index['指标ID']].value): row
        for row in src_sheet.iter_rows(min_row=4)
        if row[source_index['指标ID']].value
    }
    copy_columns = ['复核状态', '确认结论', '确认人', '确认日期', '确认备注']
    matched = 0
    pending = []
    for row in dst_sheet.iter_rows(min_row=4):
        key = row[index['指标ID']].value
        if not key:
            continue
        source_row = source_rows.get(str(key))
        if source_row is None:
            pending.append({
                'standard_number': row[index['标准编号']].value,
                'product': row[index['产品/工序']].value,
                'indicator': row[index['指标名称']].value,
                'indicator_id': key,
            })
            continue
        for column in copy_columns:
            row[index[column]].value = source_row[source_index[column]].value
        matched += 1
    note_row = dst_sheet.max_row + 2
    summary = wb['确认汇总']
    summary.cell(note_row, 1).value = (
        '合并说明：已按指标ID精确继承旧确认表中的人工确认记录；新版或替代标准中没有匹配记录的行保持“待确认”，'
        '不得把旧版确认自动视为新版确认。'
    )
    summary.merge_cells(start_row=note_row, start_column=1, end_row=note_row + 1, end_column=7)
    summary.cell(note_row, 1).alignment = summary.cell(note_row, 1).alignment.copy(wrap_text=True, vertical='top')
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    return {'matched': matched, 'pending': pending, 'output': str(output)}


def main() -> None:
    parser = argparse.ArgumentParser(description='按指标ID合并标准规则确认表，防止旧标准确认误套新版')
    parser.add_argument('base', type=Path)
    parser.add_argument('source', type=Path)
    parser.add_argument('output', type=Path)
    args = parser.parse_args()
    result = merge(args.base, args.source, args.output)
    print(result)
    for item in result['pending']:
        print('待确认:', item)


if __name__ == '__main__':
    main()
