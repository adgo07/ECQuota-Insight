from __future__ import annotations

"""Promote draft rules to ``reviewed`` after a completed confirmation workbook.

This is the first half of the release gate.  It deliberately does not publish a
package; ``publish_confirmed_rules.py`` remains the second, separate step.
Only draft indicators are required to be confirmed, so a workbook can contain
the already-published rows without forcing users to re-sign historical rules.
"""

import argparse
import json
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from uebench.domain.models import PublicationStatus, StandardDefinition


def _confirmation_rows(path: Path) -> tuple[dict[str, dict[str, object]], list[str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if "规则确认" not in workbook.sheetnames:
        raise ValueError("确认文件缺少“规则确认”工作表")
    sheet = workbook["规则确认"]
    header_row = next(sheet.iter_rows(min_row=3, max_row=3))
    headers = {str(cell.value): index for index, cell in enumerate(header_row) if cell.value}
    required = {"标准编号", "指标ID", "确认结论", "确认人", "确认日期"}
    missing = required - set(headers)
    if missing:
        raise ValueError(f"确认文件缺少列：{', '.join(sorted(missing))}")
    rows: dict[str, dict[str, object]] = {}
    errors: list[str] = []
    for row_number, row in enumerate(sheet.iter_rows(min_row=4), start=4):
        indicator_id = row[headers["指标ID"]].value
        if not indicator_id:
            continue
        key = str(indicator_id).strip()
        if key in rows:
            errors.append(f"第{row_number}行指标ID重复：{key}")
            continue
        rows[key] = {
            "row": row_number,
            "standard_number": row[headers["标准编号"]].value,
            "conclusion": row[headers["确认结论"]].value,
            "reviewer": row[headers["确认人"]].value,
            "confirmed_at": row[headers["确认日期"]].value,
        }
    return rows, errors


def _load_scope_numbers(data_dir: Path, scope_path: Path | None) -> set[str]:
    if scope_path is None:
        candidates = (data_dir / "scope-63.json", data_dir / "scope-44.json")
        scope_path = next((candidate for candidate in candidates if candidate.exists()), None)
    if scope_path is None or not scope_path.exists():
        raise ValueError(f"找不到当前范围清单：{data_dir}")
    payload = json.loads(scope_path.read_text(encoding="utf-8"))
    raw_numbers = [str(item).strip() for item in payload.get("standards", []) if str(item).strip()]
    if not raw_numbers:
        raise ValueError(f"范围清单没有标准编号：{scope_path}")
    if len(raw_numbers) != len(set(raw_numbers)):
        raise ValueError(f"范围清单存在重复标准编号：{scope_path}")
    return set(raw_numbers)


def promote(path: Path, data_dir: Path, *, apply: bool, scope_path: Path | None = None) -> dict[str, object]:
    scope_numbers = _load_scope_numbers(data_dir, scope_path)
    all_paths = sorted((data_dir / "definitions").glob("*.json"))
    all_definitions = [StandardDefinition.model_validate_json(item.read_text(encoding="utf-8")) for item in all_paths]
    if not all_definitions:
        raise ValueError(f"数据目录没有规则定义：{data_dir}")
    selected = [(definition_path, definition) for definition_path, definition in zip(all_paths, all_definitions, strict=True) if definition.number in scope_numbers]
    found_numbers = {definition.number for _, definition in selected}
    missing_numbers = sorted(scope_numbers - found_numbers)
    if missing_numbers:
        missing_text = ", ".join(missing_numbers)
        raise ValueError(f"范围清单中的标准缺少规则定义：{missing_text}")
    if len(selected) != len(found_numbers):
        raise ValueError("当前范围内存在重复标准编号的规则定义")
    definition_paths = [definition_path for definition_path, _ in selected]
    definitions = [definition for _, definition in selected]
    by_id = {indicator.id: (definition, path) for definition, path in zip(definitions, definition_paths, strict=True) for product in definition.products for indicator in product.indicators}
    draft_ids = {indicator_id for indicator_id, (definition, _) in by_id.items() if definition.publication_status is PublicationStatus.DRAFT}
    rows, errors = _confirmation_rows(path)
    for indicator_id, row in rows.items():
        if indicator_id not in by_id:
            errors.append(f"第{row['row']}行 {indicator_id} 不属于当前规则数据")
            continue
        expected_standard = by_id[indicator_id][0].number
        supplied_standard = str(row["standard_number"]).strip() if row["standard_number"] else ""
        if supplied_standard and supplied_standard != expected_standard:
            errors.append(f"第{row['row']}行 {indicator_id} 的标准编号与规则定义不一致")
        if indicator_id not in draft_ids:
            continue
        if row["conclusion"] != "同意发布":
            errors.append(f"第{row['row']}行 {indicator_id} 尚未选择“同意发布”")
        if not row["reviewer"]:
            errors.append(f"第{row['row']}行 {indicator_id} 缺少确认人")
        if not row["confirmed_at"]:
            errors.append(f"第{row['row']}行 {indicator_id} 缺少确认日期")
    confirmed_drafts = {indicator_id for indicator_id in draft_ids if indicator_id in rows and rows[indicator_id]["conclusion"] == "同意发布" and rows[indicator_id]["reviewer"] and rows[indicator_id]["confirmed_at"]}
    missing = sorted(draft_ids - confirmed_drafts)
    if missing:
        errors.append(f"仍有{len(missing)}个draft指标未完成确认")
    if errors:
        raise ValueError("\n".join(errors))

    reviewed_numbers = sorted({definition.number for indicator_id in confirmed_drafts for definition, _ in [by_id[indicator_id]]})
    result = {
        "data_dir": str(data_dir),
        "scope_count": len(scope_numbers),
        "draft_indicator_count": len(draft_ids),
        "confirmed_draft_indicator_count": len(confirmed_drafts),
        "reviewed_standards": reviewed_numbers,
        "applied": apply,
    }
    if not apply:
        return result

    stamp = datetime.now().date().isoformat()
    changed_paths: set[Path] = set()
    for definition, definition_path in zip(definitions, definition_paths, strict=True):
        if definition.publication_status is not PublicationStatus.DRAFT:
            continue
        definition_data = definition.model_dump(mode="json")
        definition_data["publication_status"] = "reviewed"
        for product in definition_data["products"]:
            product["description"] = "已完成确认表逐项确认；待发布标准包生成。"
            for indicator in product["indicators"]:
                indicator["notes"] = [note for note in indicator.get("notes", []) if note not in {"not_for_formal_evaluation", "requires_independent_review"}]
                indicator["notes"].append(f"confirmed_from_workbook_{stamp}")
                for reference in indicator.get("source_references", []):
                    reference["note"] = "已完成确认表逐项确认；待发布标准包生成。"
        StandardDefinition.model_validate(definition_data)
        definition_path.write_text(json.dumps(definition_data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        changed_paths.add(definition_path)

    catalog_path = data_dir / "catalog.json"
    if catalog_path.exists():
        catalog = json.loads(catalog_path.read_text(encoding="utf-8"))
        for item in catalog.get("standards", []):
            if item.get("number") in reviewed_numbers:
                item["status"] = "reviewed"
        catalog_path.write_text(json.dumps(catalog, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    result["changed_definition_count"] = len(changed_paths)
    return result


def main() -> None:
    parser = argparse.ArgumentParser(description="核验确认表并将完整draft标准提升为reviewed（不生成标准包）")
    parser.add_argument("confirmation", type=Path)
    parser.add_argument("--data-dir", type=Path, default=Path("data"))
    parser.add_argument("--scope", type=Path, help="当前范围清单；省略时优先使用data-dir下的scope-63.json或scope-44.json；不会默认使用历史scope-65.json")
    parser.add_argument("--apply", action="store_true", help="验证通过后实际写入reviewed状态")
    args = parser.parse_args()
    try:
        report = promote(args.confirmation.resolve(), args.data_dir.resolve(), apply=args.apply, scope_path=args.scope.resolve() if args.scope else None)
    except (OSError, ValueError) as exc:
        lines = str(exc).splitlines()
        preview = lines[:20]
        if len(lines) > len(preview):
            preview.append(f"……其余{len(lines) - len(preview)}条错误已省略")
        print("确认表校验失败：\n" + "\n".join(preview), file=sys.stderr)
        raise SystemExit(2) from exc
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
