"""选煤共享展示和 Word 出口的快照、精度与历史兼容回归。"""

from __future__ import annotations

import json
import hashlib
from decimal import Decimal, ROUND_DOWN, ROUND_UP, localcontext
from pathlib import Path
from unittest.mock import MagicMock

import pytest
from docx import Document
from docx.document import Document as DocumentClass

from uebench.application.facade import ApplicationFacade
from uebench.application.reporting import (
    ADOPTION_MARKER, MAPPING_ID, MAPPING_SHA256, _display_two_places,
    build_report_view, gb29446_business_labels,
)
from uebench.domain.engine import EvaluationEngine
from uebench.domain.models import (
    EvaluationRequest, EvaluationResult, Grade, InputMode, InputValue, StandardDefinition,
)
from uebench.infrastructure.word_report import WordReportWriter
from uebench.ui.main_window import MainWindow

ROOT = Path(__file__).resolve().parents[1]


def test_report_mapping_hash_matches_approved_utf8_lf_document():
    """采用标记绑定批准正文的 UTF-8 内容哈希，不使用 Git blob 对象哈希。"""
    mapping = ROOT / "standards/mappings" / MAPPING_ID
    source = mapping.read_bytes().decode("utf-8")
    canonical = source.replace("\r\n", "\n").encode("utf-8")
    assert hashlib.sha256(canonical).hexdigest() == MAPPING_SHA256
    assert f"sha256={MAPPING_SHA256}" in ADOPTION_MARKER
    # 工作区换行若转成 CRLF，原始字节哈希会变；按 LF 规范化后仍指向同一正文。
    crlf = source.replace("\r\n", "\n").replace("\n", "\r\n").encode("utf-8")
    assert hashlib.sha256(crlf).hexdigest() != MAPPING_SHA256
    assert hashlib.sha256(crlf.replace(b"\r\n", b"\n")).hexdigest() == MAPPING_SHA256


def _r3() -> StandardDefinition:
    return StandardDefinition.model_validate_json(
        (ROOT / "data/definitions/gb-29446-2019.json").read_text(encoding="utf-8")
    )


def _r2() -> StandardDefinition:
    fixture = json.loads((ROOT / "tests/fixtures/gb29446_record_r2.json").read_text(encoding="utf-8"))
    return StandardDefinition.model_validate_json(fixture["rule_snapshot_json"])


def _detail(standard: StandardDefinition, electricity: str, raw_coal: str, process: str = "重介"):
    request = EvaluationRequest(
        evaluation_date="2026-10-18",
        standard_id=standard.id,
        product_id=standard.products[0].id,
        input_mode=InputMode.DETAIL,
        inputs={
            "washing_process": InputValue(value=process),
            "electricity_consumption": InputValue(value=electricity, unit="kW·h"),
            "raw_coal_input": InputValue(value=raw_coal, unit="t"),
        },
        organization_name="华北地区大型选煤企业有限公司示例项目长名称用于分页检查",
    )
    return request, EvaluationEngine().evaluate(standard, request)


def _direct_r2():
    standard = _r2()
    request = EvaluationRequest(
        evaluation_date="2026-10-18",
        standard_id=standard.id,
        product_id=standard.products[0].id,
        input_mode=InputMode.DIRECT,
        inputs={
            "washing_process": InputValue(value="重介"),
            "actual.coking-coal": InputValue(value="6.272", unit="kW·h/t"),
        },
    )
    return standard, request, EvaluationEngine().evaluate(standard, request)


def test_r3_labels_come_from_definition_and_old_revision_does_not_claim_adoption():
    standard = _r3()
    product = standard.products[0]
    labels = gb29446_business_labels(standard, product)
    definitions = {d.key: d for d in (*product.input_definitions, *product.indicators[0].input_definitions)}
    assert ADOPTION_MARKER in standard.corrections
    assert labels["coal_type"] == standard.selection_schema[0].label == "选煤企业类型"
    assert labels["washing_process"] == definitions["washing_process"].label == "选煤工艺类型"
    assert labels["electricity_consumption"] == definitions["electricity_consumption"].label
    assert labels["raw_coal_input"] == definitions["raw_coal_input"].label
    assert labels["process_factor"] == next(d.label for d in product.indicators[0].display_calculations if d.key == "process_factor")
    r2, request, result = _direct_r2()
    view = build_report_view(request, result, r2)
    assert "采用状态未记录" in view.mapping_adoption
    assert view.electricity is None and view.raw_coal is None
    assert "直接录入" in view.explanation[1]
    assert "×" not in view.explanation[1]
    assert view.enterprise_type.label == "选煤企业类型"
    assert view.period.value == "未记录"


def test_full_value_boundary_display_never_changes_grade():
    standard = _r3()
    request, result = _detail(standard, "50000004", "10000000", "跳汰、浮选联合")
    view = build_report_view(request, result, standard)
    assert result.results[0].actual_value == Decimal("5.0000004")
    assert result.results[0].grade is Grade.LEVEL_2
    assert view.actual.value == "5.00"
    assert view.grade.value == "2级"
    assert "5.0000004" in " ".join(view.explanation)
    with localcontext() as context:
        context.rounding = ROUND_DOWN
        assert build_report_view(request, result, standard).actual.value == "5.00"


def test_display_rounding_is_fixed_at_half_even_independent_of_ambient_context():
    assert _display_two_places(Decimal("6.275")) == "6.28"
    for rounding in (ROUND_DOWN, ROUND_UP):
        with localcontext() as context:
            context.rounding = rounding
            assert _display_two_places(Decimal("6.275")) == "6.28"


def test_ui_and_word_share_warning_terms_without_changing_saved_warning():
    standard = _r3()
    request, result = _detail(standard, "560", "100")
    raw = "electricity_consumption 待核对；GB 29446-2019 第5.2条不变；unrecognized_key 保留"
    result.warnings.append(raw)
    view = build_report_view(request, result, standard)
    shown = MainWindow._gb29446_warning_for_display(object(), request, raw, standard)
    assert view.warnings == (shown,)
    assert "统计期内选煤电力消耗量 E_d 待核对" in shown
    assert "GB 29446-2019 第5.2条不变；unrecognized_key 保留" in shown
    assert result.warnings == [raw]


@pytest.mark.parametrize("grade", [Grade.INCOMPLETE, Grade.NOT_APPLICABLE])
def test_non_final_result_cannot_generate_formal_report(grade):
    standard = _r3()
    request, result = _detail(standard, "560", "100")
    result.results[0].grade = grade
    with pytest.raises(ValueError, match="未形成有效"):
        build_report_view(request, result, standard)


def test_historical_saved_three_part_fixture_exports_without_current_rule_or_engine(tmp_path: Path):
    fixture_path = ROOT / "tests/fixtures/gb29446_record_r2.json"
    original = fixture_path.read_bytes()
    fixture = json.loads(original)
    request = EvaluationRequest.model_validate_json(fixture["request_json"])
    result = EvaluationResult.model_validate_json(fixture["result_json"])
    snapshot = StandardDefinition.model_validate_json(fixture["rule_snapshot_json"])
    saved_models = tuple(model.model_dump_json() for model in (request, result, snapshot))
    repo = MagicMock()
    repo.get.return_value = (request, result, snapshot)
    standards = MagicMock()
    standards.get_current.side_effect = AssertionError("不得查询当前标准")
    engine = MagicMock()
    engine.evaluate.side_effect = AssertionError("不得重新计算")
    facade = ApplicationFacade(
        standards=standards, evaluations=repo, evaluation_service=engine,
        word_report_writer=WordReportWriter(),
    )
    target = tmp_path / "历史记录.docx"
    view = facade.get_report_view(result.evaluation_id)
    assert "采用状态未记录" in view.mapping_adoption
    assert view.actual_full_value == "6.272"
    assert facade.export_word_report(result.evaluation_id, target) == target
    assert target.exists() and target.stat().st_size > 0
    assert fixture_path.read_bytes() == original
    assert tuple(model.model_dump_json() for model in (request, result, snapshot)) == saved_models
    assert standards.mock_calls == []
    assert repo.save.call_count == 0
    engine.evaluate.assert_not_called()


@pytest.mark.parametrize("mutation", ["missing_electricity", "bad_unit", "wrong_revision", "wrong_indicator", "missing_threshold", "false_adoption"])
def test_report_rejects_bad_core_facts_and_does_not_infer_adoption(mutation: str):
    standard = _r3()
    request, result = _detail(standard, "560", "100")
    if mutation == "missing_electricity":
        request.inputs.pop("electricity_consumption")
    elif mutation == "bad_unit":
        request.inputs["electricity_consumption"].unit = "MWh"
    elif mutation == "wrong_revision":
        result.rule_revision += 1
    elif mutation == "wrong_indicator":
        result.results[0].indicator_id = "other"
    elif mutation == "missing_threshold":
        result.results[0].corrected_thresholds.pop("LEVEL_3")
    elif mutation == "false_adoption":
        standard.corrections.clear()
        # 标记删除后需与其新快照身份一致，模拟另一条同样叫 r3 的历史规则。
        from uebench.application.reporting import _saved_snapshot_hash
        result.rule_snapshot_sha256 = _saved_snapshot_hash(standard)
        assert "采用状态未记录" in build_report_view(request, result, standard).mapping_adoption
        return
    with pytest.raises(ValueError):
        build_report_view(request, result, standard)


def test_writer_uses_view_and_preserves_existing_file_on_failure(tmp_path: Path, monkeypatch):
    standard = _r3()
    request, result = _detail(standard, "560", "100")
    result.warnings.append("结果提醒")
    result.results[0].warnings.append("指标提醒")
    view = build_report_view(request, result, standard)
    assert view.warnings == ("结果提醒", "指标提醒")
    target = tmp_path / "中文报告.docx"
    writer = WordReportWriter()
    writer.write(view, target)
    document = Document(target)
    text = "\n".join(p.text for p in document.paragraphs) + "\n" + "\n".join(c.text for t in document.tables for row in t.rows for c in row.cells)
    assert all(term in text for term in ("机修", "等级或评价结果", "6.272", *view.warnings))
    before = target.read_bytes()
    monkeypatch.setattr(DocumentClass, "save", lambda self, path: (_ for _ in ()).throw(RuntimeError("write failed")))
    with pytest.raises(RuntimeError, match="write failed"):
        writer.write(view, target)
    assert target.read_bytes() == before


def generate_samples(directory: Path) -> list[Path]:
    """生成可视化验收样例，不由测试运行时自动写入仓库。"""
    directory.mkdir(parents=True, exist_ok=True)
    writer = WordReportWriter()
    standard = _r3()
    cases = [
        ("正常_长中文名称", *_detail(standard, "560", "100"), standard),
        ("临界_完整值5点0000004", *_detail(standard, "50000004", "10000000", "跳汰、浮选联合"), standard),
        ("未达3级", *_detail(standard, "1120", "100"), standard),
    ]
    old_standard, old_request, old_result = _direct_r2()
    cases.append(("历史DIRECT", old_request, old_result, old_standard))
    outputs: list[Path] = []
    for name, request, result, snapshot in cases:
        view = build_report_view(request, result, snapshot)
        outputs.append(writer.write(view, directory / f"{name}.docx"))
    return outputs
