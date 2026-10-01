from __future__ import annotations

import json
import zipfile
from datetime import date
from decimal import Decimal, ROUND_UP, getcontext
from pathlib import Path

import pytest
from jsonschema import Draft202012Validator
from openpyxl import Workbook, load_workbook
from pydantic import ValidationError

from uebench.domain.engine import EvaluationEngine, ExpressionEvaluator
from uebench.domain.models import (
    ComplianceCase,
    ComplianceRule,
    ComparisonDirection,
    EvaluationRequest,
    Expression,
    Grade,
    IndicatorDefinition,
    InputDefinition,
    InputMode,
    InputValue,
    ProductDefinition,
    PublicationStatus,
    SourceReference,
    StandardDefinition,
    ThresholdSet,
    parse_decimal,
)
from uebench.domain.numeric import (
    ECQUOTA_CALCULATOR_VERSION,
    ECQUOTA_DECIMAL_FULL_VALUE_V1,
    GB29446_NUMERIC_BEHAVIOR_VERSION,
)
from uebench.infrastructure.excel import ImportIssue, _decimal_from_cell


ROOT = Path(__file__).resolve().parents[3]
SCHEMA_PATH = Path(__file__).with_name("conformance_vector_v1.schema.json")
VECTORS_PATH = Path(__file__).with_name("ecquota_numeric_v1_vectors.json")
GB29446_PATH = ROOT / "data" / "definitions" / "gb-29446-2019.json"
VECTORS = json.loads(VECTORS_PATH.read_text(encoding="utf-8"))
SCHEMA = json.loads(SCHEMA_PATH.read_text(encoding="utf-8"))
GB29446 = StandardDefinition.model_validate_json(GB29446_PATH.read_text(encoding="utf-8"))
PROCESSES = {"炼焦煤": "跳汰、浮选联合", "动力煤": "干法选煤"}


SOURCE = SourceReference(
    standard_number="GB 00000-2026",
    source_file="GB 00000-2026.pdf",
    source_sha256="a" * 64,
    page=5,
    table="表1",
)


def _constant(value: str) -> Expression:
    return Expression(op="constant", value=value)


def _generic_standard() -> StandardDefinition:
    indicator = IndicatorDefinition(
        id="energy",
        name="单位产品能耗",
        unit="kgce/t",
        comparison=ComparisonDirection.LTE,
        direct_input_key="actual",
        input_definitions=[
            InputDefinition(key="actual", label="实际值", unit="kgce/t", modes=[InputMode.DIRECT]),
            InputDefinition(
                key="enterprise_status",
                label="企业状态",
                data_type="text",
                choices=["现有"],
                required=False,
                modes=[InputMode.DIRECT],
            ),
        ],
        thresholds=ThresholdSet(
            level_1=_constant("10"),
            level_2=_constant("20"),
            level_3=_constant("30"),
        ),
        compliance_rule=ComplianceRule(
            input_key="enterprise_status",
            cases=[ComplianceCase(value="现有", requirement="不高于2级限值", threshold_key="LEVEL_2")],
        ),
        source_references=[SOURCE],
    )
    return StandardDefinition(
        id="gb-00000-2026",
        number="GB 00000-2026",
        title="Numeric v1 合成边界标准",
        version="2026",
        publication_status=PublicationStatus.PUBLISHED,
        publication_date=date(2026, 1, 1),
        effective_date=date(2026, 2, 1),
        source_file=SOURCE.source_file,
        source_sha256=SOURCE.source_sha256,
        products=[ProductDefinition(id="product", name="测试产品", indicators=[indicator])],
    )


def _generic_request(actual: str, *, status: str | None = None) -> EvaluationRequest:
    inputs = {"actual": InputValue(value=actual, unit="kgce/t")}
    if status is not None:
        inputs["enterprise_status"] = InputValue(value=status)
    return EvaluationRequest(
        evaluation_date=date(2026, 2, 2),
        standard_id="gb-00000-2026",
        product_id="product",
        input_mode=InputMode.DIRECT,
        inputs=inputs,
    )


def _gb_product(coal_type: str):
    return next(item for item in GB29446.products if item.selection_values["coal_type"] == coal_type)


def _gb_direct(coal_type: str, actual: str) -> EvaluationRequest:
    product = _gb_product(coal_type)
    return EvaluationRequest(
        evaluation_date=max(GB29446.effective_date, date(2026, 10, 1)),
        standard_id=GB29446.id,
        product_id=product.id,
        input_mode=InputMode.DIRECT,
        inputs={
            "washing_process": InputValue(value=PROCESSES[coal_type]),
            product.indicators[0].direct_input_key: InputValue(value=actual, unit="kW·h/t"),
        },
    )


def _gb_detail(electricity: str, raw_coal: str = "10") -> EvaluationRequest:
    product = _gb_product("炼焦煤")
    return EvaluationRequest(
        evaluation_date=max(GB29446.effective_date, date(2026, 10, 1)),
        standard_id=GB29446.id,
        product_id=product.id,
        input_mode=InputMode.DETAIL,
        inputs={
            "washing_process": InputValue(value="跳汰、浮选联合"),
            "electricity_consumption": InputValue(value=electricity, unit="kW·h"),
            "raw_coal_input": InputValue(value=raw_coal, unit="t"),
        },
    )


def _coal_for_case(case_id: str) -> str:
    return "动力煤" if case_id.endswith(("-2", "-3", "-4_5")) else "炼焦煤"


def _execute_vector(vector: dict) -> None:
    case_id = vector["case_id"]
    if case_id == "ecq-parse-trailing-zero":
        value = parse_decimal(vector["inputs"]["value"], field_name="vector")
        assert str(value) == vector["expected_reference_value"]
        return
    if case_id == "ecq-reject-python-float":
        with pytest.raises((TypeError, ValidationError)):
            InputValue(value=vector["inputs"]["python_float"])
        return
    if case_id == "ecq-public-grade-round6-break":
        result = EvaluationEngine().evaluate(
            _generic_standard(), _generic_request(vector["inputs"]["actual"])
        )
        assert result.results[0].grade.value == vector["expected_business_result"]
        assert result.results[0].actual_value == Decimal(vector["comparison_value"])
        return
    if case_id == "ecq-public-compliance-round6-break":
        result = EvaluationEngine().evaluate(
            _generic_standard(), _generic_request(vector["inputs"]["actual"], status="现有")
        )
        assert result.results[0].compliance_result == vector["expected_business_result"]
        return
    if case_id.startswith("ecq-gb29446-t-"):
        result = EvaluationEngine().evaluate(GB29446, _gb_direct("炼焦煤", vector["inputs"]["actual"]))
        assert result.results[0].grade.value == vector["expected_business_result"]
        return
    if case_id.startswith("ecq-gb29446-break-"):
        result = EvaluationEngine().evaluate(
            GB29446, _gb_direct(_coal_for_case(case_id), vector["inputs"]["actual"])
        )
        assert result.results[0].grade.value == vector["expected_business_result"]
        comparisons = [s for s in result.results[0].calculation_trace if s.operation == "grade_comparison"]
        assert comparisons
        assert all("ROUND(" not in s.expression for s in comparisons)
        assert all(GB29446_NUMERIC_BEHAVIOR_VERSION in s.expression for s in comparisons)
        return
    if case_id == "ecq-display-separation":
        result = EvaluationEngine().evaluate(GB29446, _gb_detail(vector["inputs"]["E_d"]))
        item = result.results[0]
        assert item.actual_value == Decimal(vector["calculation_value"])
        assert item.actual_value == Decimal(vector["comparison_value"])
        assert item.grade.value == vector["expected_business_result"]
        assert format(item.actual_value.quantize(Decimal("0.01")), "f") == vector["display_value"]
        return
    if case_id == "ecq-profile-propagation":
        result = EvaluationEngine().evaluate(GB29446, _gb_direct("炼焦煤", vector["inputs"]["actual"]))
        assert result.numeric_contract_version == vector["expected_business_result"]["numeric_contract_version"]
        assert result.numeric_profile_id == vector["expected_business_result"]["numeric_profile_id"]
        assert result.calculator_version == ECQUOTA_CALCULATOR_VERSION
        assert result.rule_revision == GB29446.rule_revision
        return
    if case_id == "ecq-ambient-independence":
        context = getcontext()
        original = context.copy()
        try:
            context.prec = int(vector["ambient_profile"]["working_precision"])
            context.rounding = ROUND_UP
            result = EvaluationEngine().evaluate(
                GB29446,
                _gb_detail(vector["inputs"]["E_d"], raw_coal=vector["inputs"]["m"]),
            )
        finally:
            context.prec = original.prec
            context.rounding = original.rounding
            context.Emin = original.Emin
            context.Emax = original.Emax
            context.capitals = original.capitals
            context.clamp = original.clamp
        assert str(result.results[0].actual_value) == vector["expected_reference_value"]
        assert result.numeric_profile_id == vector["effective_profile"]
        return
    raise AssertionError(f"Frozen vector has no executor: {case_id}")


def test_frozen_vector_schema_is_exactly_executable() -> None:
    validator = Draft202012Validator(SCHEMA)
    errors = []
    for vector in VECTORS:
        errors.extend(f"{vector['case_id']}: {item.message}" for item in validator.iter_errors(vector))
    assert not errors, "\n".join(errors)


@pytest.mark.parametrize("vector", VECTORS, ids=lambda item: item["case_id"])
def test_all_frozen_numeric_v1_vectors_are_actually_executed(vector: dict) -> None:
    _execute_vector(vector)


def test_project_profile_declares_all_frozen_required_capabilities() -> None:
    profile = ECQUOTA_DECIMAL_FULL_VALUE_V1
    assert profile.numeric_profile_id == "ECQUOTA_DECIMAL_FULL_VALUE_V1"
    assert profile.numeric_contract_version == "v1"
    assert profile.representation == "Decimal"
    assert profile.working_precision == 28
    assert profile.rounding_mode == "ROUND_HALF_EVEN"
    assert profile.comparison_policy == "full-value exact"
    assert profile.explicit_rounding_policy
    assert profile.transcendental_policy == "not_applicable"
    assert profile.tolerance_policy == "none by default; no global epsilon"


def test_explicit_rule_rounding_requires_authority_metadata() -> None:
    with pytest.raises(ValidationError, match="round_places 必须声明"):
        Expression(op="constant", value="1.005", round_places=2)

    expression = Expression(
        op="constant",
        value="1.005",
        round_places=2,
        rounding={
            "stage": "intermediate",
            "mode": "ROUND_HALF_UP",
            "purpose": "business-explicit",
            "source": "synthetic conformance authority",
        },
    )
    assert ExpressionEvaluator().evaluate(expression, object(), []) == Decimal("1.01")


def test_xlsx_numeric_cell_can_cross_exact_boundary_and_guard_rejects_it(tmp_path: Path) -> None:
    """Prove the old openpyxl float route can change a formal exact-boundary result."""
    path = tmp_path / "binary-float-boundary.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = 5
    workbook.save(path)

    # Patch the XLSX numeric XML with a decimal that binary64 cannot preserve.
    patched = tmp_path / "patched.xlsx"
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(patched, "w") as target:
        for member in source.infolist():
            payload = source.read(member.filename)
            if member.filename == "xl/worksheets/sheet1.xml":
                xml = payload.decode("utf-8")
                xml = xml.replace("<v>5</v>", "<v>5.0000000000000001</v>")
                payload = xml.encode("utf-8")
            target.writestr(member, payload)

    loaded_value = load_workbook(patched, data_only=True).active["A1"].value
    assert isinstance(loaded_value, float)
    assert loaded_value == 5.0

    exact = EvaluationEngine().evaluate(GB29446, _gb_direct("炼焦煤", "5.0000000000000001"))
    materialized = EvaluationEngine().evaluate(GB29446, _gb_direct("炼焦煤", str(loaded_value)))
    assert exact.results[0].grade is Grade.LEVEL_2
    assert materialized.results[0].grade is Grade.LEVEL_1

    issues: list[ImportIssue] = []
    assert _decimal_from_cell(loaded_value, sheet="实际值", cell="C2", issues=issues) is None
    assert issues and "二进制浮点" in issues[0].message
