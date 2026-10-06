from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from uebench.application import evaluation_support
from uebench.application.services import EvaluationService
from uebench.domain.models import EvaluationRequest, InputMode, InputValue
from uebench.infrastructure.database import DatabaseManager
from uebench.infrastructure.excel import (
    EXPORT_SHEETS,
    IMPORT_SHEETS,
    WorkbookExportService,
    WorkbookImportService,
    WorkbookTemplateService,
)
from uebench.infrastructure.paths import AppPaths
from uebench.infrastructure.repositories import AuditRepository, SqlEvaluationRepository, SqlStandardRepository

from .test_engine import make_standard


def _in_formal_scope(monkeypatch, *standard_ids: str) -> None:
    """把 ``standard_ids`` 声明为正式可评价（与 ``test_ui.py::_in_formal_scope`` 同模式）。

    正式评价范围是应用层的固定常量，**与“标准库里有没有这个标准”无关**（RS05 §三）。
    本文件覆盖的是模板/导入/导出，不是范围本身，因此需要正式评价落库的用例必须显式扩展
    真正的注册表；产品行为仍然完全由注册表驱动。
    """
    extended = set(evaluation_support.SUPPORTED_EVALUATION_STANDARD_IDS) | set(standard_ids)
    monkeypatch.setattr(
        evaluation_support, "SUPPORTED_EVALUATION_STANDARD_IDS", frozenset(extended)
    )


def setup_services(tmp_path: Path):
    paths = AppPaths.from_root(tmp_path / "appdata")
    paths.ensure()
    database = DatabaseManager(paths.database)
    database.initialize()
    audit = AuditRepository(database)
    standards = SqlStandardRepository(database, audit)
    evaluations = SqlEvaluationRepository(database, audit)
    standards.install(make_standard())
    return database, audit, standards, evaluations


def test_template_and_import_round_trip(tmp_path: Path) -> None:
    database, audit, _, _ = setup_services(tmp_path)
    template = WorkbookTemplateService().create_template(tmp_path / "import.xlsx")
    workbook = load_workbook(template)
    assert workbook.sheetnames == IMPORT_SHEETS
    info = workbook["评价信息"]
    info["B3"] = date(2026, 2, 2)
    info["B4"] = "gb-00000-2026"
    info["B5"] = "product"
    info["B6"] = "DIRECT"
    actual = workbook["实际值"]
    actual.append(["actual", "单位产品能耗", "20", "kgce/t", "检测报告"])
    workbook.save(template)

    service = WorkbookImportService(database, audit)
    report = service.validate(template)
    assert report.valid, report.issues
    draft = service.commit(report.import_id)
    assert draft.request.inputs["actual"].value == "20"


def test_import_rejects_renamed_sheet(tmp_path: Path) -> None:
    database, audit, _, _ = setup_services(tmp_path)
    template = WorkbookTemplateService().create_template(tmp_path / "bad.xlsx")
    workbook = load_workbook(template)
    workbook["实际值"].title = "实际数据"
    workbook.save(template)
    report = WorkbookImportService(database, audit).validate(template)
    assert not report.valid
    assert any("工作表必须严格为" in issue.message for issue in report.issues)


def test_import_rejects_custom_column_header(tmp_path: Path) -> None:
    database, audit, _, _ = setup_services(tmp_path)
    template = WorkbookTemplateService().create_template(tmp_path / "bad-column.xlsx")
    workbook = load_workbook(template)
    workbook["实际值"]["A1"] = "自定义键"
    workbook.save(template)
    report = WorkbookImportService(database, audit).validate(template)
    assert not report.valid
    assert any(issue.sheet == "实际值" and "表头必须严格为" in issue.message for issue in report.issues)


def test_import_rejects_duplicate_input_keys(tmp_path: Path) -> None:
    database, audit, _, _ = setup_services(tmp_path)
    template = WorkbookTemplateService().create_template(tmp_path / "duplicate-input.xlsx")
    workbook = load_workbook(template)
    info = workbook["评价信息"]
    info["B3"] = date(2026, 2, 2)
    info["B4"] = "gb-00000-2026"
    info["B5"] = "product"
    info["B6"] = "DIRECT"
    actual = workbook["实际值"]
    actual.append(["actual", "单位产品能耗", "20", "kgce/t", ""])
    actual.append(["actual", "单位产品能耗", "20", "kgce/t", ""])
    workbook.save(template)
    report = WorkbookImportService(database, audit).validate(template)
    assert not report.valid
    assert any("输入键重复" in issue.message for issue in report.issues)


def test_import_reports_unknown_standard_and_product_at_info_cells(tmp_path: Path) -> None:
    database, audit, standards, _ = setup_services(tmp_path)
    template = WorkbookTemplateService().create_template(tmp_path / "unknown.xlsx")
    workbook = load_workbook(template)
    info = workbook["评价信息"]
    info["B3"] = date(2026, 2, 2)
    info["B4"] = "gb-unknown-2026"
    info["B5"] = "missing-product"
    info["B6"] = "DIRECT"
    workbook.save(template)
    report = WorkbookImportService(database, audit, standards).validate(template)
    assert not report.valid
    assert any(issue.cell == "B4" and "未找到已发布标准" in issue.message for issue in report.issues)


def test_export_matches_saved_result(tmp_path: Path, monkeypatch) -> None:
    _, audit, standards, evaluations = setup_services(tmp_path)
    # 导出需要一条**正式**记录，因此把夹具标准显式放进注册表（RS05 §三）。
    _in_formal_scope(monkeypatch, "gb-00000-2026")
    request = EvaluationRequest(
        evaluation_date=date(2026, 2, 2),
        standard_id="gb-00000-2026",
        product_id="product",
        input_mode=InputMode.DIRECT,
        inputs={"actual": InputValue(value="20", unit="kgce/t")},
        organization_name="测试单位",
    )
    result = EvaluationService(standards, evaluations).evaluate(request)
    output = WorkbookExportService(evaluations, audit).export(result.evaluation_id, tmp_path / "result.xlsx")
    workbook = load_workbook(output, data_only=False)
    assert workbook.sheetnames == EXPORT_SHEETS
    summary = workbook["对标结论"]
    assert summary["A6"].value == "单位产品能耗"
    assert summary["G6"].value == "2级"
    assert workbook["标准依据"]["E4"].value == 5
