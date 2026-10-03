"""RS03 gate — GB29446 Excel Adapter closure.

Excel is only an adapter: a workbook is validated, converted into the same
canonical ``EvaluationRequest`` the GUI produces, and evaluated by the same
``EvaluationService`` / ``EvaluationEngine``.  Excel never computes ``k``,
``e_d`` or a grade.

Business mathematics (thresholds, boundary parity) stays with the RS01 and
Numeric gates; this module proves the adapter contract, the ingress rules, the
record lifecycle and the GUI/Excel equivalence.
"""
from __future__ import annotations

import hashlib
import json
import os
import re
import subprocess
import sys
import zipfile
from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import text

from uebench.application import gb29446 as gb
from uebench.bootstrap import create_context
from uebench.domain.models import (
    EvaluationRequest,
    Grade,
    InputMode,
    InputValue,
    StandardSelectionMode,
)
from uebench.infrastructure.excel import (
    GB29446_DATA_SHEET,
    GB29446_INSTRUCTIONS_SHEET,
    GB29446_META_SHEET,
    GENERIC_PROFILE_ID,
    GB29446_PROFILE_ID,
    STATUS_EVALUATED,
    WorkbookImportService,
    WorkbookTemplateService,
)

ROOT = Path(__file__).resolve().parents[1]
STANDARD_ID = "gb-29446-2019"

DATA_ROWS = {
    gb.FIELD_ORGANIZATION_NAME: 2,
    gb.FIELD_EVALUATION_DATE: 3,
    gb.FIELD_PERIOD: 4,
    gb.FIELD_CUSTOM_PERIOD: 5,
    gb.FIELD_COAL_TYPE: 6,
    gb.FIELD_WASHING_PROCESS: 7,
    gb.FIELD_ELECTRICITY: 8,
    gb.FIELD_RAW_COAL: 9,
    gb.FIELD_NOTES: 10,
}
META_ROWS = {"adapter_id": 2, "template_version": 3, "standard_id": 4, "standard_version": 5, "rule_revision": 6}

EVALUATION_DATE = date(2026, 6, 1)


@pytest.fixture
def ctx(tmp_path):
    context = create_context(tmp_path / "中文数据")
    standard = json.loads((ROOT / "data/definitions" / "gb-29446-2019.json").read_text(encoding="utf-8"))
    from uebench.domain.models import StandardDefinition

    context.standards.install(StandardDefinition.model_validate(standard))
    try:
        yield context
    finally:
        context.database.dispose()


def make_template(ctx, tmp_path, name="rs03.xlsx") -> Path:
    return ctx.application.create_template(tmp_path / name, STANDARD_ID)


def fill(
    path: Path,
    *,
    electricity="560",
    raw_coal="100",
    coal="炼焦煤",
    process="重介",
    period=gb.PERIOD_FULL_YEAR,
    custom_period=None,
    organization="宁夏测试企业",
    note="同一内容",
    evaluation_date=EVALUATION_DATE,
    electricity_raw=None,
) -> Path:
    """Fill the GB29446 data sheet.  ``electricity_raw`` forces an exact cell value/type."""
    workbook = load_workbook(path)
    sheet = workbook[GB29446_DATA_SHEET]
    sheet.cell(row=DATA_ROWS[gb.FIELD_ORGANIZATION_NAME], column=2, value=organization)
    sheet.cell(row=DATA_ROWS[gb.FIELD_EVALUATION_DATE], column=2, value=evaluation_date)
    sheet.cell(row=DATA_ROWS[gb.FIELD_PERIOD], column=2, value=period)
    sheet.cell(row=DATA_ROWS[gb.FIELD_CUSTOM_PERIOD], column=2, value=custom_period)
    sheet.cell(row=DATA_ROWS[gb.FIELD_COAL_TYPE], column=2, value=coal)
    sheet.cell(row=DATA_ROWS[gb.FIELD_WASHING_PROCESS], column=2, value=process)
    cell = sheet.cell(row=DATA_ROWS[gb.FIELD_ELECTRICITY], column=2)
    cell.value = electricity if electricity_raw is None else electricity_raw
    if isinstance(cell.value, str) or cell.value is None:
        cell.number_format = "@"
    else:
        cell.number_format = "General"
    raw = sheet.cell(row=DATA_ROWS[gb.FIELD_RAW_COAL], column=2)
    raw.value = raw_coal
    raw.number_format = "@"
    sheet.cell(row=DATA_ROWS[gb.FIELD_NOTES], column=2, value=note)
    workbook.save(path)
    return path


def validate(ctx, path: Path):
    return ctx.application.validate_workbook(path)


def gui_request(standard, *, electricity="560", raw_coal="100", coal="炼焦煤", process="重介",
                period=gb.PERIOD_FULL_YEAR, custom_period="", note="同一内容",
                organization="宁夏测试企业") -> EvaluationRequest:
    """Build the same request the desktop form builds, using the shared mapping."""
    product = next(p for p in standard.products if p.selection_values["coal_type"] == coal)
    inputs = {gb.INPUT_WASHING_PROCESS: InputValue(value=process)}
    if electricity:
        inputs[gb.INPUT_ELECTRICITY] = InputValue(value=electricity, unit=gb.UNIT_ELECTRICITY)
    if raw_coal:
        inputs[gb.INPUT_RAW_COAL] = InputValue(value=raw_coal, unit=gb.UNIT_RAW_COAL)
    return EvaluationRequest(
        evaluation_date=EVALUATION_DATE,
        standard_id=standard.id,
        product_id=product.id,
        selection_mode=StandardSelectionMode.CURRENT,
        input_mode=InputMode.DETAIL,
        inputs=inputs,
        organization_name=organization,
        notes=gb.encode_period_notes(period, custom_period, note),
    )


def boundary_parity(ctx, path):
    """Run the same lexical boundary text through both adapters.

    The coefficient for 跳汰、浮选联合 is 1.00, so e_d = E_d × k ÷ m equals the
    lexical input exactly and the boundary value reaches the comparison intact.
    """
    standard = ctx.application.get_published_standard(STANDARD_ID)
    boundary = "5.0000004"
    report = validate(
        ctx,
        fill(
            path,
            electricity=boundary,
            raw_coal="1",
            process="跳汰、浮选联合",
        ),
    )
    assert report.valid, report.issues
    excel_result = ctx.application.evaluate_workbook(report.import_id)
    gui_result = ctx.application.evaluate(
        gui_request(standard, electricity=boundary, raw_coal="1", process="跳汰、浮选联合")
    )
    return excel_result, gui_result


# ---------------------------------------------------------------------------
# 1-3, 5: template generation, no internal ids, metadata, generic regression
# ---------------------------------------------------------------------------


def test_gb29446_template_generation(ctx, tmp_path):
    path = make_template(ctx, tmp_path)
    workbook = load_workbook(path)
    assert workbook.sheetnames == [
        GB29446_INSTRUCTIONS_SHEET,
        GB29446_DATA_SHEET,
        GB29446_META_SHEET,
    ], "GB29446 模板必须只有 填写说明 / 评价数据 / _meta 三张表"
    assert workbook[GB29446_META_SHEET].sheet_state in {"hidden", "veryHidden"}
    labels = [workbook[GB29446_DATA_SHEET].cell(row=row, column=1).value for row in range(2, 11)]
    assert labels == [
        gb.FIELD_ORGANIZATION_NAME,
        gb.FIELD_EVALUATION_DATE,
        gb.FIELD_PERIOD,
        gb.FIELD_CUSTOM_PERIOD,
        gb.FIELD_COAL_TYPE,
        gb.FIELD_WASHING_PROCESS,
        gb.FIELD_ELECTRICITY,
        gb.FIELD_RAW_COAL,
        gb.FIELD_NOTES,
    ]


def test_template_user_sheet_has_no_internal_identifiers(ctx, tmp_path):
    path = make_template(ctx, tmp_path)
    workbook = load_workbook(path)
    text_dump = []
    for sheet in workbook.worksheets:
        if sheet.title == GB29446_META_SHEET:
            continue
        for row in sheet.iter_rows(values_only=True):
            text_dump.extend(str(value) for value in row if value is not None)
    joined = "\n".join(text_dump)
    for forbidden in (
        "standard_id", "product_id", "input_mode", "rule_revision", "adapter_id",
        "template_version", "washing_process", "electricity_consumption", "raw_coal_input",
        "GB_29446-2019.coking-coal", "gb_29446-2019-coking-coal", "numeric_profile",
        "calculator", "round_places", "selection_mode",
    ):
        assert forbidden not in joined, f"普通用户表不得出现内部标识：{forbidden}"


def test_data_sheet_exposes_inputs_only_never_results(ctx, tmp_path):
    """The user sheet carries inputs; k / e_d / grade are produced by the engine."""
    path = make_template(ctx, tmp_path)
    sheet = load_workbook(path)[GB29446_DATA_SHEET]
    fields = [sheet.cell(row=row, column=1).value for row in range(2, sheet.max_row + 1)]
    fields = [value for value in fields if value]
    expected = [
        gb.FIELD_ORGANIZATION_NAME,
        gb.FIELD_EVALUATION_DATE,
        gb.FIELD_PERIOD,
        gb.FIELD_CUSTOM_PERIOD,
        gb.FIELD_COAL_TYPE,
        gb.FIELD_WASHING_PROCESS,
        gb.FIELD_ELECTRICITY,
        gb.FIELD_RAW_COAL,
        gb.FIELD_NOTES,
    ]
    assert fields == expected, "评价数据表只能包含规定输入字段"
    for forbidden in ("折算系数", "k", "e_d", "等级", "判定", "限值", "结果"):
        assert forbidden not in fields, f"评价数据表不得提供业务结果字段：{forbidden}"


def test_template_metadata_matches_definition(ctx, tmp_path):
    path = make_template(ctx, tmp_path)
    meta = load_workbook(path)[GB29446_META_SHEET]
    values = {meta.cell(row=r, column=1).value: meta.cell(row=r, column=2).value for r in META_ROWS.values()}
    standard = ctx.application.get_published_standard(STANDARD_ID)
    assert values["adapter_id"] == gb.GB29446_EXCEL_ADAPTER_ID
    assert values["template_version"] == gb.GB29446_EXCEL_TEMPLATE_VERSION
    assert values["standard_id"] == standard.id == STANDARD_ID
    assert str(values["standard_version"]) == standard.version
    assert int(values["rule_revision"]) == standard.rule_revision


def test_template_dropdowns_cover_declared_options(ctx, tmp_path):
    path = make_template(ctx, tmp_path)
    sheet = load_workbook(path)[GB29446_DATA_SHEET]
    validations = " ".join(dv.formula1 for dv in sheet.data_validations.dataValidation if dv.formula1)
    for option in gb.GB29446_PERIOD_OPTIONS:
        assert option in validations, f"核算周期下拉缺少 {option}"
    assert "炼焦煤" in validations and "动力煤" in validations
    assert "重介" in validations


def test_generic_template_regression(ctx, tmp_path):
    """The generic multi-standard template must keep working unchanged."""
    path = ctx.application.create_template(tmp_path / "generic.xlsx")
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["填写说明", "评价信息", "适用条件", "实际值", "能源明细", "产量与分摊"]
    info = workbook["评价信息"]
    info["B3"] = EVALUATION_DATE
    info["B4"] = STANDARD_ID
    info["B5"] = "gb_29446-2019-coking-coal"
    info["B6"] = "DIRECT"
    workbook["适用条件"].append(["washing_process", "选煤工艺", "重介", "", ""])
    workbook["适用条件"].append(["single_coal_single_process", "单一煤种单一工艺", "是", "", ""])
    workbook["实际值"].append(["actual.coking-coal", "单耗", "5.0", "kW·h/t", ""])
    workbook.save(path)

    report = validate(ctx, path)
    assert report.valid, report.issues
    assert report.profile_id == GENERIC_PROFILE_ID
    draft = ctx.application.commit_workbook(report.import_id)
    assert draft.request.standard_id == STANDARD_ID
    assert draft.request.product_id == "gb_29446-2019-coking-coal"


# ---------------------------------------------------------------------------
# 6-13, 27: ingress matrix
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("text", ["560", "560.25", "1e3"])
def test_text_decimal_lexical_accepted(ctx, tmp_path, text):
    report = validate(ctx, fill(make_template(ctx, tmp_path, f"ok-{abs(hash(text))}.xlsx"), electricity=text))
    assert report.valid, report.issues
    assert report.request.inputs[gb.INPUT_ELECTRICITY].value == str(Decimal(text))


def test_lexical_value_survives_without_float_round_trip(ctx, tmp_path):
    report = validate(ctx, fill(make_template(ctx, tmp_path, "lex.xlsx"), electricity="5.0000004"))
    assert report.valid, report.issues
    assert report.request.inputs[gb.INPUT_ELECTRICITY].value == "5.0000004"


@pytest.mark.parametrize("raw,numeric", [(560, True), (560.25, False)])
def test_xlsx_numeric_cell_rejected(ctx, tmp_path, raw, numeric):
    report = validate(ctx, fill(make_template(ctx, tmp_path, f"n-{raw}.xlsx"), electricity_raw=raw))
    assert not report.valid
    assert any(issue.cell == "B8" for issue in report.issues), [i.message for i in report.issues]


def test_formula_cell_rejected_without_cache(ctx, tmp_path):
    path = fill(make_template(ctx, tmp_path, "f1.xlsx"), electricity_raw=None)
    workbook = load_workbook(path)
    workbook[GB29446_DATA_SHEET]["B8"] = "=500+60"
    workbook.save(path)
    report = validate(ctx, path)
    assert not report.valid
    assert any("公式" in issue.message for issue in report.issues)


def test_formula_cell_rejected_with_cached_value(ctx, tmp_path):
    path = fill(make_template(ctx, tmp_path, "f2.xlsx"), electricity_raw=None)
    workbook = load_workbook(path)
    workbook[GB29446_DATA_SHEET]["B8"] = "=500+60"
    workbook.save(path)
    _inject_cached_formula(path, "500+60", "560")
    assert load_workbook(path, data_only=False)[GB29446_DATA_SHEET]["B8"].data_type == "f"
    assert load_workbook(path, data_only=True)[GB29446_DATA_SHEET]["B8"].value == 560
    report = validate(ctx, path)
    assert not report.valid, "缓存值存在也不得接受公式"
    assert any("公式" in issue.message for issue in report.issues)


@pytest.mark.parametrize("text", ["NaN", "sNaN", "Infinity", "-Infinity", "inf", "-inf"])
def test_non_finite_rejected(ctx, tmp_path, text):
    report = validate(ctx, fill(make_template(ctx, tmp_path, f"nf-{abs(hash(text))}.xlsx"), electricity=text))
    assert not report.valid
    assert any("有限" in issue.message for issue in report.issues)


@pytest.mark.parametrize("text", ["1,000", "560 kWh", "abc", ""])
def test_malformed_or_blank_numeric_text_rejected(ctx, tmp_path, text):
    report = validate(ctx, fill(make_template(ctx, tmp_path, f"bad-{abs(hash(text))}.xlsx"), electricity_raw=text))
    assert not report.valid


def test_boolean_cell_rejected(ctx, tmp_path):
    report = validate(ctx, fill(make_template(ctx, tmp_path, "bool.xlsx"), electricity_raw=True))
    assert not report.valid


# ---------------------------------------------------------------------------
# 14-15, 25: mapping, period, selection mode
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "coal,process,product_id",
    [
        ("炼焦煤", "重介", "gb_29446-2019-coking-coal"),
        ("炼焦煤", "跳汰、浮选联合", "gb_29446-2019-coking-coal"),
        ("动力煤", "干法选煤", "gb_29446-2019-power-coal"),
        ("动力煤", "跳汰、重介联合", "gb_29446-2019-power-coal"),
    ],
)
def test_coal_process_mapped_from_definition(ctx, tmp_path, coal, process, product_id):
    report = validate(ctx, fill(make_template(ctx, tmp_path, f"map-{product_id}-{abs(hash(process))}.xlsx"), coal=coal, process=process))
    assert report.valid, report.issues
    assert report.request.product_id == product_id
    assert report.request.inputs[gb.INPUT_WASHING_PROCESS].value == process


def test_unknown_coal_type_rejected(ctx, tmp_path):
    report = validate(ctx, fill(make_template(ctx, tmp_path, "badcoal.xlsx"), coal="无烟煤"))
    assert not report.valid


def test_coal_process_mismatch_rejected(ctx, tmp_path):
    # 干法选煤 is declared for 动力煤 only.
    report = validate(ctx, fill(make_template(ctx, tmp_path, "mismatch.xlsx"), coal="炼焦煤", process="干法选煤"))
    assert not report.valid
    assert any("工艺" in issue.message for issue in report.issues)


@pytest.mark.parametrize("period", [gb.PERIOD_FULL_YEAR, "6月", "1月", "12月"])
def test_period_selection(ctx, tmp_path, period):
    report = validate(ctx, fill(make_template(ctx, tmp_path, f"p-{abs(hash(period))}.xlsx"), period=period))
    assert report.valid, report.issues
    assert gb.decode_period_notes(report.request.notes)[0] == period


def test_custom_period_requires_text(ctx, tmp_path):
    report = validate(ctx, fill(make_template(ctx, tmp_path, "cp1.xlsx"), period=gb.PERIOD_CUSTOM, custom_period=None))
    assert not report.valid

    report = validate(
        ctx,
        fill(make_template(ctx, tmp_path, "cp2.xlsx"), period=gb.PERIOD_CUSTOM, custom_period="2026年6月"),
    )
    assert report.valid, report.issues
    period, custom, note = gb.decode_period_notes(report.request.notes)
    assert period == gb.PERIOD_CUSTOM and custom == "2026年6月" and note == "同一内容"


def test_excel_evaluation_is_current_only(ctx, tmp_path):
    report = validate(ctx, fill(make_template(ctx, tmp_path, "cur.xlsx")))
    assert report.valid, report.issues
    assert report.request.selection_mode is StandardSelectionMode.CURRENT


# ---------------------------------------------------------------------------
# 4: stale revision / metadata gate
# ---------------------------------------------------------------------------


def test_stale_rule_revision_rejected(ctx, tmp_path):
    path = make_template(ctx, tmp_path, "stale.xlsx")
    fill(path)
    workbook = load_workbook(path)
    workbook[GB29446_META_SHEET].cell(row=META_ROWS["rule_revision"], column=2, value=1)
    workbook.save(path)
    report = validate(ctx, path)
    assert not report.valid
    assert any("重新生成模板" in issue.message for issue in report.issues)


def test_wrong_standard_id_in_meta_rejected(ctx, tmp_path):
    path = make_template(ctx, tmp_path, "sid.xlsx")
    fill(path)
    workbook = load_workbook(path)
    workbook[GB29446_META_SHEET].cell(row=META_ROWS["standard_id"], column=2, value="gb-00000-2026")
    workbook.save(path)
    report = validate(ctx, path)
    assert not report.valid


def test_wrong_template_version_rejected(ctx, tmp_path):
    path = make_template(ctx, tmp_path, "tv.xlsx")
    fill(path)
    workbook = load_workbook(path)
    workbook[GB29446_META_SHEET].cell(row=META_ROWS["template_version"], column=2, value="9.9")
    workbook.save(path)
    report = validate(ctx, path)
    assert not report.valid


def test_missing_meta_field_rejected(ctx, tmp_path):
    path = make_template(ctx, tmp_path, "meta-missing.xlsx")
    fill(path)
    workbook = load_workbook(path)
    workbook[GB29446_META_SHEET].delete_rows(META_ROWS["adapter_id"])
    workbook.save(path)
    report = validate(ctx, path)
    assert not report.valid
    assert any("_meta" in issue.message for issue in report.issues)


# ---------------------------------------------------------------------------
# 12, 28: structural strictness
# ---------------------------------------------------------------------------


def test_renamed_sheet_rejected(ctx, tmp_path):
    path = fill(make_template(ctx, tmp_path, "rename.xlsx"))
    workbook = load_workbook(path)
    workbook[GB29446_DATA_SHEET].title = "评价数据表"
    workbook.save(path)
    report = validate(ctx, path)
    assert not report.valid


def test_missing_sheet_rejected(ctx, tmp_path):
    path = fill(make_template(ctx, tmp_path, "drop.xlsx"))
    workbook = load_workbook(path)
    del workbook[GB29446_META_SHEET]
    workbook.save(path)
    report = validate(ctx, path)
    assert not report.valid


def test_extra_sheet_rejected(ctx, tmp_path):
    path = fill(make_template(ctx, tmp_path, "extra.xlsx"))
    workbook = load_workbook(path)
    workbook.create_sheet("附加表")
    workbook.save(path)
    report = validate(ctx, path)
    assert not report.valid


def test_changed_header_rejected(ctx, tmp_path):
    path = fill(make_template(ctx, tmp_path, "header.xlsx"))
    workbook = load_workbook(path)
    workbook[GB29446_DATA_SHEET]["A1"] = "自定义字段"
    workbook.save(path)
    report = validate(ctx, path)
    assert not report.valid


def test_real_extra_column_rejected(ctx, tmp_path):
    """A genuinely appended column must be rejected."""
    path = fill(make_template(ctx, tmp_path, "col.xlsx"))
    workbook = load_workbook(path)
    sheet = workbook[GB29446_DATA_SHEET]
    extra_column = sheet.max_column + 1
    sheet.cell(row=1, column=extra_column, value="额外列")
    sheet.cell(row=2, column=extra_column, value="x")
    workbook.save(path)
    assert load_workbook(path)[GB29446_DATA_SHEET].max_column == extra_column
    report = validate(ctx, path)
    assert not report.valid
    assert any("自定义列" in issue.message for issue in report.issues)


# ---------------------------------------------------------------------------
# 16-17: validation must not evaluate
# ---------------------------------------------------------------------------


def test_validation_does_not_run_engine_or_create_record(ctx, tmp_path, monkeypatch):
    from uebench.domain import engine as engine_mod

    calls = {"engine": 0}
    original = engine_mod.EvaluationEngine.evaluate

    def spy(self, *args, **kwargs):
        calls["engine"] += 1
        return original(self, *args, **kwargs)

    monkeypatch.setattr(engine_mod.EvaluationEngine, "evaluate", spy)
    report = validate(ctx, fill(make_template(ctx, tmp_path, "novalidate.xlsx")))
    assert report.valid, report.issues
    assert calls["engine"] == 0, "validate 不得运行 Engine"
    with ctx.database.engine.connect() as connection:
        count = connection.execute(text("SELECT COUNT(*) FROM evaluations")).scalar_one()
    assert count == 0, "validate 不得创建正式记录"


# ---------------------------------------------------------------------------
# 18-19: source integrity
# ---------------------------------------------------------------------------


def test_changed_after_validation_rejected(ctx, tmp_path):
    path = fill(make_template(ctx, tmp_path, "changed.xlsx"), electricity="560")
    report = validate(ctx, path)
    assert report.valid, report.issues
    workbook = load_workbook(path)
    workbook[GB29446_DATA_SHEET]["B8"] = "999"
    workbook.save(path)
    with pytest.raises(ValueError, match="已发生变化"):
        ctx.application.evaluate_workbook(report.import_id)


def test_missing_after_validation_rejected(ctx, tmp_path):
    path = fill(make_template(ctx, tmp_path, "gone.xlsx"))
    report = validate(ctx, path)
    assert report.valid, report.issues
    path.unlink()
    with pytest.raises(ValueError, match="重新选择文件"):
        ctx.application.evaluate_workbook(report.import_id)


def test_failed_evaluation_keeps_batch_retryable(ctx, tmp_path):
    """A batch must not be stranded by a failed formal evaluation.

    The failure is injected into the real ``EvaluationService.evaluate`` call
    that ``evaluate_workbook`` performs, so the retry path is genuinely
    exercised rather than asserted on paper.
    """
    path = fill(make_template(ctx, tmp_path, "retry.xlsx"))
    report = validate(ctx, path)
    assert report.valid, report.issues

    service = ctx.application._evaluation  # noqa: SLF001 - deliberate injection point
    original = service.evaluate
    calls = {"n": 0}

    def flaky(request):
        calls["n"] += 1
        if calls["n"] == 1:
            raise RuntimeError("注入的正式评价失败")
        return original(request)

    service.evaluate = flaky
    try:
        with pytest.raises(RuntimeError, match="注入的正式评价失败"):
            ctx.application.evaluate_workbook(report.import_id)
    finally:
        service.evaluate = original

    with ctx.database.engine.connect() as connection:
        status = connection.execute(
            text("SELECT status FROM import_batches WHERE import_id=:i"), {"i": report.import_id}
        ).scalar_one()
    assert status == "validated", "正式评价失败后批次必须保持 validated 以便重试"

    # A second attempt must succeed and only then advance the batch.
    result = ctx.application.evaluate_workbook(report.import_id)
    assert result.evaluation_id
    with ctx.database.engine.connect() as connection:
        status = connection.execute(
            text("SELECT status FROM import_batches WHERE import_id=:i"), {"i": report.import_id}
        ).scalar_one()
    assert status == STATUS_EVALUATED


def test_prepare_keeps_batch_retryable(ctx, tmp_path):
    """``prepare`` alone must not advance the batch status."""
    path = fill(make_template(ctx, tmp_path, "prepare-only.xlsx"))
    report = validate(ctx, path)
    service = WorkbookImportService(ctx.database, ctx.audit, ctx.standards)
    draft = service.prepare(report.import_id)
    assert draft.request.standard_id == STANDARD_ID
    with ctx.database.engine.connect() as connection:
        status = connection.execute(
            text("SELECT status FROM import_batches WHERE import_id=:i"), {"i": report.import_id}
        ).scalar_one()
    assert status == "validated"


def test_submit_rejects_template_when_rule_revision_changed(ctx, tmp_path):
    """A rule revision installed after validation must not be applied silently.

    Reproduces the reported P1: validate against r2, install a current r3,
    then submit.  ``source_sha256`` is unchanged, so the file hash alone cannot
    detect this; the metadata/resolver gate must.
    """
    from uebench.domain.models import StandardDefinition

    definition = json.loads((ROOT / "data/definitions" / "gb-29446-2019.json").read_text(encoding="utf-8"))
    path = fill(make_template(ctx, tmp_path, "stale-submit.xlsx"))
    report = validate(ctx, path)
    assert report.valid, report.issues
    assert report.rule_revision == 2, "fixture expects the r2 definition"

    # Install a newer revision as the currently applicable definition.  The
    # thresholds change too, so accepting the stale template would change the grade.
    updated = json.loads(json.dumps(definition))
    updated["rule_revision"] = 3
    for product in updated["products"]:
        for indicator in product["indicators"]:
            indicator["thresholds"] = {
                "level_1": {"op": "constant", "value": "9.0", "unit": "kW·h/t"},
                "level_2": {"op": "constant", "value": "9.5", "unit": "kW·h/t"},
                "level_3": {"op": "constant", "value": "9.9", "unit": "kW·h/t"},
            }
            indicator["base_thresholds"] = dict(indicator["thresholds"])
    ctx.standards.install(StandardDefinition.model_validate(updated))

    with pytest.raises(ValueError, match="重新生成模板"):
        ctx.application.evaluate_workbook(report.import_id)

    with ctx.database.engine.connect() as connection:
        status = connection.execute(
            text("SELECT status FROM import_batches WHERE import_id=:i"), {"i": report.import_id}
        ).scalar_one()
        records = connection.execute(text("SELECT COUNT(*) FROM evaluations")).scalar_one()
    assert status == "validated", "拒绝后批次必须保持可重试"
    assert records == 0, "拒绝时不得写入任何正式记录"


def test_legacy_commit_does_not_bypass_revision_gate(ctx, tmp_path):
    """The legacy generic commit path must not be a way around the gate.

    ``commit_workbook`` is retained only for the other standards' generic
    template; a GB29446 batch submitted through it still has to pass the
    revision gate at evaluation time.
    """
    from uebench.domain.models import StandardDefinition

    definition = json.loads((ROOT / "data/definitions" / "gb-29446-2019.json").read_text(encoding="utf-8"))
    path = fill(make_template(ctx, tmp_path, "legacy-commit.xlsx"))
    report = validate(ctx, path)
    assert report.valid, report.issues

    updated = json.loads(json.dumps(definition))
    updated["rule_revision"] = 3
    ctx.standards.install(StandardDefinition.model_validate(updated))

    draft = ctx.application.commit_workbook(report.import_id)
    assert draft.request.standard_id == STANDARD_ID
    # Even though the legacy path marked the batch committed, the formal
    # evaluation入口 still re-runs the gate.
    with pytest.raises(ValueError, match="重新生成模板"):
        ctx.application.evaluate_workbook(report.import_id)


# ---------------------------------------------------------------------------
# 20-23, 26: request/result equivalence, boundary parity, record lifecycle
# ---------------------------------------------------------------------------


def test_gui_request_equals_excel_request(ctx, tmp_path):
    standard = ctx.application.get_published_standard(STANDARD_ID)
    report = validate(
        ctx,
        fill(make_template(ctx, tmp_path, "eq.xlsx"), period=gb.PERIOD_CUSTOM, custom_period="2026年6月"),
    )
    assert report.valid, report.issues
    excel_request = report.request
    gui = gui_request(standard, period=gb.PERIOD_CUSTOM, custom_period="2026年6月")
    assert excel_request.evaluation_date == gui.evaluation_date
    assert excel_request.standard_id == gui.standard_id
    assert excel_request.product_id == gui.product_id
    assert excel_request.selection_mode == gui.selection_mode
    assert excel_request.input_mode == gui.input_mode
    assert excel_request.organization_name == gui.organization_name
    assert excel_request.notes == gui.notes
    assert {k: v.value for k, v in excel_request.inputs.items()} == {k: v.value for k, v in gui.inputs.items()}
    assert excel_request == gui, "规范化后的 GUI 与 Excel 请求必须完全相等"


def test_real_gui_form_request_equals_excel_request(ctx, tmp_path, monkeypatch):
    """Equality against the request the actual desktop form produces.

    The previous case compared Excel against a hand-built request.  Here the
    real ``MainWindow`` widgets are filled through the real Qt slots, and the
    request that ``_collect_gb29446_request`` returns must equal the Excel
    adapter's canonical request.
    """
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    from PySide6.QtWidgets import QApplication

    from uebench.ui.main_window import MainWindow

    app = QApplication.instance() or QApplication([])
    window = MainWindow(ctx)
    try:
        window.refresh_all()
        # Drive the real controls rather than assigning the request directly.
        window.eval_standard.setCurrentIndex(window.eval_standard.findData(STANDARD_ID))
        window._standard_changed()
        window.gb29446_coal_type.setCurrentIndex(window.gb29446_coal_type.findData("gb_29446-2019-coking-coal"))
        window._gb29446_coal_changed()
        window.gb29446_process.setCurrentIndex(window.gb29446_process.findData("重介"))
        window.gb29446_organization.setText("宁夏测试企业")
        window.gb29446_period.setCurrentIndex(window.gb29446_period.findData(gb.PERIOD_CUSTOM))
        window.gb29446_custom_period.setText("2026年6月")
        window.gb29446_electricity.setText("560")
        window.gb29446_raw_coal.setText("100")
        window.gb29446_notes.setText("同一内容")
        gui = window._collect_gb29446_request()
    finally:
        window.close()

    report = validate(
        ctx,
        fill(make_template(ctx, tmp_path, "real-eq.xlsx"), period=gb.PERIOD_CUSTOM, custom_period="2026年6月"),
    )
    assert report.valid, report.issues
    excel_request = report.request
    assert {k: v.value for k, v in gui.inputs.items()} == {k: v.value for k, v in excel_request.inputs.items()}
    assert gui.organization_name == excel_request.organization_name
    assert gui.notes == excel_request.notes
    assert gui.product_id == excel_request.product_id
    assert gui.input_mode == excel_request.input_mode
    assert gui.selection_mode == excel_request.selection_mode
    # ``evaluation_date`` is the day the GUI is used, so pin it for comparison.
    assert excel_request.model_copy(update={"evaluation_date": gui.evaluation_date}).model_dump() == (
        gui.model_dump()
    ), "真实 GUI 表单产生的请求必须与 Excel 适配器的规范请求一致"


def test_gui_result_equals_excel_result_projection(ctx, tmp_path):
    standard = ctx.application.get_published_standard(STANDARD_ID)
    report = validate(ctx, fill(make_template(ctx, tmp_path, "res.xlsx")))
    assert report.valid, report.issues
    excel_result = ctx.application.evaluate_workbook(report.import_id)
    gui_result = ctx.application.evaluate(gui_request(standard))

    def projection(result):
        def thresholds(item):
            data = item.corrected_thresholds
            if hasattr(data, "model_dump"):
                data = data.model_dump()
            return {key: str(getattr(value, "value", value)) for key, value in dict(data).items()}

        return {
            "standard_id": result.standard_id,
            "standard_version": result.standard_version,
            "rule_revision": result.rule_revision,
            "product_id": result.product_id,
            "numeric_contract_version": result.numeric_contract_version,
            "numeric_profile_id": result.numeric_profile_id,
            "calculator_version": result.calculator_version,
            "numeric_behavior_version": result.numeric_behavior_version,
            "results": [
                {
                    "indicator_id": item.indicator_id,
                    "actual_value": str(item.actual_value),
                    "grade": str(item.grade),
                    "display_values": {k: str(getattr(v, "value", v)) for k, v in item.display_values.items()},
                    "corrected_thresholds": thresholds(item),
                    "sources": [(s.page, s.clause, s.table) for s in item.source_references],
                }
                for item in result.results
            ],
        }

    assert projection(excel_result) == projection(gui_result), "GUI 与 Excel 的业务投影必须一致"
    # Identity/timestamp are allowed to differ: they are two formal records.
    assert excel_result.evaluation_id != gui_result.evaluation_id


def test_full_value_boundary_parity(ctx, tmp_path):
    excel_result, gui_result = boundary_parity(ctx, make_template(ctx, tmp_path, "boundary.xlsx"))
    assert excel_result.results[0].actual_value == gui_result.results[0].actual_value == Decimal("5.0000004")
    assert excel_result.results[0].grade is gui_result.results[0].grade
    # T+δ must not collapse onto T through implicit ROUND6.
    assert excel_result.results[0].grade is Grade.LEVEL_2


def test_formal_excel_evaluation_creates_record(ctx, tmp_path):
    report = validate(ctx, fill(make_template(ctx, tmp_path, "record.xlsx")))
    result = ctx.application.evaluate_workbook(report.import_id)
    loaded = ctx.application.get_evaluation(result.evaluation_id)
    assert loaded is not None
    request, stored, snapshot = loaded
    assert request.standard_id == STANDARD_ID
    assert stored.results[0].grade == result.results[0].grade
    assert snapshot.rule_revision == ctx.application.get_published_standard(STANDARD_ID).rule_revision
    with ctx.database.engine.connect() as connection:
        status = connection.execute(
            text("SELECT status FROM import_batches WHERE import_id=:i"), {"i": report.import_id}
        ).scalar_one()
    assert status == STATUS_EVALUATED


def test_import_batch_links_to_evaluation_audit(ctx, tmp_path):
    report = validate(ctx, fill(make_template(ctx, tmp_path, "audit.xlsx")))
    result = ctx.application.evaluate_workbook(report.import_id)
    with ctx.database.engine.connect() as connection:
        rows = connection.execute(
            text("SELECT details_json FROM audit_log WHERE action='WORKBOOK_EVALUATION'")
        ).scalars().all()
    assert rows, "必须写入 WORKBOOK_EVALUATION 审计"
    matched = [json.loads(row) for row in rows if json.loads(row).get("evaluation_id") == result.evaluation_id]
    assert matched, "审计事件必须能反查该正式记录"
    details = matched[0]
    assert details["import_id"] == report.import_id
    assert details["profile_id"] == GB29446_PROFILE_ID
    assert re.fullmatch(r"[0-9a-f]{64}", details["source_sha256"])
    # The audit link must not require a new evaluations column.
    with ctx.database.engine.connect() as connection:
        columns = {row[1] for row in connection.execute(text("PRAGMA table_info(evaluations)"))}
    assert "import_id" not in columns


def test_excel_record_lifecycle_survives_restart(ctx, tmp_path):
    """RS02 lifecycle inheritance across a real process boundary.

    The record is created through the Excel adapter, the context is disposed, and
    a **separate Python process** reopens the same data directory and restores
    Request / Result / Rule Snapshot.  An in-process rebuild would not prove the
    lifecycle, and neither would a read that silently recalculated.
    """
    report = validate(ctx, fill(make_template(ctx, tmp_path, "life.xlsx")))
    result = ctx.application.evaluate_workbook(report.import_id)
    evaluation_id = result.evaluation_id
    data_root = tmp_path / "中文数据"
    script = _restart_probe_script(evaluation_id, data_root)
    ctx.database.dispose()

    completed = subprocess.run(
        [sys.executable, "-B", str(script)],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        cwd=str(ROOT),
        timeout=180,
    )
    assert completed.returncode == 0, f"subprocess failed:\n{completed.stdout}\n{completed.stderr}"
    payload = json.loads(completed.stdout.strip().splitlines()[-1])
    assert payload["restored"] is True, payload
    assert payload["standard_id"] == STANDARD_ID
    assert payload["evaluation_id"] == evaluation_id
    assert payload["grade"] == str(result.results[0].grade)
    assert payload["numeric_profile_id"] == result.numeric_profile_id
    assert payload["rule_snapshot_revision"] == result.rule_revision
    assert payload["engine_calls"] == 0, "只读查看历史不得重新运行 Engine"


def _restart_probe_script(evaluation_id: str, data_root: Path) -> Path:
    """Write a standalone script that reopens the data directory in a new process."""
    script = data_root / "_restart_probe.py"
    script.parent.mkdir(parents=True, exist_ok=True)
    script.write_text(
        "import json, sys\n"
        "from pathlib import Path\n"
        "sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))\n"
        "from uebench.bootstrap import create_context\n"
        "from uebench.domain import engine as engine_mod\n"
        "\n"
        "calls = {'n': 0}\n"
        "original = engine_mod.EvaluationEngine.evaluate\n"
        "def spy(self, *a, **k):\n"
        "    calls['n'] += 1\n"
        "    return original(self, *a, **k)\n"
        "engine_mod.EvaluationEngine.evaluate = spy\n"
        "\n"
        f"context = create_context(Path(r'{data_root}'))\n"
        "try:\n"
        f"    loaded = context.application.get_evaluation({evaluation_id!r})\n"
        "    if loaded is None:\n"
        "        print(json.dumps({'restored': False}))\n"
        "        raise SystemExit(0)\n"
        "    request, result, snapshot = loaded\n"
        "    print(json.dumps({\n"
        "        'restored': True,\n"
        "        'standard_id': request.standard_id,\n"
        "        'evaluation_id': result.evaluation_id,\n"
        "        'grade': str(result.results[0].grade),\n"
        "        'numeric_profile_id': result.numeric_profile_id,\n"
        "        'rule_snapshot_revision': snapshot.rule_revision,\n"
        "        'engine_calls': calls['n'],\n"
        "    }, ensure_ascii=False))\n"
        "finally:\n"
        "    context.database.dispose()\n",
        encoding="utf-8",
    )
    return script


def test_import_to_export_consistency(ctx, tmp_path):
    report = validate(ctx, fill(make_template(ctx, tmp_path, "exp.xlsx")))
    result = ctx.application.evaluate_workbook(report.import_id)
    output = ctx.application.export_evaluation(result.evaluation_id, tmp_path / "export.xlsx")
    workbook = load_workbook(output, data_only=True)
    summary = workbook["对标结论"]
    header_row = 5
    headers = [summary.cell(row=header_row, column=c).value for c in range(1, 8)]
    assert headers[:3] == ["指标", "实际值", "单位"]
    indicator_row = header_row + 1
    assert summary.cell(row=indicator_row, column=1).value == result.results[0].indicator_name
    assert str(summary.cell(row=indicator_row, column=2).value) == str(result.results[0].actual_value)
    from uebench.domain.models import GRADE_LABELS

    assert summary.cell(row=indicator_row, column=7).value == GRADE_LABELS[result.results[0].grade]
    # Export must read the stored record, not recalculate.
    from uebench.domain import engine as engine_mod

    original = engine_mod.EvaluationEngine.evaluate
    engine_mod.EvaluationEngine.evaluate = lambda *a, **k: (_ for _ in ()).throw(
        AssertionError("导出不得重新计算")
    )
    try:
        ctx.application.export_evaluation(result.evaluation_id, tmp_path / "export2.xlsx")
    finally:
        engine_mod.EvaluationEngine.evaluate = original


# ---------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------


def _inject_cached_formula(path: Path, formula: str, cached: str) -> None:
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        payload = {name: archive.read(name) for name in names}
    target = "xl/worksheets/sheet2.xml"
    text = payload[target].decode("utf-8")
    pattern = re.compile(r'(<c r="B8"[^>]*>)(<f>)([^<]*)(</f>)(<v\s*/>)')
    match = pattern.search(text)
    assert match, "unexpected sheet XML layout"
    replacement = f"{match.group(1)}<f>{formula}</f><v>{cached}</v>"
    payload[target] = (text[: match.start()] + replacement + text[match.end():]).encode("utf-8")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name in names:
            archive.writestr(name, payload[target] if name == target else payload[name])
