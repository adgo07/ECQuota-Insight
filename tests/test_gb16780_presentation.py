"""水泥四模式录入、保存快照与 Word 同源展示的业务回归。"""
import os
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from pathlib import Path
from unittest.mock import Mock

import pytest
from PySide6.QtWidgets import QApplication

from uebench.domain.models import StandardDefinition, InputValue, EvaluationRequest, InputMode, CementDetails, CementFuelLine
from uebench.domain.gb16780 import Gb16780Calculator, DEDUCTION_KEYS, MODE_INDICATORS
from uebench.application.cement_reporting import build_cement_report_view, ALTITUDE_DISCLOSURE
from uebench.infrastructure.word_report import WordReportWriter
from uebench.ui.cement_panel import CementEvaluationPanel
from uebench.ui.record_detail import RecordDetailPage

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture
def standard():
    return StandardDefinition.model_validate_json((ROOT / "data/definitions/gb-16780-2021.json").read_text(encoding="utf-8"))


def grinding_request():
    return EvaluationRequest(evaluation_date="2026-10-29", standard_id="gb-16780-2021", product_id="cement",
        input_mode=InputMode.DETAIL, inputs={"production_mode": InputValue(value="grinding_station"),
            "evaluation_period": InputValue(value="2025年"), "production_line_id": InputValue(value="A"),
            "cement_production_t": InputValue(value="1000", unit="t"),
            "grinding_electricity_kwh": InputValue(value="23000", unit="kW·h")})


def integrated_request(*, altitude="1500", ratio="75"):
    request = grinding_request()
    request.inputs["production_mode"] = InputValue(value="integrated")
    fields = {
        "altitude_m": (altitude, "m"), "clinker_production_t": ("1000", "t"),
        "clinker_consumption_t": ("750", "t"), "clinker_ratio_percent": (ratio, "%"),
        "clinker_electricity_kwh": ("50000", "kW·h"),
        "waste_heat_generation_kwh": ("10000", "kW·h"), "waste_heat_self_use_kwh": ("0", "kW·h"),
        "cement_grinding_packaging_kwh": ("20000", "kW·h"),
        "admixture_preparation_kwh": ("1000", "kW·h"), "gypsum_preparation_kwh": ("1000", "kW·h"),
        "auxiliary_electricity_kwh": ("2000", "kW·h"), "cement_fuel_kgce": ("1000", "kgce"),
    }
    request.inputs.update({key: InputValue(value=value, unit=unit) for key, (value, unit) in fields.items()})
    if altitude != "1500":
        request.inputs["pressure_pa"] = InputValue(value="80000", unit="Pa")
    request.cement_details = CementDetails(fuel_lines=[CementFuelLine(line_id="f1", name="燃煤",
        amount_kg="100000", low_heat_kj_per_kg="29307.6")])
    return request


@pytest.fixture
def panel(standard):
    app = QApplication.instance() or QApplication([])
    widget = CementEvaluationPanel(Mock())
    widget.set_standard(standard)
    yield widget
    widget.close()
    app.processEvents()


def test_four_deductions_start_blank_with_approved_placeholder(panel):
    for key in DEDUCTION_KEYS:
        assert panel.fields[key].text() == ""
        assert panel.fields[key].placeholderText() == "若已扣除，则不填"
    assert "standard_coal_heat_kj_per_kgce" not in panel.fields


def test_switching_mode_discards_hidden_values_and_only_grinding_inputs_remain(panel):
    panel.fields["clinker_production_t"].setText("1000")
    panel.fields[DEDUCTION_KEYS[0]].setText("1")
    panel.mode.setCurrentIndex(panel.mode.findData("grinding_station"))
    assert set(panel.fields) == {"cement_production_t", "grinding_electricity_kwh"}
    panel.line.setText("A")
    panel.period.setText("2025年")
    panel.fields["cement_production_t"].setText("1000")
    panel.fields["grinding_electricity_kwh"].setText("23000")
    request = panel.collect_request()
    assert "clinker_production_t" not in request.inputs
    assert not any(key in request.inputs for key in DEDUCTION_KEYS)
    assert request.cement_details.fuel_lines == []
    panel.mode.setCurrentIndex(panel.mode.findData("integrated"))
    assert panel.fields["clinker_production_t"].text() == ""
    assert panel.fields[DEDUCTION_KEYS[0]].text() == ""


def test_pressure_only_visible_when_triggered_and_old_value_cleared(panel):
    panel.fields["altitude_m"].setText("1600")
    assert not panel.fields["pressure_pa"].isHidden()
    panel.fields["pressure_pa"].setText("80000")
    panel.fields["altitude_m"].setText("1500")
    assert panel.fields["pressure_pa"].isHidden()
    assert panel.fields["pressure_pa"].text() == ""


def test_snapshot_report_does_not_recalculate_and_grinding_has_one_indicator(standard, monkeypatch):
    request = grinding_request()
    result = Gb16780Calculator().evaluate(standard, request)
    monkeypatch.setattr(Gb16780Calculator, "evaluate", lambda *a, **k: pytest.fail("报告不得重新计算"))
    view = build_cement_report_view(request, result, standard)
    assert len(view.indicators) == 1
    assert view.indicators[0].key == MODE_INDICATORS["grinding_station"][0]
    assert view.indicators[0].actual == "23.00"
    assert view.indicators[0].grade == "1级"
    assert "总等级" not in str(view)
    assert not any("海拔修正" in line for line in view.disclosures)


def test_report_rejects_partial_and_snapshot_identity_mismatch(standard):
    request = grinding_request()
    del request.inputs["grinding_electricity_kwh"]
    result = Gb16780Calculator().evaluate(standard, request)
    with pytest.raises(ValueError, match="完整"):
        build_cement_report_view(request, result, standard)
    assert not build_cement_report_view(request, result, standard, formal=False).complete
    standard.rule_revision += 1
    with pytest.raises(ValueError, match="身份"):
        build_cement_report_view(request, result, standard)


def test_word_and_two_tab_record_detail_use_same_saved_view(standard, tmp_path):
    from docx import Document
    app = QApplication.instance() or QApplication([])
    request = grinding_request()
    result = Gb16780Calculator().evaluate(standard, request)
    view = build_cement_report_view(request, result, standard)
    path = WordReportWriter().write(view, tmp_path / "水泥报告.docx")
    document = Document(path)
    text = "\n".join([p.text for p in document.paragraphs] + [c.text for t in document.tables for r in t.rows for c in r.cells])
    assert view.indicators[0].label in text
    assert "23.00" in text and "1级" in text
    assert "总等级" not in text
    assert "clinker" not in text and "grinding_electricity" not in text
    page = RecordDetailPage()
    page.show_record((request, result, standard), view, "")
    assert page.tabs.count() == 2
    assert "23.00" in page.result_content.toPlainText()
    assert page.export_button.isEnabled()
    page.close()
    app.processEvents()


def test_calculate_and_save_suppresses_double_click_but_allows_new_record(panel, standard, monkeypatch):
    request = grinding_request()
    result = Gb16780Calculator().evaluate(standard, request)
    panel.application.evaluate_or_preview.return_value = (result, True)
    monkeypatch.setattr(panel, "collect_request", lambda: request)
    times = iter((1.0, 1.1, 2.0))
    monkeypatch.setattr("uebench.ui.cement_panel.monotonic", lambda: next(times))
    panel.calculate_and_save()
    panel.calculate_and_save()
    assert panel.application.evaluate_or_preview.call_count == 1
    panel.calculate_and_save()
    assert panel.application.evaluate_or_preview.call_count == 2
    assert panel.export_button.isEnabled()


def test_partial_preview_does_not_offer_formal_word(panel, standard, monkeypatch):
    request = grinding_request()
    del request.inputs["grinding_electricity_kwh"]
    result = Gb16780Calculator().evaluate(standard, request)
    panel.application.evaluate_or_preview.return_value = (result, False)
    monkeypatch.setattr(panel, "collect_request", lambda: request)
    panel.calculate_and_save()
    assert panel.record_id is None
    assert not panel.export_button.isEnabled()
    assert "未保存" in panel.message.text()


def test_shared_system_keeps_original_quantity_and_participating_line_detail(panel):
    panel.mode.setCurrentIndex(panel.mode.findData("grinding_station"))
    panel.line.setText("A")
    panel.period.setText("2025年")
    panel.add_allocation()
    editor = panel.allocations[0]
    editor.amount.setText("12000")
    editor.lines.item(0, 0).setText("A")
    editor.lines.item(0, 1).setText("1000")
    editor.add_line()
    editor.lines.item(1, 0).setText("B")
    editor.lines.item(1, 1).setText("2000")
    request = panel.collect_request()
    allocation = request.cement_details.shared_allocations[0]
    assert allocation.original_amount == "12000"
    assert allocation.target_line_id == "A"
    assert allocation.line_production_t == {"A": "1000", "B": "2000"}
    assert allocation.product == "cement"
    panel.mode.setCurrentIndex(panel.mode.findData("clinker_only"))
    assert panel.allocations == []


def test_untriggered_correction_has_no_ordinary_technical_disclosure(standard):
    request = integrated_request()
    result = Gb16780Calculator().evaluate(standard, request)
    saved_before = result.model_dump(mode="json")
    view = build_cement_report_view(request, result, standard)
    assert view.disclosures == ()
    detail = next(w for item in result.results for w in item.warnings if "ECQ-STD-GB16780-003" in w)
    assert detail in view.basis
    assert result.model_dump(mode="json") == saved_before


def test_actual_corrections_show_short_business_text_and_keep_full_detail(standard):
    app = QApplication.instance() or QApplication([])
    request = integrated_request(altitude="1600", ratio="80")
    request.inputs[DEDUCTION_KEYS[0]] = InputValue(value="1", unit="tce")
    result = Gb16780Calculator().evaluate(standard, request)
    saved_before = result.model_dump(mode="json")
    view = build_cement_report_view(request, result, standard)
    assert view.disclosures.count(ALTITUDE_DISCLOSURE) == 1
    assert any("先作海拔修正" in text for text in view.disclosures)
    assert any("扣除" in text for text in view.disclosures)
    assert all("ECQ-STD-" not in text and "2019" not in text for text in view.disclosures)
    technical = set(w for item in result.results for w in item.warnings if "ECQ-STD-" in w or "2019" in w)
    assert technical and technical.issubset(view.basis)
    assert result.model_dump(mode="json") == saved_before
    page = RecordDetailPage()
    page.show_record((request, result, standard), view, "")
    assert "ECQ-STD-" not in page.result_content.toPlainText()
    assert "2019" not in page.result_content.toPlainText()
    assert ALTITUDE_DISCLOSURE in page.result_content.toPlainText()
    assert all(text in page.input_content.toPlainText() for text in technical)
    page.close()
    app.processEvents()
