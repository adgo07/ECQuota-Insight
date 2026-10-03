"""ECQ-RS04 — GB29446 Product Golden Gate (v1).

Golden expects come from the standard text (表1 / 表2 / 式(1) / 表A.1), the
published ``StandardDefinition`` and Numeric Contract v1 — never from current
software output.  The gate re-derives every expected value from the declared
authority, evaluates through the formal Application/Engine path, and compares a
**whitelist business projection** only.

Gate A  Business Golden      - every case, formal path, projection == golden
Gate B  Product Lifecycle    - formal record, real restart, history projection
Gate C  Excel Golden         - GUI request == Excel request, both == golden
Gate D  Regression           - existing gates keep passing (run separately)
"""
from __future__ import annotations

import json
import re
import subprocess
import sys
from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import text

from uebench.application.gb29446 import decode_period_notes
from uebench.bootstrap import create_context
from uebench.domain.models import (
    EvaluationRequest,
    InputMode,
    InputValue,
    StandardDefinition,
    StandardSelectionMode,
)

ROOT = Path(__file__).resolve().parents[1]
GOLDEN_PATH = ROOT / "tests" / "golden" / "gb29446_product_golden_v1.json"
DEFINITION_PATH = ROOT / "data" / "definitions" / "gb-29446-2019.json"
ISSUE_REGISTER_PATH = ROOT / "STANDARD_ISSUES_REGISTER.md"
STANDARD_ID = "gb-29446-2019"
PROFILE_ID = "ECQUOTA_DECIMAL_FULL_VALUE_V1"
EVALUATION_DATE = date(2026, 6, 1)

GOLDEN = json.loads(GOLDEN_PATH.read_text(encoding="utf-8"))

#: Frozen business keys.  Everything not listed is deliberately excluded so the
#: Golden does not pin environment, identity or presentation details.
FROZEN_TOP_LEVEL = (
    "standard_id",
    "standard_number",
    "standard_version",
    "rule_revision",
    "product_id",
    "numeric_contract_version",
    "numeric_profile_id",
)


# ---------------------------------------------------------------------------
# Golden whitelist projection
# ---------------------------------------------------------------------------


def _dec(value) -> str | None:
    if value is None:
        return None
    if not isinstance(value, Decimal):
        value = Decimal(str(value))
    return format(value, "f")


def _dec_text(value) -> str | None:
    """Canonical decimal text: numerically equal values compare equal.

    ``Decimal('5.00')`` and ``Decimal('5.0')`` are the same business value but
    format differently, so every comparison normalises trailing zeros.  This
    keeps the Golden pinned to business values rather than to scale.
    """
    if value is None:
        return None
    return format(Decimal(str(value)).normalize(), "f")


def business_projection(result, request: EvaluationRequest) -> dict:
    """Deterministic, whitelist-only view of a formal result.

    Explicitly excludes: evaluation_id, evaluated_at, SQLite ids, absolute
    paths, source reference page numbers, JSON key order, rule_snapshot_sha256,
    whole-definition SHA, full trace text, complete warning wording,
    calculator_version, numeric_behavior_version and UI details.
    """
    period, custom_period, note = decode_period_notes(request.notes)
    projection = {key: getattr(result, key) for key in FROZEN_TOP_LEVEL}
    projection.update(
        {
            "process": request.inputs["washing_process"].value,
            "e_d_input": request.inputs["electricity_consumption"].value,
            "m_input": request.inputs["raw_coal_input"].value,
            "period": period,
            "custom_period": custom_period,
            "note": note,
            "organization_name": request.organization_name,
            "indicators": [
                {
                    "indicator_id": item.indicator_id,
                    "actual_value": _dec(item.actual_value),
                    "unit": item.unit,
                    "k": _dec(item.display_values.get("process_factor")),
                    "base_thresholds": {
                        key: _dec(value) for key, value in sorted(item.base_thresholds.items())
                    },
                    "corrected_thresholds": {
                        key: _dec(value) for key, value in sorted(item.corrected_thresholds.items())
                    },
                    "grade": str(item.grade),
                    "warnings_empty": item.warnings == [],
                    "sources": sorted(
                        {(ref.standard_number, ref.clause, ref.table) for ref in item.source_references}
                    ),
                    "trace_operations": [step.operation for step in item.calculation_trace],
                    "has_grade_comparison": any(
                        step.operation == "grade_comparison" for step in item.calculation_trace
                    ),
                    "has_explicit_rounding_step": any(
                        "round" in step.operation.lower() for step in item.calculation_trace
                    ),
                }
                for item in result.results
            ],
        }
    )
    # Freeze "no implicit rounding" as an invariant over the trace text too.
    projection["round_places_declared"] = any(
        re.search(r"round_places\s*=\s*\d+", json.dumps(step.model_dump(), ensure_ascii=False, default=str))
        for item in result.results
        for step in item.calculation_trace
    )
    return projection


def compared_projection(projection: dict) -> dict:
    """Drop non-frozen diagnostics before comparing against the Golden."""
    return {
        key: projection[key]
        for key in (
            "standard_id",
            "standard_number",
            "standard_version",
            "rule_revision",
            "product_id",
            "numeric_contract_version",
            "numeric_profile_id",
            "process",
            "e_d_input",
            "m_input",
            "period",
            "custom_period",
            "note",
            "organization_name",
        )
    } | {
        "indicators": [
            {
                key: item[key]
                for key in (
                    "indicator_id",
                    "actual_value",
                    "unit",
                    "k",
                    "base_thresholds",
                    "corrected_thresholds",
                    "grade",
                    "warnings_empty",
                    "has_grade_comparison",
                    "has_explicit_rounding_step",
                )
            }
            | {"sources": [list(entry) for entry in item["sources"]]}
            for item in projection["indicators"]
        ]
    }


# ---------------------------------------------------------------------------
# Independent derivation from the declared authority
# ---------------------------------------------------------------------------


def definition_json() -> dict:
    return json.loads(DEFINITION_PATH.read_text(encoding="utf-8"))


def definition_standard() -> StandardDefinition:
    return StandardDefinition.model_validate_json(DEFINITION_PATH.read_text(encoding="utf-8"))


def product_of(definition: dict, coal_type: str) -> dict:
    return next(
        p for p in definition["products"] if p["selection_values"]["coal_type"] == coal_type
    )


def coefficient_of(definition: dict, coal_type: str, process: str) -> Decimal:
    """Read k from 附录A 表A.1 as transcribed in the definition."""
    for indicator in product_of(definition, coal_type)["indicators"]:
        for display in indicator.get("display_calculations", []):
            if display.get("key") != "process_factor":
                continue
            for row in display["formula"]["rows"]:
                if row["condition"]["value"] == process:
                    return Decimal(row["expression"]["value"])
    raise AssertionError(f"未找到系数：{coal_type} / {process}")


def coefficients_of(definition: dict, coal_type: str) -> dict[str, str]:
    """All k values declared for a coal type, as canonical decimal text."""
    out: dict[str, str] = {}
    for indicator in product_of(definition, coal_type)["indicators"]:
        for display in indicator.get("display_calculations", []):
            if display.get("key") != "process_factor":
                continue
            for row in display["formula"]["rows"]:
                out[row["condition"]["value"]] = _dec_text(row["expression"]["value"])
    return out


def thresholds_of(definition: dict, coal_type: str) -> dict[str, Decimal]:
    raw = product_of(definition, coal_type)["indicators"][0]["thresholds"]
    return {key: Decimal(value["value"]) for key, value in raw.items() if value}


def independently_derive(definition: dict, case: dict) -> dict:
    """Recompute the Golden expected from 原文-derived facts, not software output."""
    inputs = case["inputs"]
    if inputs["electricity_consumption"] is None:
        return {"k": None, "e_d": None, "grade": None}
    k = coefficient_of(definition, inputs["coal_type"], inputs["washing_process"])
    e_d = (Decimal(inputs["electricity_consumption"]) * k) / Decimal(inputs["raw_coal_input"])
    limits = thresholds_of(definition, inputs["coal_type"])
    if e_d <= limits["level_1"]:
        grade = "LEVEL_1"
    elif e_d <= limits["level_2"]:
        grade = "LEVEL_2"
    elif e_d <= limits["level_3"]:
        grade = "LEVEL_3"
    else:
        grade = "NOT_QUALIFIED"
    return {"k": _dec_text(k), "e_d": _dec_text(e_d), "grade": grade}


# ---------------------------------------------------------------------------
# Applicability Gate
# ---------------------------------------------------------------------------


ISSUE_ROW_RE = re.compile(r"^\|\s*状态\s*\|\s*(?P<value>.+?)\s*\|\s*$", re.MULTILINE)


def _issue_status(issue_id: str) -> str | None:
    """Read the status code of one registered Standard Issue.

    The register is a Markdown table, so the status row is matched with a
    tolerant regex rather than by exact line prefix.
    """
    lines = ISSUE_REGISTER_PATH.read_text(encoding="utf-8").splitlines()
    heading = next(
        (index for index, line in enumerate(lines) if line.startswith(f"### {issue_id} ")),
        None,
    )
    if heading is None:
        return None
    section: list[str] = []
    for line in lines[heading + 1:]:
        if line.startswith("### "):
            break
        section.append(line)
    match = ISSUE_ROW_RE.search("\n".join(section))
    if not match:
        return None
    raw = match.group("value").strip()
    # Strip emphasis/backticks and any trailing Chinese parenthetical note.
    return raw.strip("`* ").replace("`", "").split("（")[0].strip()


def applicability_report() -> dict:
    definition = definition_json()
    return {
        "standard_id": definition["id"],
        "standard_version": definition["version"],
        "standard_source_sha256": definition["source_sha256"],
        "rule_revision": definition["rule_revision"],
        "numeric_profile_id": PROFILE_ID,
        "standard_issue_status": _issue_status(GOLDEN["standard_issue_ref"]),
        "golden_declared": {
            "standard_id": GOLDEN["standard_id"],
            "standard_version": GOLDEN["standard_version"],
            "standard_source_sha256": GOLDEN["standard_source_sha256"],
            "rule_revision": GOLDEN["rule_revision"],
            "numeric_profile_id": GOLDEN["numeric_profile_id"],
            "standard_issue_status": GOLDEN["standard_issue_expected_status"],
        },
    }


def _applicability_mismatches() -> list[str]:
    report = applicability_report()
    declared = report["golden_declared"]
    return [
        f"{key}: 当前={report[key]!r} Golden声明={declared[key]!r}"
        for key in (
            "standard_id",
            "standard_version",
            "standard_source_sha256",
            "rule_revision",
            "numeric_profile_id",
            "standard_issue_status",
        )
        if str(report[key]) != str(declared[key])
    ]


# ---------------------------------------------------------------------------
# Fixtures and helpers
# ---------------------------------------------------------------------------


@pytest.fixture(scope="module")
def definition() -> dict:
    return definition_json()


@pytest.fixture
def ctx(tmp_path):
    context = create_context(tmp_path / "中文数据")
    context.standards.install(definition_standard())
    try:
        yield context
    finally:
        context.database.dispose()


def normal_cases() -> list[dict]:
    """Business cases: they freeze a concrete grade."""
    return [case for case in GOLDEN["cases"] if "grade_in" not in case["expected"]]


def safety_cases() -> list[dict]:
    """Safety cases: they only freeze "no official grade may be produced"."""
    return [case for case in GOLDEN["cases"] if "grade_in" in case["expected"]]


def case_by_id(case_id: str) -> dict:
    return next(case for case in GOLDEN["cases"] if case["case_id"] == case_id)


def build_request(
    case: dict,
    standard: StandardDefinition,
    *,
    organization: str = "Golden 验证企业",
    notes: str = "核算周期：全年\n备注：RS04 Product Golden v1",
    include_single_process_flag: bool = True,
    include_enterprise_status: bool = True,
) -> EvaluationRequest:
    """Build the canonical request for a Golden case.

    Two optional product-level inputs are switchable because the GB29446 Excel
    template exposes the nine user fields (企业名称 / 评价日期 / 核算周期 /
    自定义周期 / 煤种 / 选煤工艺 / E_d / m / 备注) and no product-level field:

    * ``single_coal_single_process`` — the engine exempts it from GB29446
      validation, so it cannot change the outcome;
    * ``enterprise_status`` — it drives only the 限定值/准入值 compliance
      conclusion, which is ``None`` in DETAIL mode and which this Golden
      deliberately does not freeze (the mandate keeps 企业属性 compliance out of
      the Golden because RS01/RS02 own that semantics).

    Gate C therefore compares the Excel request against a request built from the
    same business inputs that can actually reach the workbook.
    """
    inputs = case["inputs"]
    product = next(
        p for p in standard.products if p.selection_values["coal_type"] == inputs["coal_type"]
    )
    supplied: dict[str, InputValue] = {
        "washing_process": InputValue(value=inputs["washing_process"]),
    }
    if include_single_process_flag:
        supplied["single_coal_single_process"] = InputValue(
            value=inputs["single_coal_single_process"]
        )
    if include_enterprise_status and inputs.get("enterprise_status"):
        supplied["enterprise_status"] = InputValue(value=inputs["enterprise_status"])
    if inputs["electricity_consumption"] is not None:
        supplied["electricity_consumption"] = InputValue(
            value=str(inputs["electricity_consumption"]), unit="kW·h"
        )
    if inputs["raw_coal_input"] is not None:
        supplied["raw_coal_input"] = InputValue(value=str(inputs["raw_coal_input"]), unit="t")
    return EvaluationRequest(
        evaluation_date=EVALUATION_DATE,
        standard_id=standard.id,
        product_id=product.id,
        selection_mode=StandardSelectionMode.CURRENT,
        input_mode=InputMode.DETAIL,
        inputs=supplied,
        organization_name=organization,
        notes=notes,
    )


# ===========================================================================
# Applicability Gate
# ===========================================================================


def test_golden_declares_required_metadata() -> None:
    for key in (
        "golden_id",
        "golden_version",
        "standard_id",
        "standard_version",
        "rule_revision",
        "standard_source_sha256",
        "numeric_profile_id",
        "standard_issue_ref",
        "standard_issue_expected_status",
        "approval_basis",
        "derivation_method",
        "business_data_summary",
        "cases",
    ):
        assert key in GOLDEN, f"Golden 缺少必需字段：{key}"
    assert GOLDEN["golden_version"] == 1
    assert GOLDEN["standard_issue_expected_status"] == "PROVISIONAL"


def test_golden_upstream_basis_matches_repository() -> None:
    """Applicability Gate, not a business regression.

    If the upstream basis changed, the correct response is to re-review and
    re-approve the Golden — not to adjust the software until this passes.
    """
    mismatches = _applicability_mismatches()
    assert not mismatches, (
        "Golden 上游依据已变化，需要重新复核/批准：\n  " + "\n  ".join(mismatches)
    )


def test_standard_issue_stays_provisional() -> None:
    assert GOLDEN["standard_issue_ref"] == "ECQ-STD-GB29446-001"
    assert _issue_status("ECQ-STD-GB29446-001") == "PROVISIONAL"


def test_golden_does_not_claim_official_interpretation() -> None:
    marker = "无发布机构正式解释"
    trap = case_by_id("gb29446-coking-full-value-trap")
    assert marker in trap["confirmation_degree"]
    assert "ECQ-STD-GB29446-001" in trap["software_interpretation_ref"]


# ===========================================================================
# Business Data Summary — business fact drift detection
# ===========================================================================


def test_business_data_summary_matches_definition(definition: dict) -> None:
    summary = GOLDEN["business_data_summary"]
    for coal_type, expected in summary["coefficients_by_coal_type"].items():
        actual = coefficients_of(definition, coal_type)
        assert {key: _dec_text(value) for key, value in actual.items()} == {
            key: _dec_text(value) for key, value in expected.items()
        }, f"系数漂移：{coal_type}"
    for product_id, expected in summary["thresholds_kwh_per_t"].items():
        product = next(p for p in definition["products"] if p["id"] == product_id)
        raw = product["indicators"][0]["thresholds"]
        actual = {key: _dec_text(raw[key]["value"]) for key in expected}
        assert actual == {key: _dec_text(value) for key, value in expected.items()}, (
            f"阈值漂移：{product_id}"
        )
    definition_text = json.dumps(definition, ensure_ascii=False)
    for clause in summary["key_clauses"]:
        assert clause.split(" ")[0] in definition_text, f"条款漂移：{clause}"
    for table in summary["key_tables"]:
        assert table in definition_text, f"表号漂移：{table}"


def test_all_twelve_coefficients_are_covered(definition: dict) -> None:
    used: set[tuple[str, str, str]] = set()
    for case in normal_cases():
        inputs = case["inputs"]
        k = coefficient_of(definition, inputs["coal_type"], inputs["washing_process"])
        assert _dec_text(case["expected"]["k"]) == _dec_text(k), (
            f"Golden 声明的 k 与 Definition 不一致：{case['case_id']}"
        )
        used.add((inputs["coal_type"], inputs["washing_process"], _dec_text(k)))
    recipes = {
        ("炼焦煤", "跳汰", "1.26"),
        ("炼焦煤", "跳汰、浮选联合", "1"),
        ("炼焦煤", "重介", "1.12"),
        ("炼焦煤", "重介、浮选联合", "0.83"),
        ("炼焦煤", "重介、跳汰、浮选联合", "0.78"),
        ("动力煤", "干法选煤", "1.04"),
        ("动力煤", "跳汰", "0.94"),
        ("动力煤", "跳汰、浮选联合", "0.8"),
        ("动力煤", "跳汰、重介联合", "0.85"),
        ("动力煤", "重介", "0.89"),
        ("动力煤", "重介、浮选联合", "0.76"),
        ("动力煤", "重介、跳汰、浮选联合", "0.72"),
    }
    assert len(recipes) == 12
    assert recipes <= used, f"附录A 系数未被覆盖：{sorted(recipes - used)}"


def test_expected_values_are_independently_derivable(definition: dict) -> None:
    """Every expected value must be reproducible from 原文-derived facts alone."""
    for case in normal_cases():
        derived = independently_derive(definition, case)
        assert _dec_text(derived["k"]) == _dec_text(case["expected"]["k"]), case["case_id"]
        assert _dec_text(derived["e_d"]) == _dec_text(case["expected"]["e_d"]), case["case_id"]
        assert derived["grade"] == case["expected"]["grade"], case["case_id"]


def test_coverage_dimensions_are_met() -> None:
    grades = {case["expected"]["grade"] for case in normal_cases()}
    assert {"LEVEL_1", "LEVEL_2", "LEVEL_3", "NOT_QUALIFIED"} <= grades
    assert {case["inputs"]["coal_type"] for case in normal_cases()} == {"炼焦煤", "动力煤"}
    exact_keys = {
        (case["inputs"]["coal_type"], case["expected"]["e_d"])
        for case in normal_cases()
        if case["expected"]["e_d"] in {"5.0", "7.0", "8.5", "2.0", "3.0", "4.5"}
    }
    for required in (
        ("炼焦煤", "5.0"),
        ("炼焦煤", "7.0"),
        ("炼焦煤", "8.5"),
        ("动力煤", "2.0"),
        ("动力煤", "3.0"),
        ("动力煤", "4.5"),
    ):
        assert required in exact_keys, f"缺少 exact threshold 案例：{required}"
    assert {case["inputs"]["raw_coal_input"] for case in normal_cases()} != {"100"}
    trap = case_by_id("gb29446-coking-full-value-trap")
    assert trap["expected"]["e_d"] == "5.0000004"
    assert trap["expected"]["grade"] == "LEVEL_2"


def test_every_case_declares_source_and_purpose() -> None:
    allowed = set(GOLDEN["source_type_legend"])
    seen: set[str] = set()
    for case in GOLDEN["cases"]:
        for key in (
            "case_id",
            "coverage_dimension",
            "inputs",
            "expected",
            "source_type",
            "confirmation_degree",
            "source_clause",
            "source_table",
            "derivation",
            "software_interpretation_ref",
        ):
            assert key in case, f"{case.get('case_id')} 缺少 {key}"
        assert case["case_id"] not in seen, f"case_id 重复：{case['case_id']}"
        seen.add(case["case_id"])
        assert case["source_type"] in allowed, case["case_id"]
        assert case["derivation"].strip(), f"{case['case_id']} 缺少推导"
        assert case["coverage_dimension"].strip(), f"{case['case_id']} 未说明关闭什么风险"


def test_normal_case_count_is_twelve_coefficient_clusters_plus_trap() -> None:
    """12 coefficient clusters + 1 full-value trap case = 13 normal cases.

    The 12 clusters cover 附录A 的全部 12 个 k（炼焦煤 5 + 动力煤 7）以及 6 个
    exact threshold。第 13 个是 full-value trap：它必须使用非阈值精确值的
    ``500.00004 × 1.00 ÷ 100 = 5.0000004``，无法并入任何一个 exact-threshold
    案例，因此是满足任务书全部覆盖要求所需的最小案例数。
    """
    cases = normal_cases()
    assert len(cases) == 13, f"正常案例应为 12 个系数簇 + 1 个 trap，实际 {len(cases)}"
    cluster_keys = {(case["inputs"]["coal_type"], case["inputs"]["washing_process"]) for case in cases}
    assert len(cluster_keys) == 12, f"系数簇数量应为 12，实际 {len(cluster_keys)}"
    trap = case_by_id("gb29446-coking-full-value-trap")
    assert trap in cases


# ===========================================================================
# Gate A — Business Golden
# ===========================================================================


def assert_matches_golden(projection: dict, case: dict) -> None:
    expected = case["expected"]
    indicator = projection["indicators"][0]
    assert projection["standard_id"] == GOLDEN["standard_id"]
    assert projection["standard_version"] == GOLDEN["standard_version"]
    assert projection["rule_revision"] == GOLDEN["rule_revision"]
    assert projection["numeric_profile_id"] == GOLDEN["numeric_profile_id"]
    assert projection["numeric_contract_version"] == GOLDEN["numeric_contract_version"]
    assert projection["process"] == case["inputs"]["washing_process"]
    assert _dec_text(indicator["k"]) == _dec_text(expected["k"]), f"k 不符：{case['case_id']}"
    assert _dec_text(indicator["actual_value"]) == _dec_text(expected["e_d"]), (
        f"e_d 不符：{case['case_id']}"
    )
    assert indicator["grade"] == expected["grade"], f"等级不符：{case['case_id']}"
    assert indicator["unit"] == GOLDEN["business_data_summary"]["unit"]
    assert indicator["warnings_empty"] is True, f"正常案例不应有 warning：{case['case_id']}"
    assert indicator["has_grade_comparison"] is True
    assert indicator["has_explicit_rounding_step"] is False
    assert projection["round_places_declared"] is False
    limits = GOLDEN["business_data_summary"]["thresholds_kwh_per_t"][projection["product_id"]]
    assert {
        key.upper(): _dec_text(value) for key, value in indicator["corrected_thresholds"].items()
    } == {key.upper(): _dec_text(value) for key, value in limits.items()}
    sources = {tuple(entry) for entry in indicator["sources"]}
    assert any(entry[0] == "GB 29446-2019" for entry in sources), sources
    assert any(entry[2] in {"表1", "表2"} for entry in sources), sources
    assert any(entry[2] == "表A.1" for entry in sources), sources


@pytest.mark.parametrize("case", normal_cases(), ids=lambda c: c["case_id"])
def test_gate_a_business_golden(ctx, case) -> None:
    standard = ctx.application.get_published_standard(STANDARD_ID)
    request = build_request(case, standard)
    result = ctx.application.evaluate(request)
    projection = business_projection(result, request)
    assert_matches_golden(projection, case)
    compared_projection(projection)  # keeps the whitelist honest/importable


@pytest.mark.parametrize("case", safety_cases(), ids=lambda c: c["case_id"])
def test_gate_a_safety_semantics(ctx, case) -> None:
    """Only freezes "must not produce any official grade".

    Whether the attempt is persisted as an INCOMPLETE record is deliberately not
    frozen here.
    """
    standard = ctx.application.get_published_standard(STANDARD_ID)
    request = build_request(case, standard)
    result = ctx.application.evaluate(request)
    grade = str(result.results[0].grade)
    expected = case["expected"]
    assert grade in expected["grade_in"], f"{case['case_id']}: {grade}"
    for forbidden in expected["must_not_be_any_of"]:
        assert grade != forbidden, f"{case['case_id']} 不得产生 {forbidden}"


# ===========================================================================
# Gate B — Product Lifecycle
# ===========================================================================


def test_gate_b_product_lifecycle_real_restart(ctx, tmp_path) -> None:
    """Formal record → real process restart → history projection == Golden."""
    case = case_by_id("gb29446-coking-grade2-exact-l2")
    standard = ctx.application.get_published_standard(STANDARD_ID)
    result = ctx.application.evaluate(build_request(case, standard))
    evaluation_id = result.evaluation_id
    data_root = tmp_path / "中文数据"
    script = _restart_script(evaluation_id, data_root)
    ctx.database.dispose()

    completed = subprocess.run(
        [sys.executable, "-B", str(script)],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        cwd=str(ROOT),
        timeout=240,
    )
    assert completed.returncode == 0, f"{completed.stdout}\n{completed.stderr}"
    payload = json.loads(completed.stdout.strip().splitlines()[-1])
    assert payload["engine_calls"] == 0, "查看历史不得重新运行 Engine"
    assert payload["restored"] is True
    assert payload["evaluation_id"] == evaluation_id
    # The restored business projection still equals the Golden.
    assert payload["grade"] == case["expected"]["grade"]
    assert _dec_text(payload["actual_value"]) == _dec_text(case["expected"]["e_d"])
    assert _dec_text(payload["k"]) == _dec_text(case["expected"]["k"])
    assert payload["rule_revision"] == GOLDEN["rule_revision"]
    assert payload["numeric_profile_id"] == GOLDEN["numeric_profile_id"]
    assert payload["step_operations_include_grade_comparison"] is True
    assert payload["has_explicit_rounding_step"] is False
    assert payload["rule_snapshot_revision"] == GOLDEN["rule_revision"]


def _restart_script(evaluation_id: str, data_root: Path) -> Path:
    script = data_root / "_rs04_restart_probe.py"
    script.parent.mkdir(parents=True, exist_ok=True)
    script.write_text(
        "import json, sys\n"
        "from pathlib import Path\n"
        "sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))\n"
        "from uebench.bootstrap import create_context\n"
        "from uebench.domain import engine as engine_mod\n"
        "\n"
        "calls = {'n': 0}\n"
        "original = engine_mod.EvaluationEngine.evaluate\n"
        "def spy(self, *a, **k):\n"
        "    calls['n'] += 1\n"
        "    return original(self, *a, **k)\n"
        "engine_mod.EvaluationEngine.evaluate = spy\n"
        "\n"
        f"context = create_context(Path(r'{data_root}'))\n"
        "try:\n"
        f"    loaded = context.application.get_evaluation({evaluation_id!r})\n"
        "    if loaded is None:\n"
        "        print(json.dumps({'restored': False}))\n"
        "        raise SystemExit(0)\n"
        "    request, result, snapshot = loaded\n"
        "    item = result.results[0]\n"
        "    ops = [step.operation for step in item.calculation_trace]\n"
        "    print(json.dumps({\n"
        "        'restored': True,\n"
        "        'evaluation_id': result.evaluation_id,\n"
        "        'grade': str(item.grade),\n"
        "        'actual_value': format(item.actual_value, 'f'),\n"
        "        'k': format(item.display_values['process_factor'], 'f'),\n"
        "        'rule_revision': result.rule_revision,\n"
        "        'numeric_profile_id': result.numeric_profile_id,\n"
        "        'rule_snapshot_revision': snapshot.rule_revision,\n"
        "        'step_operations_include_grade_comparison': 'grade_comparison' in ops,\n"
        "        'has_explicit_rounding_step': any('round' in op.lower() for op in ops),\n"
        "        'engine_calls': calls['n'],\n"
        "    }, ensure_ascii=False))\n"
        "finally:\n"
        "    context.database.dispose()\n",
        encoding="utf-8",
    )
    return script


# ===========================================================================
# Gate C — Excel Golden
# ===========================================================================


EXCEL_DATA_ROWS = {
    "企业名称": 2,
    "评价日期": 3,
    "核算周期": 4,
    "自定义周期": 5,
    "煤种": 6,
    "选煤工艺": 7,
    "统计期选煤电力消耗量 E_d": 8,
    "统计期入选原煤量 m": 9,
    "备注": 10,
}


def fill_excel(path: Path, case: dict, *, evaluation_date: date) -> Path:
    inputs = case["inputs"]
    workbook = load_workbook(path)
    sheet = workbook["评价数据"]
    sheet.cell(row=EXCEL_DATA_ROWS["企业名称"], column=2, value="Golden 验证企业")
    sheet.cell(row=EXCEL_DATA_ROWS["评价日期"], column=2, value=evaluation_date)
    sheet.cell(row=EXCEL_DATA_ROWS["核算周期"], column=2, value="全年")
    sheet.cell(row=EXCEL_DATA_ROWS["煤种"], column=2, value=inputs["coal_type"])
    sheet.cell(row=EXCEL_DATA_ROWS["选煤工艺"], column=2, value=inputs["washing_process"])
    sheet.cell(
        row=EXCEL_DATA_ROWS["统计期选煤电力消耗量 E_d"],
        column=2,
        value=str(inputs["electricity_consumption"]),
    )
    sheet.cell(
        row=EXCEL_DATA_ROWS["统计期入选原煤量 m"],
        column=2,
        value=str(inputs["raw_coal_input"]),
    )
    sheet.cell(row=EXCEL_DATA_ROWS["备注"], column=2, value="RS04 Product Golden v1")
    workbook.save(path)
    return path


@pytest.mark.parametrize(
    "case_id",
    ["gb29446-power-grade2-exact-l2", "gb29446-coking-full-value-trap"],
)
def test_gate_c_excel_golden(ctx, tmp_path, case_id) -> None:
    """One ordinary Golden case and one full-value boundary case via Excel."""
    case = case_by_id(case_id)
    standard = ctx.application.get_published_standard(STANDARD_ID)
    # The Excel template exposes the nine user fields; build the GUI-side
    # reference from the same business inputs that can reach the workbook.
    gui_request = build_request(
        case, standard, include_single_process_flag=False, include_enterprise_status=False
    )

    path = tmp_path / f"{case_id}.xlsx"
    ctx.application.create_template(path, STANDARD_ID)
    fill_excel(path, case, evaluation_date=gui_request.evaluation_date)
    report = ctx.application.validate_workbook(path)
    assert report.valid, report.issues
    excel_request = report.request

    # GUI Request == Excel Request on every field the template exposes.
    assert excel_request.evaluation_date == gui_request.evaluation_date
    assert excel_request.standard_id == gui_request.standard_id
    assert excel_request.product_id == gui_request.product_id
    assert excel_request.selection_mode == gui_request.selection_mode
    assert excel_request.input_mode == gui_request.input_mode
    assert excel_request.organization_name == gui_request.organization_name
    assert {k: v.value for k, v in excel_request.inputs.items()} == {
        k: v.value for k, v in gui_request.inputs.items()
    }

    excel_result = ctx.application.evaluate_workbook(report.import_id)
    gui_result = ctx.application.evaluate(gui_request)
    excel_projection = business_projection(excel_result, excel_request)
    gui_projection = business_projection(gui_result, gui_request)

    assert_matches_golden(excel_projection, case)
    assert_matches_golden(gui_projection, case)
    assert compared_projection(excel_projection) == compared_projection(gui_projection)
    # identity/timestamp are the only permitted differences
    assert excel_result.evaluation_id != gui_result.evaluation_id


def test_gate_c_excel_request_equals_gui_form_request(ctx, tmp_path) -> None:
    """The desktop form and the Excel adapter must produce the same canonical request.

    The form additionally collects the optional product-level flag; because the
    engine exempts it from GB29446 validation, the business projection is still
    identical, and this test proves that equivalence explicitly.
    """
    case = case_by_id("gb29446-coking-grade2-exact-l2")
    standard = ctx.application.get_published_standard(STANDARD_ID)
    gui_with_flag = build_request(case, standard, include_single_process_flag=True)
    gui_without_flag = build_request(
        case, standard, include_single_process_flag=False, include_enterprise_status=False
    )

    path = tmp_path / "parity.xlsx"
    ctx.application.create_template(path, STANDARD_ID)
    fill_excel(path, case, evaluation_date=EVALUATION_DATE)
    report = ctx.application.validate_workbook(path)
    assert report.valid, report.issues
    excel_request = report.request

    assert {k: v.value for k, v in excel_request.inputs.items()} == {
        k: v.value for k, v in gui_without_flag.inputs.items()
    }

    excel_result = ctx.application.evaluate_workbook(report.import_id)
    with_flag = ctx.application.evaluate(gui_with_flag)
    without_flag = ctx.application.evaluate(gui_without_flag)
    for result, request in ((with_flag, gui_with_flag), (without_flag, gui_without_flag)):
        assert compared_projection(business_projection(result, request)) == compared_projection(
            business_projection(excel_result, excel_request)
        )


def test_gate_c_excel_ingress_safety(ctx, tmp_path) -> None:
    """XLSX numeric/formula cells must never become authoritative Decimal input."""
    case = case_by_id("gb29446-power-grade2-exact-l2")
    safety = GOLDEN["excel_safety_semantics"]

    path = tmp_path / "numeric-cell.xlsx"
    ctx.application.create_template(path, STANDARD_ID)
    workbook = load_workbook(path)
    sheet = workbook["评价数据"]
    sheet.cell(row=EXCEL_DATA_ROWS["企业名称"], column=2, value="Golden 验证企业")
    sheet.cell(row=EXCEL_DATA_ROWS["评价日期"], column=2, value=EVALUATION_DATE)
    sheet.cell(row=EXCEL_DATA_ROWS["核算周期"], column=2, value="全年")
    sheet.cell(row=EXCEL_DATA_ROWS["煤种"], column=2, value=case["inputs"]["coal_type"])
    sheet.cell(row=EXCEL_DATA_ROWS["选煤工艺"], column=2, value=case["inputs"]["washing_process"])
    sheet.cell(
        row=EXCEL_DATA_ROWS["统计期选煤电力消耗量 E_d"],
        column=2,
        value=375,  # numeric int cell
    )
    sheet.cell(row=EXCEL_DATA_ROWS["统计期入选原煤量 m"], column=2, value="117.5")
    workbook.save(path)
    report = ctx.application.validate_workbook(path)
    assert not report.valid, "XLSX numeric cell 不得成为 authoritative 输入"

    formula_path = tmp_path / "formula-cell.xlsx"
    ctx.application.create_template(formula_path, STANDARD_ID)
    workbook = load_workbook(formula_path)
    sheet = workbook["评价数据"]
    sheet.cell(row=EXCEL_DATA_ROWS["企业名称"], column=2, value="Golden 验证企业")
    sheet.cell(row=EXCEL_DATA_ROWS["评价日期"], column=2, value=EVALUATION_DATE)
    sheet.cell(row=EXCEL_DATA_ROWS["核算周期"], column=2, value="全年")
    sheet.cell(row=EXCEL_DATA_ROWS["煤种"], column=2, value=case["inputs"]["coal_type"])
    sheet.cell(row=EXCEL_DATA_ROWS["选煤工艺"], column=2, value=case["inputs"]["washing_process"])
    sheet.cell(
        row=EXCEL_DATA_ROWS["统计期选煤电力消耗量 E_d"], column=2, value="=300+75"
    )
    sheet.cell(row=EXCEL_DATA_ROWS["统计期入选原煤量 m"], column=2, value="117.5")
    workbook.save(formula_path)
    report = ctx.application.validate_workbook(formula_path)
    assert not report.valid, "Excel 公式不得成为 authoritative 输入"
    assert safety["must_accept"] == ["decimal lexical text"]


def test_gate_c_decimal_lexical_text_is_accepted(ctx, tmp_path) -> None:
    case = case_by_id("gb29446-power-grade2-exact-l2")
    path = tmp_path / "lexical.xlsx"
    ctx.application.create_template(path, STANDARD_ID)
    fill_excel(path, case, evaluation_date=EVALUATION_DATE)
    report = ctx.application.validate_workbook(path)
    assert report.valid, report.issues
    assert report.request.inputs["electricity_consumption"].value == "375"
